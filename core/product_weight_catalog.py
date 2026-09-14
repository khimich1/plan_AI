#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Справочник весов ФБС/ЛС/ЛМ в plita.db: парсинг 1С-xlsx, upsert, resolve."""

from __future__ import annotations

import logging
import os
import re
import sqlite3
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone
from typing import Optional, Sequence

from openpyxl import load_workbook

_log = logging.getLogger(__name__)

DENSITY_MIN_KG_M3 = 1500.0
DENSITY_MAX_KG_M3 = 3500.0
PRODUCT_WEIGHT_TYPES = frozenset({"fbs", "steps", "marches"})

_NAME_PREFIX_RE = re.compile(
    r"^(?:блоки|лестничные\s+ступени|лестничные\s+марши|лестничные\s+площадки)\s+",
    re.IGNORECASE | re.UNICODE,
)
_LS_PREFIX_RE = re.compile(r"^(?:LS|LС|ЛS)", re.IGNORECASE | re.UNICODE)
_LS_FAMILY_RE = re.compile(r"^ЛC-?(\d+)", re.IGNORECASE | re.UNICODE)
_ONE_C_CODE_RE = re.compile(r"^\d{2}-\d+$")
_KB_MARK_RE = re.compile(r"\(КБ|\bКБ[\s\-]?\d", re.IGNORECASE | re.UNICODE)
_LP_MARK_RE = re.compile(r"\d*ЛП", re.IGNORECASE | re.UNICODE)
_LS_MARK_RE = re.compile(r"\bЛС", re.IGNORECASE | re.UNICODE)
_LM_MARK_RE = re.compile(r"\d*ЛМ", re.IGNORECASE | re.UNICODE)


@dataclass(frozen=True)
class ProductWeightRecord:
    mark: str
    mark_norm: str
    product_type: str
    weight_kg: float
    volume_m3: Optional[float]
    display_name: str


@dataclass(frozen=True)
class ParseIssue:
    display_name: str
    reason: str
    mark: Optional[str] = None


@dataclass(frozen=True)
class ProductWeightParseResult:
    records: list[ProductWeightRecord]
    quarantine: list[ParseIssue]
    skipped: list[ParseIssue]


def normalize_product_mark(mark: str) -> str:
    """Ключ lookup: C↔С, T↔Т, B↔В, без пробелов, запятая→точка."""
    text = str(mark or "").strip().upper().replace("Ё", "Е")
    ls_prefix = _LS_PREFIX_RE.match(text)
    if ls_prefix:
        text = "ЛС" + text[ls_prefix.end() :]
    text = text.replace("С", "C").replace("В", "B").replace("Т", "T")
    text = text.replace(",", ".")
    return re.sub(r"\s+", "", text)


def extract_product_mark(display_name: str) -> str:
    """Короткая марка: «Блоки ФБС 24.6.6-Т» → «ФБС 24.6.6-Т»."""
    text = str(display_name or "").strip()
    stripped = _NAME_PREFIX_RE.sub("", text).strip()
    return stripped or text


def classify_product_scope(display_name: str) -> Optional[str]:
    """Тип строки файла: fbs/steps/marches либо kb/lp (вне scope)."""
    text = str(display_name or "").strip().upper().replace("Ё", "Е")
    if not text:
        return None
    if "КРОСС" in text or _KB_MARK_RE.search(text):
        return "kb"
    if _LP_MARK_RE.search(text) or "ПЛОЩАДК" in text:
        return "lp"
    if "ФБС" in text:
        return "fbs"
    if _LS_MARK_RE.search(text) or "СТУПЕН" in text:
        return "steps"
    if _LM_MARK_RE.search(text) or "МАРШ" in text:
        return "marches"
    return None


def _as_float(value: object) -> Optional[float]:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", ".").replace("\u00a0", "")
    if not text or text in {"-", "—", "–"}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _looks_like_1c_code(value: object) -> bool:
    if isinstance(value, bool) or isinstance(value, (int, float)):
        return False
    text = str(value or "").strip()
    if not text or _as_float(value) is not None:
        return False
    return bool(_ONE_C_CODE_RE.match(text)) or text.startswith("00-")


def _cell_text(value: object) -> str:
    return str(value or "").strip().lower().replace("ё", "е")


def _find_header_row(rows: list[tuple]) -> int | None:
    for idx, row in enumerate(rows):
        texts = [_cell_text(cell) for cell in row]
        if any("наименован" in t for t in texts) and any("вес" in t for t in texts):
            return idx
    return None


def _column_index(header: tuple, *needles: str) -> int | None:
    for idx, cell in enumerate(header):
        text = _cell_text(cell)
        if any(needle in text for needle in needles):
            return idx
    return None


def _density_kg_m3(weight_kg: float, volume_m3: Optional[float]) -> Optional[float]:
    if volume_m3 is None or volume_m3 <= 0:
        return None
    return weight_kg / volume_m3


def _density_in_range(weight_kg: float, volume_m3: Optional[float]) -> bool:
    density = _density_kg_m3(weight_kg, volume_m3)
    if density is None:
        return volume_m3 is None
    return DENSITY_MIN_KG_M3 <= density <= DENSITY_MAX_KG_M3


def _skip_reason_for_scope(scope: Optional[str], requested: str) -> str | None:
    if scope == requested:
        return None
    if scope == "kb":
        return "тип вне scope (кросс-блок)"
    if scope == "lp":
        return "тип вне scope (ЛП)"
    if scope is None:
        return "тип вне scope"
    return f"тип вне scope ({scope})"


def _ls_family_keys(mark_norm: str) -> tuple[str, ...]:
    match = _LS_FAMILY_RE.match(mark_norm)
    if not match:
        return ()
    number = match.group(1)
    return (f"ЛC-{number}", f"ЛC{number}")


def parse_product_weights_from_xlsx(
    xlsx_path: str,
    product_type: str,
) -> ProductWeightParseResult:
    """Разобрать выгрузку 1С (Наименование / Вес / Объём / Код) с валидацией."""
    requested = str(product_type or "").strip().lower()
    if requested not in PRODUCT_WEIGHT_TYPES:
        raise ValueError(f"Неизвестный тип продукции: {product_type}")

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        raw_rows = [tuple(row) for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()

    header_idx = _find_header_row(raw_rows)
    if header_idx is None:
        return ProductWeightParseResult(records=[], quarantine=[], skipped=[])

    header = raw_rows[header_idx]
    name_col = _column_index(header, "наименован")
    weight_col = _column_index(header, "вес")
    volume_col = _column_index(header, "объем")
    if name_col is None or weight_col is None:
        return ProductWeightParseResult(records=[], quarantine=[], skipped=[])

    quarantine: list[ParseIssue] = []
    skipped: list[ParseIssue] = []
    candidates: list[ProductWeightRecord] = []

    for row in raw_rows[header_idx + 1 :]:
        if not row:
            continue
        name_val = row[name_col] if name_col < len(row) else None
        display_name = str(name_val).strip() if name_val is not None else ""
        if not display_name or display_name.lower() in {"none", "-"}:
            continue

        weight_val = row[weight_col] if weight_col < len(row) else None
        volume_val = (
            row[volume_col] if volume_col is not None and volume_col < len(row) else None
        )
        weight = _as_float(weight_val)
        if weight is None or weight <= 0:
            skipped.append(ParseIssue(display_name=display_name, reason="нет веса"))
            continue

        if _looks_like_1c_code(volume_val):
            quarantine.append(
                ParseIssue(
                    display_name=display_name,
                    reason="код 1С в колонке объёма",
                    mark=extract_product_mark(display_name),
                )
            )
            continue

        volume = _as_float(volume_val)
        if volume is not None and volume <= 0:
            quarantine.append(
                ParseIssue(
                    display_name=display_name,
                    reason="плотность вне 1500–3500 кг/м³",
                    mark=extract_product_mark(display_name),
                )
            )
            continue
        if not _density_in_range(weight, volume):
            quarantine.append(
                ParseIssue(
                    display_name=display_name,
                    reason="плотность вне 1500–3500 кг/м³",
                    mark=extract_product_mark(display_name),
                )
            )
            continue

        scope = classify_product_scope(display_name)
        skip_reason = _skip_reason_for_scope(scope, requested)
        if skip_reason is not None:
            skipped.append(ParseIssue(display_name=display_name, reason=skip_reason))
            continue

        mark = extract_product_mark(display_name)
        candidates.append(
            ProductWeightRecord(
                mark=mark,
                mark_norm=normalize_product_mark(mark),
                product_type=requested,
                weight_kg=weight,
                volume_m3=volume,
                display_name=display_name,
            )
        )

    records: list[ProductWeightRecord] = []
    grouped: dict[str, list[ProductWeightRecord]] = defaultdict(list)
    for record in candidates:
        grouped[record.mark_norm].append(record)
    for group in grouped.values():
        weights = {round(item.weight_kg, 6) for item in group}
        if len(weights) > 1:
            for item in group:
                quarantine.append(
                    ParseIssue(
                        display_name=item.display_name,
                        reason="дубли марки с разным весом",
                        mark=item.mark,
                    )
                )
            continue
        records.append(group[0])

    return ProductWeightParseResult(
        records=records,
        quarantine=quarantine,
        skipped=skipped,
    )


def ensure_product_weight_schema(db_path: str) -> None:
    """Создать ``product_weight_catalog``, если таблицы ещё нет."""
    conn = sqlite3.connect(os.fspath(db_path))
    try:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS product_weight_catalog (
                mark_norm TEXT PRIMARY KEY,
                product_type TEXT,
                weight_kg REAL NOT NULL,
                volume_m3 REAL,
                display_name TEXT,
                source_file TEXT,
                imported_at TEXT
            )
            """
        )
        conn.commit()
    finally:
        conn.close()


def upsert_product_weights(
    db_path: str,
    records: Sequence[ProductWeightRecord],
    *,
    source_file: str = "",
) -> tuple[int, int]:
    """Upsert по ``mark_norm``. Возвращает (inserted, updated)."""
    ensure_product_weight_schema(db_path)
    imported_at = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    inserted = 0
    updated = 0
    conn = sqlite3.connect(os.fspath(db_path))
    try:
        cur = conn.cursor()
        for record in records:
            cur.execute(
                "SELECT 1 FROM product_weight_catalog WHERE mark_norm = ?",
                (record.mark_norm,),
            )
            exists = cur.fetchone() is not None
            cur.execute(
                """
                INSERT INTO product_weight_catalog (
                    mark_norm, product_type, weight_kg, volume_m3,
                    display_name, source_file, imported_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(mark_norm) DO UPDATE SET
                    product_type = excluded.product_type,
                    weight_kg = excluded.weight_kg,
                    volume_m3 = excluded.volume_m3,
                    display_name = excluded.display_name,
                    source_file = excluded.source_file,
                    imported_at = excluded.imported_at
                """,
                (
                    record.mark_norm,
                    record.product_type,
                    record.weight_kg,
                    record.volume_m3,
                    record.display_name,
                    source_file,
                    imported_at,
                ),
            )
            if exists:
                updated += 1
            else:
                inserted += 1
        conn.commit()
    finally:
        conn.close()
    return inserted, updated


def resolve_product_weight_kg(
    mark: str,
    product_type: str,
    db_path: str,
) -> Optional[float]:
    """Точное совпадение mark_norm; для ЛС — fallback на семейство ЛС-N."""
    key = normalize_product_mark(mark)
    if not key or not db_path:
        return None
    try:
        conn = sqlite3.connect(os.fspath(db_path))
    except sqlite3.Error:
        return None
    try:
        cur = conn.cursor()
        cur.execute(
            "SELECT weight_kg FROM product_weight_catalog WHERE mark_norm = ?",
            (key,),
        )
        row = cur.fetchone()
        if row is not None:
            return float(row[0])
        if str(product_type or "").strip().lower() != "steps":
            return None
        for family_key in _ls_family_keys(key):
            if family_key == key:
                continue
            cur.execute(
                "SELECT weight_kg FROM product_weight_catalog WHERE mark_norm = ?",
                (family_key,),
            )
            family = cur.fetchone()
            if family is None:
                continue
            _log.info(
                "product_weight_catalog: семейный fallback %s → %s",
                mark,
                family_key,
            )
            return float(family[0])
        return None
    except sqlite3.Error:
        return None
    finally:
        conn.close()
