"""Pure Т-13 parser: xlsx workbook → attendance rows (no DB)."""

from __future__ import annotations

import calendar
import io
import re
from dataclasses import dataclass
from datetime import date
from typing import BinaryIO

from openpyxl.utils import column_index_from_string
from openpyxl.workbook.workbook import Workbook
from openpyxl import load_workbook

# Days 1–15 live on row r, days 16–30 on r+2, in these columns (not consecutive).
_CODE_COLUMNS = (
    "I",
    "K",
    "M",
    "N",
    "P",
    "R",
    "T",
    "V",
    "X",
    "Z",
    "AB",
    "AD",
    "AF",
    "AH",
    "AK",
)
_DAY31_COLUMN = "AL"
_CODE_INDEXES = tuple(column_index_from_string(col) for col in _CODE_COLUMNS)
_DAY31_INDEX = column_index_from_string(_DAY31_COLUMN)

_FILENAME_PERIOD = re.compile(r"от\s+(\d{1,2})\.(\d{1,2})\.(\d{4})")


class TimesheetParseError(ValueError):
    """Workbook is not a recognizable Т-13. Service maps this to HTTP 400."""


@dataclass(frozen=True, slots=True)
class TimesheetRow:
    """One person's attendance grid parsed from a Т-13 sheet."""

    full_name: str
    personnel_number: str
    codes: tuple[tuple[date, str], ...]


@dataclass(frozen=True, slots=True)
class UnmatchedTimesheetRow:
    """A named row whose personnel number could not be parsed."""

    full_name: str
    raw_personnel: str


@dataclass(frozen=True, slots=True)
class ParsedTimesheet:
    """Parsed Т-13. ``unmatched`` is rows with a name but no readable number."""

    period: tuple[date, date]
    rows: tuple[TimesheetRow, ...]
    unmatched: tuple[UnmatchedTimesheetRow, ...]


def parse_personnel_number(raw: str | None) -> str | None:
    """'СЭ00-00470' → '470'; 'СЭ0В-00001' → '1'; garbage → None.

    Digits after the first hyphen, leading zeros stripped. The series
    (including a Cyrillic letter such as ``В``) is discarded with the hyphen.
    """
    text = str(raw or "").strip()
    if not text or "-" not in text:
        return None
    tail = text.split("-", 1)[1]
    digits = "".join(ch for ch in tail if ch.isdigit())
    if not digits:
        return None
    return str(int(digits))


def period_from_filename(filename: str | None) -> tuple[date, date] | None:
    """``... от DD.MM.YYYY ...`` → that calendar month. Garbage → None."""
    match = _FILENAME_PERIOD.search(filename or "")
    if match is None:
        return None
    day_n, month_n, year_n = (int(match.group(i)) for i in (1, 2, 3))
    try:
        date(year_n, month_n, day_n)
    except ValueError:
        return None
    last = calendar.monthrange(year_n, month_n)[1]
    return date(year_n, month_n, 1), date(year_n, month_n, last)


def parse_timesheet_xlsx(
    source: bytes | BinaryIO,
    *,
    filename: str = "",
    month: date | None = None,
) -> ParsedTimesheet:
    """Read a Т-13 workbook. Header mismatch raises ``TimesheetParseError``."""
    period = period_from_filename(filename)
    if period is None and month is not None:
        last = calendar.monthrange(month.year, month.month)[1]
        period = (date(month.year, month.month, 1), date(month.year, month.month, last))
    if period is None:
        raise TimesheetParseError(
            "Укажите месяц табеля: в имени файла нет даты «от ДД.ММ.ГГГГ»."
        )

    try:
        workbook = _load(source)
    except TimesheetParseError:
        raise
    except Exception as exc:
        raise TimesheetParseError("Файл не является книгой Excel (.xlsx).") from exc

    sheet = _find_sheet(workbook)
    _require_day_headers(sheet)
    rows, unmatched = _read_people(sheet, period)
    return ParsedTimesheet(period=period, rows=tuple(rows), unmatched=tuple(unmatched))


def _load(source: bytes | BinaryIO) -> Workbook:
    if isinstance(source, (bytes, bytearray)):
        stream: BinaryIO = io.BytesIO(source)
    else:
        stream = source
    return load_workbook(stream, data_only=True, read_only=False)


def _find_sheet(workbook: Workbook):
    for sheet in workbook.worksheets:
        if _contains_tabel(sheet):
            return sheet
    raise TimesheetParseError("Файл не похож на табель Т-13: нет шапки «ТАБЕЛЬ».")


def _contains_tabel(sheet) -> bool:
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is not None and "ТАБЕЛЬ" in str(cell.value).upper():
                return True
    return False


def _require_day_headers(sheet) -> None:
    if _find_header_row(sheet, start_day=1) is None:
        raise TimesheetParseError(
            "Файл не похож на табель Т-13: колонки дней 1–15 не совпадают."
        )
    if _find_header_row(sheet, start_day=16) is None:
        raise TimesheetParseError(
            "Файл не похож на табель Т-13: колонки дней 16–31 не совпадают."
        )


def _find_header_row(sheet, *, start_day: int) -> int | None:
    max_row = int(sheet.max_row or 0)
    for row_idx in range(1, max_row + 1):
        if _row_has_day_numbers(sheet, row_idx, start_day):
            return row_idx
    return None


def _row_has_day_numbers(sheet, row_idx: int, start_day: int) -> bool:
    for offset, col in enumerate(_CODE_INDEXES):
        if _as_int(sheet.cell(row_idx, col).value) != start_day + offset:
            return False
    if start_day == 16:
        return _as_int(sheet.cell(row_idx, _DAY31_INDEX).value) == 31
    return True


def _read_people(sheet, period: tuple[date, date]):
    rows: list[TimesheetRow] = []
    unmatched: list[UnmatchedTimesheetRow] = []
    max_row = int(sheet.max_row or 0)
    consumed_until = 0
    for row_idx in range(1, max_row + 1):
        if row_idx <= consumed_until:
            continue
        if _row_has_day_numbers(sheet, row_idx, 1) or _row_has_day_numbers(sheet, row_idx, 16):
            continue
        name_raw = _cell_text(sheet, row_idx, 3)
        raw_number = _cell_text(sheet, row_idx, 5)
        parsed = parse_personnel_number(raw_number)
        if parsed is None and not _looks_like_name(name_raw):
            continue
        codes = _read_codes(sheet, row_idx, period)
        full_name = _clean_name(name_raw)
        if parsed is None:
            unmatched.append(
                UnmatchedTimesheetRow(full_name=full_name, raw_personnel=raw_number)
            )
        else:
            rows.append(
                TimesheetRow(
                    full_name=full_name,
                    personnel_number=parsed,
                    codes=codes,
                )
            )
        # Hours sit on r+1 and r+3; the second half of the grid is r+2.
        consumed_until = row_idx + 3
    return rows, unmatched


def _read_codes(sheet, row_idx: int, period: tuple[date, date]) -> tuple[tuple[date, str], ...]:
    period_from, period_to = period
    last_day = period_to.day
    found: list[tuple[date, str]] = []
    for offset, col in enumerate(_CODE_INDEXES):
        day_num = offset + 1
        if day_num > last_day:
            break
        code = _code_text(sheet.cell(row_idx, col).value)
        if code:
            found.append((date(period_from.year, period_from.month, day_num), code))
    second = row_idx + 2
    for offset, col in enumerate(_CODE_INDEXES):
        day_num = offset + 16
        if day_num > last_day:
            break
        code = _code_text(sheet.cell(second, col).value)
        if code:
            found.append((date(period_from.year, period_from.month, day_num), code))
    if last_day >= 31:
        code = _code_text(sheet.cell(second, _DAY31_INDEX).value)
        if code:
            found.append((date(period_from.year, period_from.month, 31), code))
    return tuple(found)


def _looks_like_name(text: str) -> bool:
    if not any(ch.isalpha() for ch in text):
        return False
    return " " in text or "\n" in text


def _clean_name(raw: str) -> str:
    return " ".join(raw.replace("\n", " ").split())


def _cell_text(sheet, row_idx: int, col: int) -> str:
    value = sheet.cell(row_idx, col).value
    if value is None:
        return ""
    return str(value).strip()


def _code_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text


def _as_int(value: object) -> int | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value == int(value):
        return int(value)
    text = str(value).strip()
    if text.isdigit():
        return int(text)
    return None
