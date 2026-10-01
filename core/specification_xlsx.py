"""Книга спецификации из бланка № 650. Возвращает байты и файл не пишет."""

from __future__ import annotations

import io
from copy import copy
from dataclasses import dataclass
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.worksheet import Worksheet

from core.commercial_pricing import vat_included_from_line_sums
from core.specification_buyer import buyer_signature_heading
from core.specification_text import (
    SUPPLIER_DIRECTOR_GENITIVE,
    SUPPLIER_DIRECTOR_SHORT,
    SUPPLIER_NAME,
    SpecificationParagraphs,
    date_in_words,
    format_date,
    format_rub,
    total_amount_words,
)

_TEMPLATE = Path(__file__).resolve().parents[1] / "templates" / "specification.xlsx"
_KOPECK = Decimal("0.01")
_SUPPLIER_FULL_NAME = "Общество с ограниченной ответственностью «ЖБК СТАРТ»"
_ITEM_VALUE_COUNT = 8


@dataclass(frozen=True)
class SpecificationHeader:
    invoice_number: str | None
    spec_date: date
    contract_number: str
    contract_date: date
    buyer_preamble: str
    buyer_short_name: str
    buyer_inn: str
    buyer_kpp: str
    buyer_signatory_position: str
    buyer_signatory_short: str
    concrete_grade: str | None = None


@dataclass(frozen=True)
class SpecificationPayableLine:
    name: str
    qty: int
    price: Decimal
    amount: Decimal | None = None


def build_specification_xlsx(
    header: SpecificationHeader,
    lines: list[SpecificationPayableLine],
    paragraphs: SpecificationParagraphs,
) -> bytes:
    """Лист «Лист_1» из бланка. Дата берётся из шапки, сборщик свою не ставит."""
    workbook = load_workbook(_TEMPLATE)
    sheet = workbook["Лист_1"]
    prototype = _replicate_items(sheet, len(lines))
    _fill_title(sheet, header)
    _fill_items(sheet, prototype, lines)
    _fill_totals(sheet, lines)
    _fill_conditions(sheet, paragraphs)
    _fill_buyer(sheet, header)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _replicate_items(sheet: Worksheet, count: int) -> int:
    """Копирует строку-образец. Объединения переносятся явно: insert_rows их не чинит."""
    prototype = _prototype_row(sheet)
    extra = max(count, 1) - 1
    if extra < 1:
        return prototype
    footer_start = prototype + 1
    last_row = sheet.max_row
    prototype_cells = _row_snapshot(sheet, prototype)
    prototype_merges = _single_row_merges(sheet, prototype)
    prototype_height = sheet.row_dimensions[prototype].height
    footer_cells = {
        (row, column): _snap(cell)
        for (row, column), cell in list(sheet._cells.items())
        if row >= footer_start
    }
    footer_merges = [
        CellRange(str(merged))
        for merged in sheet.merged_cells.ranges
        if merged.min_row >= footer_start
    ]
    footer_heights = {
        row: sheet.row_dimensions[row].height for row in range(footer_start, last_row + 1)
    }
    _unmerge_from(sheet, footer_start)
    _drop_cells_from(sheet, footer_start)
    for offset in range(1, extra + 1):
        _paste_row(sheet, prototype + offset, prototype_cells)
        _merge_shifted(sheet, prototype_merges, offset)
        sheet.row_dimensions[prototype + offset].height = prototype_height
    for (row, column), saved in footer_cells.items():
        _paste_cell(sheet, row + extra, column, saved)
    _merge_shifted(sheet, footer_merges, extra)
    for row, height in footer_heights.items():
        sheet.row_dimensions[row + extra].height = height
    return prototype


def _fill_title(sheet: Worksheet, header: SpecificationHeader) -> None:
    _find_containing(sheet, "Спецификация по счету").value = _title(header)
    city = _find_exact(sheet, "г. Кострома")
    date_cell = _right_merge_anchor(sheet, city.row, city.column)
    date_cell.value = date_in_words(header.spec_date).removesuffix(" года") + "\u00a0г."
    preamble = _find_containing(sheet, "«Поставщик»")
    preamble.value = _preamble(header.buyer_preamble)
    _grow_merged_row(sheet, preamble)


def _fill_items(sheet: Worksheet, prototype: int, lines: list[SpecificationPayableLine]) -> None:
    spans = _single_row_merges(sheet, prototype)
    if len(spans) < _ITEM_VALUE_COUNT:
        raise ValueError("В бланке неполная строка товара")
    for index, line in enumerate(lines):
        line_sum = _line_sum(line)
        values = (
            index + 1,
            line.name,
            line.qty,
            "шт",
            _as_number(line.price),
            22,
            _as_number(vat_included_from_line_sums([line_sum])),
            _as_number(line_sum),
        )
        row = prototype + index
        for span, value in zip(spans, values, strict=True):
            sheet.cell(row, span.min_col).value = value


def _fill_totals(sheet: Worksheet, lines: list[SpecificationPayableLine]) -> None:
    sums = [_line_sum(line) for line in lines]
    total = sum(sums, Decimal("0.00"))
    vat = vat_included_from_line_sums(sums) if sums else Decimal("0.00")
    _write_amount(sheet, "Итого:", total)
    _write_amount(sheet, "В т.ч. НДС (22%):", vat)
    _write_amount(sheet, "Итого с НДС:", total)
    words = _find_containing(sheet, "Всего наименований")
    words.value = _totals_sentence(len(lines), total, vat)
    _grow_merged_row(sheet, words)


def _fill_conditions(sheet: Worksheet, paragraphs: SpecificationParagraphs) -> None:
    for anchor, paragraph in (
        ("Условия оплаты:", paragraphs.payment),
        ("Условия поставки:", paragraphs.term),
        ("Вид поставки:", paragraphs.delivery),
    ):
        cell = _find_containing(sheet, anchor)
        body = (paragraph or "").strip()
        cell.value = f"{anchor} {body}" if body else anchor
        _grow_merged_row(sheet, cell)


def _fill_buyer(sheet: Worksheet, header: SpecificationHeader) -> None:
    buyer = _find_exact(sheet, "Покупатель")
    _write(sheet, buyer.row + 1, buyer.column, buyer_signature_heading(header.buyer_short_name))
    inn = (header.buyer_inn or "").strip()
    kpp = (header.buyer_kpp or "").strip()
    inn_line = f"ИНН {inn}" + (f", КПП {kpp}" if kpp else "")
    _write(sheet, buyer.row + 2, buyer.column, inn_line)
    _write(sheet, buyer.row + 3, buyer.column, (header.buyer_signatory_position or "").strip())
    _buyer_signature_cell(sheet).value = (header.buyer_signatory_short or "").strip()


def _title(header: SpecificationHeader) -> str:
    number = (header.invoice_number or "").strip() or "_____"
    return (
        f"Спецификация по счету № {number} от {format_date(header.spec_date)}г. "
        f"к Договору поставки № {header.contract_number} от {format_date(header.contract_date)}г."
    )


def _preamble(buyer_preamble: str) -> str:
    return (
        f"{_SUPPLIER_FULL_NAME} (далее - {SUPPLIER_NAME}), "
        "именуемое в дальнейшем «Поставщик», "
        f"в лице {SUPPLIER_DIRECTOR_GENITIVE}, действующего на основании Устава, "
        "с одной стороны, и "
        f"{buyer_preamble.strip()}, с другой стороны, вместе именуемые «Стороны», "
        "согласовали настоящую Спецификацию о нижеследующем: "
        "Поставщик обязуется изготовить и отгрузить, а Покупатель принять и оплатить "
        "нижеуказанные изделия из бетона:"
    )


def _totals_sentence(count: int, total: Decimal, vat: Decimal) -> str:
    total_kop = _kopecks(total)
    vat_kop = _kopecks(vat)
    return (
        f"Всего наименований {count}, на сумму {format_rub(total_kop)} "
        f"({total_amount_words(total_kop)}), в т.ч. НДС — {format_rub(vat_kop)} руб."
    )


def _write_amount(sheet: Worksheet, anchor: str, amount: Decimal) -> None:
    label = _find_containing(sheet, anchor)
    _amount_cell(sheet, label).value = _as_number(amount)


def _amount_cell(sheet: Worksheet, label):
    merged = _merge_containing(sheet, label.row, label.column)
    start = merged.max_col + 1 if merged is not None else label.column + 1
    candidates = [
        item
        for item in sheet.merged_cells.ranges
        if item.min_row == label.row and item.max_row == label.row and item.min_col >= start
    ]
    if candidates:
        chosen = min(candidates, key=lambda item: item.min_col)
        return sheet.cell(chosen.min_row, chosen.min_col)
    return sheet.cell(label.row, start)


def _buyer_signature_cell(sheet: Worksheet):
    supplier = _find_exact(sheet, SUPPLIER_DIRECTOR_SHORT)
    chosen = None
    for (row, column), cell in sheet._cells.items():
        if isinstance(cell, MergedCell) or row != supplier.row or column <= supplier.column:
            continue
        bottom = cell.border.bottom
        if cell.alignment.horizontal == "right" and bottom is not None and bottom.style:
            if chosen is None or column > chosen.column:
                chosen = cell
    if chosen is None:
        raise ValueError("В бланке нет подписи покупателя")
    return chosen


def _right_merge_anchor(sheet: Worksheet, row: int, after_column: int):
    chosen = None
    for merged in sheet.merged_cells.ranges:
        if merged.min_row == row and merged.min_col > after_column:
            if chosen is None or merged.min_col > chosen.min_col:
                chosen = merged
    if chosen is None:
        raise ValueError("В бланке нет даты спецификации")
    return sheet.cell(chosen.min_row, chosen.min_col)


def _prototype_row(sheet: Worksheet) -> int:
    total = _find_containing(sheet, "Итого:")
    for row in range(total.row - 1, 0, -1):
        if len(_single_row_merges(sheet, row)) >= 6:
            return row
    raise ValueError("В бланке нет строки товара")


def _find_containing(sheet: Worksheet, needle: str):
    found = None
    for cell in _writable(sheet):
        if isinstance(cell.value, str) and needle in cell.value:
            if found is None or len(cell.value) < len(str(found.value)):
                found = cell
    if found is None:
        raise ValueError(f"В бланке нет ячейки {needle}")
    return found


def _find_exact(sheet: Worksheet, text: str):
    for cell in _writable(sheet):
        if isinstance(cell.value, str) and cell.value.strip() == text:
            return cell
    raise ValueError(f"В бланке нет ячейки {text}")


def _write(sheet: Worksheet, row: int, column: int, value: str) -> None:
    cell = sheet.cell(row, column)
    if isinstance(cell, MergedCell):
        raise ValueError("Значение покупателя попало не в верхнюю левую ячейку объединения")
    cell.value = value


def _writable(sheet: Worksheet):
    return [cell for cell in sheet._cells.values() if not isinstance(cell, MergedCell)]


def _single_row_merges(sheet: Worksheet, row: int) -> list[CellRange]:
    return sorted(
        (
            merged
            for merged in sheet.merged_cells.ranges
            if merged.min_row == row and merged.max_row == row
        ),
        key=lambda merged: merged.min_col,
    )


def _merge_containing(sheet: Worksheet, row: int, column: int):
    for merged in sheet.merged_cells.ranges:
        if merged.min_row <= row <= merged.max_row and merged.min_col <= column <= merged.max_col:
            return merged
    return None


def _row_snapshot(sheet: Worksheet, row: int) -> dict[int, tuple]:
    return {
        column: _snap(cell)
        for (cell_row, column), cell in sheet._cells.items()
        if cell_row == row
    }


def _snap(cell) -> tuple:
    style = copy(cell._style) if cell.has_style else None
    value = None if isinstance(cell, MergedCell) else cell.value
    return value, style


def _paste_row(sheet: Worksheet, row: int, cells: dict[int, tuple]) -> None:
    for column, saved in cells.items():
        _paste_cell(sheet, row, column, saved)


def _paste_cell(sheet: Worksheet, row: int, column: int, saved: tuple) -> None:
    value, style = saved
    cell = sheet.cell(row, column)
    if style is not None:
        cell._style = copy(style)
    if not isinstance(cell, MergedCell):
        cell.value = value


def _merge_shifted(sheet: Worksheet, merges: list[CellRange], offset: int) -> None:
    for merged in merges:
        moved = CellRange(str(merged))
        moved.shift(row_shift=offset)
        sheet.merge_cells(str(moved))


def _unmerge_from(sheet: Worksheet, row: int) -> None:
    for ref in [str(merged) for merged in list(sheet.merged_cells.ranges) if merged.min_row >= row]:
        sheet.unmerge_cells(ref)


def _drop_cells_from(sheet: Worksheet, row: int) -> None:
    for key in [key for key in list(sheet._cells) if key[0] >= row]:
        del sheet._cells[key]


def _grow_merged_row(sheet: Worksheet, cell) -> None:
    merged = _merge_containing(sheet, cell.row, cell.column)
    if merged is None:
        return
    width = 0.0
    for column in range(merged.min_col, merged.max_col + 1):
        width += sheet.column_dimensions[get_column_letter(column)].width or 3.5
    capacity = max(int(width * 0.8), 20)
    lines = 0
    for part in str(cell.value or "").splitlines() or [""]:
        lines += max(1, (len(part) + capacity - 1) // capacity)
    needed = 15 * lines + 4
    rows = range(merged.min_row, merged.max_row + 1)
    current = sum((sheet.row_dimensions[row].height or 15) for row in rows)
    if needed <= current:
        return
    last = merged.max_row
    sheet.row_dimensions[last].height = (sheet.row_dimensions[last].height or 15) + (needed - current)


def _line_sum(line: SpecificationPayableLine) -> Decimal:
    if line.amount is not None:
        return line.amount.quantize(_KOPECK, rounding=ROUND_HALF_UP)
    qty = line.qty if isinstance(line.qty, int) and line.qty > 0 else 0
    return (line.price * Decimal(qty)).quantize(_KOPECK, rounding=ROUND_HALF_UP)


def _kopecks(amount: Decimal) -> int:
    return int((amount * Decimal(100)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def _as_number(amount: Decimal) -> float:
    return float(amount.quantize(_KOPECK, rounding=ROUND_HALF_UP))
