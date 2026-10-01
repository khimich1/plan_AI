"""Подстановка даты и стороны-2 в бланк соглашения об ЭДО."""

from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from core.supply_contract import (
    DEFAULT_SIGNATORY_VERB,
    LEGAL_FORMS_REQUIRING_KPP,
    build_preamble,
)

_TEMPLATE = Path(__file__).resolve().parents[1] / "templates" / "edo_agreement.docx"
_MONTHS = (
    "",
    "января",
    "февраля",
    "марта",
    "апреля",
    "мая",
    "июня",
    "июля",
    "августа",
    "сентября",
    "октября",
    "ноября",
    "декабря",
)
_DATE_SHORT = "{{EDO_DATE_SHORT}}"
_DATE_LONG = "{{EDO_DATE_LONG}}"
_PREAMBLE = "{{SIDE2_PREAMBLE}}"
_REQUISITES = "{{SIDE2_REQUISITES}}"
_PARTY = "Сторона-2"


def buyer_okved_line(okved: str | None) -> str | None:
    code = (okved or "").strip()
    if not code:
        return None
    return f"ОКВЭД {code}"


def side2_preamble(fields: dict[str, Any]) -> str:
    """Преамбула стороны-2. У ИП основание — ОГРНИП, даже если в договоре устав."""
    legal_form = str(fields.get("legal_form") or "")
    if legal_form == "ip":
        verb = str(fields.get("signatory_verb") or DEFAULT_SIGNATORY_VERB)
        named = "именуемая" if verb == "действующей" else "именуемый"
        name = str(fields.get("full_name") or "").strip()
        ogrn = str(fields.get("ogrn") or "").strip()
        return (
            f"Индивидуальный предприниматель {name}, {named} в дальнейшем «{_PARTY}», "
            f"{verb} на основании ОГРНИП {ogrn}"
        )
    return build_preamble(
        legal_form=legal_form,
        full_name=str(fields.get("full_name") or ""),
        signatory_name=str(fields.get("signatory_name") or ""),
        authority_basis=str(fields.get("authority_basis") or "устав"),
        signatory_position=fields.get("signatory_position"),
        signatory_verb=str(fields.get("signatory_verb") or DEFAULT_SIGNATORY_VERB),
        poa_number=fields.get("poa_number"),
        poa_date=fields.get("poa_date"),
        party_title=_PARTY,
    )


def buyer_edo_line(fields: dict[str, Any]) -> str | None:
    operator = str(fields.get("edo_operator") or "").strip()
    edo_id = str(fields.get("edo_id") or "").strip()
    if not operator and not edo_id:
        return None
    if operator and edo_id:
        return f"Оператор ЭДО: {operator}, идентификатор {edo_id}"
    if operator:
        return f"Оператор ЭДО: {operator}"
    return f"Оператор ЭДО: идентификатор {edo_id}"


def side2_requisite_lines(fields: dict[str, Any]) -> list[str]:
    """Реквизиты стороны-2. Сторона-1 в этот список не входит."""
    legal_form = str(fields.get("legal_form") or "")
    full_name = str(fields.get("full_name") or "").strip()
    legal_address = str(fields.get("legal_address") or "").strip()
    lines = [_side2_title(legal_form, full_name)]
    inn = str(fields.get("inn") or "").strip()
    if inn:
        lines.append(f"ИНН {inn}")
    if legal_form in LEGAL_FORMS_REQUIRING_KPP:
        kpp = str(fields.get("kpp") or "").strip()
        if kpp:
            lines.append(f"КПП {kpp}")
    ogrn = str(fields.get("ogrn") or "").strip()
    if ogrn:
        label = "ОГРНИП" if legal_form == "ip" else "ОГРН"
        lines.append(f"{label} {ogrn}")
    if legal_address:
        lines.append(f"Юридический адрес: {legal_address}")
    postal = str(fields.get("postal_address") or "").strip()
    if postal and postal != legal_address:
        lines.append(f"Почтовый адрес: {postal}")
    email = str(fields.get("email") or "").strip()
    if email:
        lines.append(f"E-mail: {email}")
    okved = buyer_okved_line(fields.get("okved"))
    if okved:
        lines.append(okved)
    edo = buyer_edo_line(fields)
    if edo:
        lines.append(edo)
    if legal_form == "ip":
        lines.append("Индивидуальный предприниматель")
    else:
        position = str(fields.get("signatory_position") or "").strip()
        if position:
            lines.append(position)
    signatory = str(fields.get("signatory_name") or "").strip()
    lines.append("________________")
    if signatory:
        lines.append(signatory)
    return lines


def render_edo_agreement_docx(
    fields: dict[str, Any],
    *,
    template_path: Path | None = None,
) -> bytes:
    """Собирает docx. Пункты бланка не переписываются."""
    doc = Document(str(template_path or _TEMPLATE))
    on = _contract_date(fields["contract_date"])
    _replace_token(doc, _DATE_SHORT, f"{on.day:02d}.{on.month:02d}.{on.year}")
    _replace_token(doc, _DATE_LONG, f"«{on.day:02d}» {_MONTHS[on.month]} {on.year}")
    _replace_token(doc, _PREAMBLE, side2_preamble(fields))
    _expand_token(doc, _REQUISITES, side2_requisite_lines(fields))
    leftover = [
        paragraph.text
        for paragraph in _iter_paragraphs(doc)
        if "{{" in paragraph.text
    ]
    if leftover:
        raise ValueError("в соглашении остался плейсхолдер")
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _side2_title(legal_form: str, full_name: str) -> str:
    if legal_form == "ooo":
        return f"Общество с ограниченной ответственностью «{full_name}»"
    if legal_form == "ao":
        return f"Акционерное общество «{full_name}»"
    if legal_form == "ip":
        return f"Индивидуальный предприниматель {full_name}"
    if legal_form == "kfh":
        return f"Крестьянское (фермерское) хозяйство «{full_name}»"
    return full_name


def _contract_date(value: date | str) -> date:
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value)[:10])


def _iter_paragraphs(doc: DocumentObject):
    yield from doc.paragraphs
    for section in doc.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def _replace_span(paragraph: Paragraph, start: int, end: int, replacement: str) -> None:
    cursor = 0
    placed = False
    for run in paragraph.runs:
        run_start = cursor
        run_end = cursor + len(run.text)
        cursor = run_end
        if run_end <= start or run_start >= end:
            continue
        local_start = max(0, start - run_start)
        local_end = min(len(run.text), end - run_start)
        if not placed:
            run.text = run.text[:local_start] + replacement + run.text[local_end:]
            placed = True
        else:
            run.text = run.text[:local_start] + run.text[local_end:]


def _replace_once(paragraph: Paragraph, old: str, new: str) -> bool:
    text = paragraph.text
    start = text.find(old)
    if start < 0:
        return False
    _replace_span(paragraph, start, start + len(old), new)
    return True


def _replace_token(doc: DocumentObject, token: str, value: str) -> None:
    found = False
    for paragraph in _iter_paragraphs(doc):
        while token in paragraph.text:
            if not _replace_once(paragraph, token, value):
                break
            found = True
    if not found:
        raise ValueError(token)


def _set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
        return
    paragraph.add_run(text)


def _insert_paragraph_after(paragraph: Paragraph, text: str) -> Paragraph:
    new_p = paragraph._p.makeelement(qn("w:p"), {})
    paragraph._p.addnext(new_p)
    created = Paragraph(new_p, paragraph._parent)
    run = created.add_run(text)
    if paragraph.runs:
        source = paragraph.runs[0]
        run.bold = source.bold
        run.font.name = source.font.name
        run.font.size = source.font.size
    return created


def _expand_token(doc: DocumentObject, token: str, lines: list[str]) -> None:
    targets = [paragraph for paragraph in _iter_paragraphs(doc) if token in paragraph.text]
    if not targets:
        raise ValueError(token)
    body = lines or [""]
    for paragraph in targets:
        if paragraph.text.strip() != token:
            _replace_once(paragraph, token, " ".join(body))
            continue
        _set_paragraph_text(paragraph, body[0])
        anchor = paragraph
        for line in body[1:]:
            anchor = _insert_paragraph_after(anchor, line)
