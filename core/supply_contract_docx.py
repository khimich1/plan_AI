"""Подстановка реквизитов покупателя в бланк договора поставки."""

from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from core.supply_contract import build_preamble

_TEMPLATE = Path(__file__).resolve().parents[1] / "templates" / "supply_contract.docx"
_MONTHS = (
    "",
    "января",
    "февраля",
    "марта",
    "апреля",
    "мая",
    "июня",
    "июля",
    "августа",
    "сентября",
    "октября",
    "ноября",
    "декабря",
)
_BUYER_PLACEHOLDER = "Общество с ограниченной ответственностью ,"
_TITLE_BLANK = "ДОГОВОР N /26"
_HEADER_NUMBER_BLANK = "№ /26"
_EMAIL_LABEL = "адрес Покупателя по e-mail:"


def render_supply_contract_docx(
    fields: dict[str, Any],
    *,
    template_path: Path | None = None,
) -> bytes:
    """Собирает docx из сохранённых полей. Блок поставщика и М-2 не меняются."""
    doc = Document(str(template_path or _TEMPLATE))
    number = str(fields["number"])
    on = _contract_date(fields["contract_date"])
    _fill_number(doc, number)
    _fill_dates(doc, on)
    _fill_preamble(doc, fields)
    _fill_buyer_email(doc, str(fields["email"]))
    _fill_buyer_block(doc, fields)
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _contract_date(value: date | str) -> date:
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value)[:10])


def _iter_paragraphs(doc: DocumentObject):
    yield from doc.paragraphs
    for section in doc.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def _replace_span(paragraph: Paragraph, start: int, end: int, replacement: str) -> None:
    cursor = 0
    placed = False
    for run in paragraph.runs:
        run_start = cursor
        run_end = cursor + len(run.text)
        cursor = run_end
        if run_end <= start or run_start >= end:
            continue
        local_start = max(0, start - run_start)
        local_end = min(len(run.text), end - run_start)
        if not placed:
            run.text = run.text[:local_start] + replacement + run.text[local_end:]
            placed = True
        else:
            run.text = run.text[:local_start] + run.text[local_end:]


def _replace_once(paragraph: Paragraph, old: str, new: str) -> bool:
    text = paragraph.text
    start = text.find(old)
    if start < 0:
        return False
    _replace_span(paragraph, start, start + len(old), new)
    return True


def _fill_number(doc: DocumentObject, number: str) -> None:
    for paragraph in _iter_paragraphs(doc):
        if _TITLE_BLANK in paragraph.text:
            _replace_once(paragraph, _TITLE_BLANK, f"ДОГОВОР {number}")
        elif _HEADER_NUMBER_BLANK in paragraph.text:
            _replace_once(paragraph, _HEADER_NUMBER_BLANK, f"№ {number}")


def _fill_dates(doc: DocumentObject, on: date) -> None:
    day = f"{on.day:02d}"
    month_name = _MONTHS[on.month]
    numeric = f"{day}.{on.month:02d}"
    for paragraph in _iter_paragraphs(doc):
        text = paragraph.text
        if "«" in text and "2026" in text and "года" in text:
            start = text.find("«")
            year_at = text.find("2026", start)
            if start >= 0 and year_at > start:
                _replace_span(paragraph, start, year_at, f"«{day}» {month_name} ")
        if "от .2026" in paragraph.text:
            _replace_once(paragraph, "от .2026", f"от {numeric}.2026")


def _fill_preamble(doc: DocumentObject, fields: dict[str, Any]) -> None:
    preamble = build_preamble(
        legal_form=str(fields["legal_form"]),
        full_name=str(fields["full_name"]),
        signatory_name=str(fields["signatory_name"]),
        authority_basis=str(fields["authority_basis"]),
        signatory_position=fields.get("signatory_position"),
        signatory_verb=str(fields.get("signatory_verb") or "действующего"),
        poa_number=fields.get("poa_number"),
        poa_date=fields.get("poa_date"),
    )
    for paragraph in doc.paragraphs:
        text = paragraph.text
        start = text.find(_BUYER_PLACEHOLDER)
        other = text.find(", с другой", start if start >= 0 else 0)
        if start < 0 or other < 0:
            continue
        _replace_span(paragraph, start, other, preamble)
        return


def _fill_buyer_email(doc: DocumentObject, email: str) -> None:
    for paragraph in doc.paragraphs:
        if _EMAIL_LABEL not in paragraph.text:
            continue
        label_end = paragraph.text.find(_EMAIL_LABEL) + len(_EMAIL_LABEL)
        tail = paragraph.text[label_end:].strip()
        if tail:
            return
        if paragraph.runs:
            paragraph.runs[-1].text = paragraph.runs[-1].text.rstrip() + " " + email
        return


def _buyer_lines(fields: dict[str, Any]) -> list[str]:
    short = str(fields["short_name"]).strip()
    legal_form = str(fields["legal_form"])
    heading = {
        "ooo": f"ООО «{short}»",
        "ao": f"АО «{short}»",
        "ip": f"ИП {short}",
        "kfh": f"КФХ «{short}»",
        "person": short,
    }.get(legal_form, short)
    kpp = str(fields.get("kpp") or "").strip()
    kpp_part = f" КПП {kpp}" if kpp else ""
    lines = [
        heading,
        str(fields["full_name"]).strip(),
        f"ИНН {fields['inn']}{kpp_part} ОГРН {fields['ogrn']}",
        f"Место нахождения: {fields['legal_address']}",
    ]
    postal = str(fields.get("postal_address") or "").strip()
    if postal and postal != str(fields["legal_address"]).strip():
        lines.append(f"Почтовый адрес: {postal}")
    phone = str(fields.get("phone") or "").strip()
    contact = f"e-mail: {fields['email']}"
    if phone:
        contact = f"телефон: {phone}, {contact}"
    lines.extend(
        [
            f"расчётный счёт {fields['account']} в {fields['bank_name']}",
            f"кор/счёт {fields['corr_account']}, БИК {fields['bik']}",
            contact,
        ]
    )
    edo_operator = str(fields.get("edo_operator") or "").strip()
    edo_id = str(fields.get("edo_id") or "").strip()
    if edo_operator or edo_id:
        lines.append(
            f"оператор ЭДО – {edo_operator}, идентификатор участника ЭДО - {edo_id}".rstrip()
        )
    position = str(fields.get("signatory_position") or "").strip()
    if position:
        lines.append(position)
    lines.append(str(fields["signatory_name"]).strip())
    return lines


def _set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
        return
    paragraph.add_run(text)


def _fill_buyer_block(doc: DocumentObject, fields: dict[str, Any]) -> None:
    paragraphs = list(doc.paragraphs)
    start = next(
        (index for index, paragraph in enumerate(paragraphs) if paragraph.text.strip() == "ПОКУПАТЕЛЬ:"),
        None,
    )
    if start is None:
        return
    lines = _buyer_lines(fields)
    blanks = [
        paragraph
        for paragraph in paragraphs[start + 1 :]
        if not paragraph.text.strip()
    ]
    # Стоп перед приложением: пустые абзацы только до следующего непустого.
    limited: list[Paragraph] = []
    for paragraph in paragraphs[start + 1 :]:
        if paragraph.text.strip():
            break
        limited.append(paragraph)
    targets = limited or blanks
    for paragraph, line in zip(targets, lines):
        _set_paragraph_text(paragraph, line)
    if len(lines) <= len(targets):
        return
    anchor = targets[-1] if targets else paragraphs[start]
    for line in lines[len(targets) :]:
        anchor = _insert_paragraph_after(anchor, line)


def _insert_paragraph_after(paragraph: Paragraph, text: str) -> Paragraph:
    new_p = paragraph._p.makeelement(qn("w:p"), {})
    paragraph._p.addnext(new_p)
    created = Paragraph(new_p, paragraph._parent)
    created.add_run(text)
    return created
