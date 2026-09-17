#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Импорт цены за штуку из Excel-очереди GPS-001 в прайс-таблицы pb.db.

Читает только блок «💰 ввести цену» листа «Задачи». Колонка
«Цена за шт (ввод)» — единственный источник цены; подсказка
«как у ближайшего соседа» не подставляется. Сваи: одна цена на все
классы бетона, уже лежащие у марки (или все известные, если марки нет).
"""

from __future__ import annotations

import math
import re
import sqlite3
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Callable, Optional

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from core.bridge_pile_price_db import (
    BRIDGE_PILE_GRADE_CODES,
    init_bridge_pile_prices_schema,
    normalize_bridge_pile_mark_for_lookup,
)
from core.fbs_price_db import (
    FBS_GRADE_CODES,
    init_fbs_prices_schema,
    normalize_fbs_mark_for_lookup,
)
from core.march_price_db import GRADE_CODES as MARCH_GRADE_CODES
from core.march_price_db import init_march_prices_schema, normalize_march_mark
from core.pile_catalog import normalize_pile_mark_key
from core.pile_price_db import GRADE_CODES as PILE_GRADE_CODES
from core.pile_price_db import init_pile_prices_schema
from core.price_db import _connect
from core.step_price_db import (
    extract_step_mark,
    init_step_prices_schema,
    normalize_step_mark,
)

TASKS_SHEET = "Задачи"
PRICE_SECTION_TITLE = "💰 ввести цену"
PRICE_INPUT_HEADER = "Цена за шт (ввод)"
GROUP_HEADER = "Группа"
NAME_1C_HEADER = "Наименование 1С"
GUID_HEADER = "GUID 1С"

REQUIRED_HEADERS = (GROUP_HEADER, NAME_1C_HEADER, PRICE_INPUT_HEADER)

UNIT_PRICE_MAX = 10_000_000.0

GROUP_TITLE_TO_KIND = {
    "сваи": "pile",
    "сваи мостовые": "bridge_pile",
    "блоки фбс": "fbs",
    "марши": "stair_flight",
    "ступени": "stair_step",
}

KIND_KNOWN_GRADES = {
    "pile": PILE_GRADE_CODES,
    "bridge_pile": BRIDGE_PILE_GRADE_CODES,
    "fbs": FBS_GRADE_CODES,
    "stair_flight": MARCH_GRADE_CODES,
}

_PILE_TOKEN_RE = re.compile(
    r"([СC]\s*\d+(?:[.,]\d+)?\s*\.\s*\d+(?:\s*-\s*[^\s,;]+)?)",
    re.I,
)
_PILE_DOT_LOAD_RE = re.compile(
    r"([СC]\s*\d+(?:[.,]\d+)?)\.(20|30|35|40|45)\.(\d+)",
    re.I,
)
_BRIDGE_1C_BODY_RE = re.compile(r"^с(\d+)\.(\d+)-([твt])(\d+)$")
_BRIDGE_OUR_RE = re.compile(
    r"^[CcСс]\s*(\d+)[-\.](\d+)\s*([TТtтBВbв])\s*(\d+)$"
)
_FBS_MARK_RE = re.compile(r"ФБС\s*\d", re.I)
_STEP_TOKEN_RE = re.compile(r"ЛС\s*\S+", re.I | re.UNICODE)

KeyFn = Callable[[str], str]


class QueueFormatError(ValueError):
    """Лист «Задачи» не похож на очередь GPS-001."""


@dataclass(frozen=True)
class QueuePriceRow:
    row_index: int
    group: str
    name_1c: str
    guid_1c: Optional[str]
    price_raw: object


@dataclass(frozen=True)
class PriceImportError:
    row_index: int
    group: str
    name_1c: str
    mark: Optional[str]
    reason: str


@dataclass(frozen=True)
class PriceImportWrite:
    product_kind: str
    mark: str
    grades: tuple[str, ...]
    price: float
    created: bool


@dataclass(frozen=True)
class PriceImportReport:
    written: tuple[PriceImportWrite, ...] = ()
    errors: tuple[PriceImportError, ...] = ()
    skipped_empty: int = 0


@dataclass
class _KindIndex:
    by_key: dict[str, str]
    grades: dict[str, list[str]]


def parse_queue_price_rows(xlsx_path: str | Path) -> list[QueuePriceRow]:
    """Прочитать блок 💰 с листа «Задачи». 🏭/⚠️/⏸️/🚫/плиты не читаются."""
    path = Path(xlsx_path)
    if not path.is_file():
        raise QueueFormatError(f"Файл очереди не найден: {path}")
    wb = load_workbook(path, data_only=False)
    try:
        if TASKS_SHEET not in wb.sheetnames:
            raise QueueFormatError(f"Нет листа «{TASKS_SHEET}»")
        return _rows_from_price_section(wb[TASKS_SHEET])
    finally:
        wb.close()


def import_queue_prices(
    xlsx_path: str | Path,
    db_path: str | Path,
    *,
    dry_run: bool = False,
) -> PriceImportReport:
    """Валидировать заполненные цены и записать их в прайс-таблицы."""
    queue_rows = parse_queue_price_rows(xlsx_path)
    db = str(db_path)
    _ensure_price_schemas(db)
    conn = _connect(db)
    try:
        indexes = _load_indexes(conn)
        errors: list[PriceImportError] = []
        skipped_empty = 0
        pending: list[PriceImportWrite] = []
        for row in queue_rows:
            outcome = _plan_row(row, indexes)
            if outcome == "empty":
                skipped_empty += 1
                continue
            if isinstance(outcome, PriceImportError):
                errors.append(outcome)
                continue
            pending.append(outcome)
            _remember_write(indexes, outcome)
        if not dry_run:
            imported_at = datetime.now().isoformat(timespec="seconds")
            for item in pending:
                _apply_write(conn, item, imported_at)
            conn.commit()
        return PriceImportReport(
            written=tuple(pending),
            errors=tuple(errors),
            skipped_empty=skipped_empty,
        )
    finally:
        conn.close()


def extract_mark_for_kind(product_kind: str, name_1c: str) -> Optional[str]:
    """Достать марку прайса из наименования 1С (или уже короткой марки)."""
    extractors = {
        "pile": _extract_pile_mark,
        "bridge_pile": _extract_bridge_mark,
        "fbs": _extract_fbs_mark,
        "stair_flight": _extract_march_mark,
        "stair_step": _extract_step_queue_mark,
    }
    extract = extractors.get(product_kind)
    if extract is None:
        return None
    return extract(name_1c)


def apply_unit_price(
    db_path: str | Path,
    product_kind: str,
    mark: str,
    price: float,
) -> PriceImportWrite:
    """Записать цену за штуку в прайс-таблицы (та же семантика, что Excel-импорт)."""
    kind = str(product_kind or "").strip().lower()
    if kind not in KIND_KNOWN_GRADES and kind != "stair_step":
        raise ValueError(f"Неизвестная группа изделий: {product_kind}")
    validated = validate_unit_price(price)
    extracted = extract_mark_for_kind(kind, mark) or str(mark or "").strip()
    if not extracted:
        raise ValueError("неизвестная марка")
    db = str(db_path)
    _ensure_price_schemas(db)
    conn = _connect(db)
    try:
        indexes = _load_indexes(conn)
        stored, created = _resolve_stored_mark(kind, extracted, indexes)
        grades = _grades_for_write(kind, stored, created, indexes)
        item = PriceImportWrite(
            product_kind=kind,
            mark=stored,
            grades=grades,
            price=validated,
            created=created,
        )
        _apply_write(conn, item, datetime.now().isoformat(timespec="seconds"))
        conn.commit()
        return item
    finally:
        conn.close()


def validate_unit_price(value: object) -> float:
    """Цена > 0, конечная, не выше потолка 10_000_000."""
    parsed = _parse_unit_price(value)
    if parsed == "empty":
        raise ValueError("укажите цену")
    if parsed == "invalid":
        raise ValueError("некорректная цена")
    price = float(parsed)
    if price <= 0:
        raise ValueError("цена должна быть > 0")
    if not math.isfinite(price):
        raise ValueError("некорректная цена")
    if price > UNIT_PRICE_MAX:
        raise ValueError(f"цена выше потолка {int(UNIT_PRICE_MAX)}")
    return price


def format_report(report: PriceImportReport, *, dry_run: bool = False) -> str:
    prefix = "dry-run (запись не выполнялась)\n" if dry_run else ""
    lines = [
        prefix + f"written: {len(report.written)}",
        f"errors: {len(report.errors)}",
        f"skipped_empty: {report.skipped_empty}",
    ]
    for item in report.written:
        grades = ",".join(item.grades) or "—"
        action = "insert" if item.created else "update"
        lines.append(
            f"  {action} {item.product_kind} {item.mark} "
            f"[{grades}] = {item.price}"
        )
    for err in report.errors:
        mark = err.mark or "—"
        lines.append(
            f"  row {err.row_index}: {err.reason} "
            f"({err.group!r} {err.name_1c!r} mark={mark})"
        )
    return "\n".join(lines)


def _rows_from_price_section(ws: Worksheet) -> list[QueuePriceRow]:
    header_row = _find_price_header_row(ws)
    columns = _header_map(ws, header_row)
    missing = [name for name in REQUIRED_HEADERS if name not in columns]
    if missing:
        raise QueueFormatError(
            "В блоке «💰 ввести цену» нет колонок: " + ", ".join(missing)
        )
    rows: list[QueuePriceRow] = []
    for excel_row in range(header_row + 1, ws.max_row + 1):
        group = _cell_text(ws, excel_row, columns[GROUP_HEADER])
        name_1c = _cell_text(ws, excel_row, columns[NAME_1C_HEADER])
        if not group and not name_1c:
            continue
        guid_col = columns.get(GUID_HEADER)
        guid = _cell_text(ws, excel_row, guid_col) if guid_col else ""
        rows.append(
            QueuePriceRow(
                row_index=excel_row,
                group=group,
                name_1c=name_1c,
                guid_1c=guid or None,
                price_raw=ws.cell(excel_row, columns[PRICE_INPUT_HEADER]).value,
            )
        )
    return rows


def _find_price_header_row(ws: Worksheet) -> int:
    title_row: Optional[int] = None
    for row_idx in range(1, ws.max_row + 1):
        value = _cell_text(ws, row_idx, 1)
        if PRICE_SECTION_TITLE in value:
            title_row = row_idx
            break
    if title_row is None:
        raise QueueFormatError(
            f"На листе «{TASKS_SHEET}» нет блока «{PRICE_SECTION_TITLE}»"
        )
    for row_idx in range(title_row + 1, min(title_row + 5, ws.max_row + 1)):
        if _cell_text(ws, row_idx, 1) == GROUP_HEADER:
            return row_idx
    raise QueueFormatError("После «💰 ввести цену» нет строки заголовков")


def _header_map(ws: Worksheet, header_row: int) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        title = _cell_text(ws, header_row, col)
        if title and title not in mapping:
            mapping[title] = col
    return mapping


def _cell_text(ws: Worksheet, row: int, col: int) -> str:
    value = ws.cell(row, col).value
    if value is None:
        return ""
    return str(value).strip()


def _ensure_price_schemas(db_path: str) -> None:
    init_pile_prices_schema(db_path)
    init_bridge_pile_prices_schema(db_path)
    init_fbs_prices_schema(db_path)
    init_march_prices_schema(db_path)
    init_step_prices_schema(db_path)


def _load_indexes(conn: sqlite3.Connection) -> dict[str, _KindIndex]:
    indexes = {
        "pile": _graded_index(conn, "pile_prices", normalize_pile_mark_key),
        "bridge_pile": _graded_index(
            conn, "bridge_pile_prices", normalize_bridge_pile_mark_for_lookup
        ),
        "fbs": _graded_index(conn, "fbs_prices", normalize_fbs_mark_for_lookup),
        "stair_flight": _graded_index(conn, "march_prices", _march_lookup_key),
        "stair_step": _step_index(conn),
    }
    return indexes


def _graded_index(
    conn: sqlite3.Connection,
    table: str,
    key_fn: KeyFn,
) -> _KindIndex:
    by_key: dict[str, str] = {}
    grades: dict[str, list[str]] = {}
    sql = f"SELECT mark, concrete_grade FROM {table}"
    for mark, grade in conn.execute(sql):
        stored = str(mark)
        by_key[key_fn(stored)] = stored
        bucket = grades.setdefault(stored, [])
        code = str(grade)
        if code not in bucket:
            bucket.append(code)
    return _KindIndex(by_key=by_key, grades=grades)


def _step_index(conn: sqlite3.Connection) -> _KindIndex:
    by_key: dict[str, str] = {}
    grades: dict[str, list[str]] = {}
    for (mark,) in conn.execute("SELECT mark FROM step_prices"):
        stored = str(mark)
        by_key[_step_lookup_key(stored)] = stored
        grades[stored] = []
    return _KindIndex(by_key=by_key, grades=grades)


def _plan_row(
    row: QueuePriceRow,
    indexes: dict[str, _KindIndex],
) -> PriceImportWrite | PriceImportError | str:
    parsed = _parse_unit_price(row.price_raw)
    if parsed == "empty":
        return "empty"
    if parsed == "invalid":
        return _error(row, "некорректная цена")
    price = parsed
    if price <= 0:
        return _error(row, "цена должна быть > 0")

    kind = _kind_from_group(row.group)
    if kind is None:
        return _error(row, "неизвестная группа")

    extracted = extract_mark_for_kind(kind, row.name_1c)
    if not extracted:
        return _error(row, "неизвестная марка")

    stored, created = _resolve_stored_mark(kind, extracted, indexes)
    grades = _grades_for_write(kind, stored, created, indexes)
    return PriceImportWrite(
        product_kind=kind,
        mark=stored,
        grades=grades,
        price=price,
        created=created,
    )


def _parse_unit_price(value: object) -> float | str:
    if value is None:
        return "empty"
    if isinstance(value, bool):
        return "invalid"
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("\u00a0", " ")
    if not text:
        return "empty"
    compact = text.replace(" ", "").replace(",", ".")
    try:
        return float(compact)
    except ValueError:
        return "invalid"


def _kind_from_group(group: str) -> Optional[str]:
    key = str(group or "").strip().lower()
    return GROUP_TITLE_TO_KIND.get(key)


def _resolve_stored_mark(
    kind: str,
    extracted: str,
    indexes: dict[str, _KindIndex],
) -> tuple[str, bool]:
    stored = indexes[kind].by_key.get(_key_fn_for_kind(kind)(extracted))
    if stored:
        return stored, False
    return extracted, True


def _grades_for_write(
    kind: str,
    stored_mark: str,
    created: bool,
    indexes: dict[str, _KindIndex],
) -> tuple[str, ...]:
    if kind == "stair_step":
        return ()
    known = tuple(KIND_KNOWN_GRADES[kind])
    if created:
        return known
    existing = tuple(indexes[kind].grades.get(stored_mark) or ())
    return existing or known


def _remember_write(
    indexes: dict[str, _KindIndex],
    item: PriceImportWrite,
) -> None:
    kind_index = indexes[item.product_kind]
    key = _key_fn_for_kind(item.product_kind)(item.mark)
    kind_index.by_key[key] = item.mark
    kind_index.grades[item.mark] = list(item.grades)


def _key_fn_for_kind(kind: str) -> KeyFn:
    return {
        "pile": normalize_pile_mark_key,
        "bridge_pile": normalize_bridge_pile_mark_for_lookup,
        "fbs": normalize_fbs_mark_for_lookup,
        "stair_flight": _march_lookup_key,
        "stair_step": _step_lookup_key,
    }[kind]


def _apply_write(
    conn: sqlite3.Connection,
    item: PriceImportWrite,
    imported_at: str,
) -> None:
    if item.product_kind == "stair_step":
        _write_step(conn, item, imported_at)
        return
    _write_graded(conn, item, imported_at)


def _write_graded(
    conn: sqlite3.Connection,
    item: PriceImportWrite,
    imported_at: str,
) -> None:
    table = {
        "pile": "pile_prices",
        "bridge_pile": "bridge_pile_prices",
        "fbs": "fbs_prices",
        "stair_flight": "march_prices",
    }[item.product_kind]
    for grade in item.grades:
        exists = conn.execute(
            f"SELECT 1 FROM {table} WHERE mark = ? AND concrete_grade = ?",
            (item.mark, grade),
        ).fetchone()
        if exists:
            conn.execute(
                f"UPDATE {table} SET price = ?, imported_at = ? "
                "WHERE mark = ? AND concrete_grade = ?",
                (item.price, imported_at, item.mark, grade),
            )
            continue
        _insert_graded_row(conn, item, grade, imported_at)


def _insert_graded_row(
    conn: sqlite3.Connection,
    item: PriceImportWrite,
    grade: str,
    imported_at: str,
) -> None:
    if item.product_kind == "pile":
        conn.execute(
            "INSERT INTO pile_prices "
            "(mark, concrete_grade, price, price_list_date, imported_at) "
            "VALUES (?, ?, ?, NULL, ?)",
            (item.mark, grade, item.price, imported_at),
        )
        return
    if item.product_kind == "bridge_pile":
        conn.execute(
            "INSERT INTO bridge_pile_prices "
            "(mark, concrete_grade, price, variant_group, display_name, "
            " price_list_date, imported_at) "
            "VALUES (?, ?, ?, NULL, ?, NULL, ?)",
            (item.mark, grade, item.price, item.mark, imported_at),
        )
        return
    if item.product_kind == "fbs":
        conn.execute(
            "INSERT INTO fbs_prices "
            "(mark, concrete_grade, price, display_name, "
            " price_list_date, imported_at) "
            "VALUES (?, ?, ?, ?, NULL, ?)",
            (item.mark, grade, item.price, item.mark, imported_at),
        )
        return
    conn.execute(
        "INSERT INTO march_prices "
        "(mark, concrete_grade, price, display_name, "
        " price_list_date, imported_at) "
        "VALUES (?, ?, ?, ?, NULL, ?)",
        (
            item.mark,
            grade,
            item.price,
            f"Лестничные марши {item.mark}",
            imported_at,
        ),
    )


def _write_step(
    conn: sqlite3.Connection,
    item: PriceImportWrite,
    imported_at: str,
) -> None:
    exists = conn.execute(
        "SELECT 1 FROM step_prices WHERE mark = ?",
        (item.mark,),
    ).fetchone()
    if exists:
        conn.execute(
            "UPDATE step_prices SET price = ?, imported_at = ? WHERE mark = ?",
            (item.price, imported_at, item.mark),
        )
        return
    conn.execute(
        "INSERT INTO step_prices "
        "(mark, price, display_name, price_list_date, imported_at) "
        "VALUES (?, ?, ?, NULL, ?)",
        (item.mark, item.price, item.mark, imported_at),
    )


def _error(row: QueuePriceRow, reason: str) -> PriceImportError:
    return PriceImportError(
        row_index=row.row_index,
        group=row.group,
        name_1c=row.name_1c,
        mark=None,
        reason=reason,
    )


def _extract_pile_mark(name_1c: str) -> Optional[str]:
    text = re.sub(r"^сваи\s+", "", str(name_1c or "").strip(), flags=re.I)
    dashed = _PILE_TOKEN_RE.search(text)
    if dashed and "-" in dashed.group(1):
        return _canon_pile_token(dashed.group(1))
    dotted = _PILE_DOT_LOAD_RE.search(text)
    if dotted:
        token = f"{dotted.group(1)}.{dotted.group(2)}-{dotted.group(3)}"
        return _canon_pile_token(token)
    if dashed:
        return _canon_pile_token(dashed.group(1))
    return None


def _canon_pile_token(token: str) -> str:
    text = re.sub(r"^[CcСс]\s*", "С", token.strip())
    return re.sub(r"\s+", "", text)


def _extract_bridge_mark(name_1c: str) -> Optional[str]:
    raw = str(name_1c or "").strip()
    our = _BRIDGE_OUR_RE.match(raw)
    if our:
        return _bridge_canon(*our.groups())
    compact = re.sub(r"^сваимостовые", "", _compact_cyr(raw))
    body = _BRIDGE_1C_BODY_RE.match(compact)
    if body:
        return _bridge_canon(*body.groups())
    return None


def _bridge_canon(length: str, section: str, kind: str, num: str) -> str:
    letter = "T" if kind.lower() in {"t", "т"} else "B"
    return f"C{length}-{section}{letter}{num}"


def _extract_fbs_mark(name_1c: str) -> Optional[str]:
    text = re.sub(r"^блоки\s+", "", str(name_1c or "").strip(), flags=re.I)
    if not _FBS_MARK_RE.search(text):
        return None
    return text


def _extract_march_mark(name_1c: str) -> Optional[str]:
    mark = normalize_march_mark(name_1c)
    upper = mark.upper()
    if upper.startswith("1ЛМ") or upper.startswith("ЛМ"):
        return mark
    return None


def _extract_step_queue_mark(name_1c: str) -> Optional[str]:
    cleaned = re.sub(
        r"^лестничные\s+ступени\s+",
        "",
        str(name_1c or "").strip(),
        flags=re.I,
    )
    token_match = _STEP_TOKEN_RE.search(cleaned)
    mark = extract_step_mark(cleaned) or extract_step_mark(name_1c)
    if not mark:
        return None
    mark = re.sub(r"^ЛС-", "ЛС", mark)
    if not token_match:
        return mark
    rest = cleaned[token_match.end() :].strip()
    if not rest:
        return mark
    suffix = normalize_step_mark(rest)
    if suffix and not mark.endswith(suffix):
        return f"{mark}{suffix}"
    return mark


def _march_lookup_key(mark: str) -> str:
    return re.sub(r"\s+", "", normalize_march_mark(mark).upper())


def _step_lookup_key(mark: str) -> str:
    return re.sub(r"^ЛС-", "ЛС", normalize_step_mark(mark))


def _compact_cyr(value: str) -> str:
    text = str(value).strip().lower().replace("ё", "е")
    text = text.translate(str.maketrans({"c": "с", "t": "т", "b": "в"}))
    return re.sub(r"\s+", "", text)


__all__ = [
    "GROUP_TITLE_TO_KIND",
    "PRICE_INPUT_HEADER",
    "PRICE_SECTION_TITLE",
    "PriceImportError",
    "PriceImportReport",
    "PriceImportWrite",
    "QueueFormatError",
    "QueuePriceRow",
    "TASKS_SHEET",
    "apply_unit_price",
    "extract_mark_for_kind",
    "format_report",
    "import_queue_prices",
    "parse_queue_price_rows",
    "validate_unit_price",
]
