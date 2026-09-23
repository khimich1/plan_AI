#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Комплектация составных свай (серия 1.011.1-10 вып.8) в pb.db."""

from __future__ import annotations

import sqlite3
from dataclasses import dataclass
from pathlib import Path
from typing import Optional, Tuple

from openpyxl import load_workbook

from core.composite_pile_text_normalizer import canonicalize_composite_pile_mark
from core.price_db import DEFAULT_DB, _connect

SHEET_NAME = "Комплектация"
AssemblyPair = Tuple[str, str]  # upper_mark, lower_mark


@dataclass(frozen=True)
class CompositePileAssemblyRow:
    pile_mark: str
    joint_type: str
    upper_mark: str
    lower_mark: str


def init_composite_pile_assemblies_schema(db_path: str = DEFAULT_DB) -> None:
    conn = _connect(db_path)
    try:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS composite_pile_assemblies (
                pile_mark TEXT PRIMARY KEY,
                joint_type TEXT NOT NULL,
                upper_mark TEXT NOT NULL,
                lower_mark TEXT NOT NULL
            )
            """
        )
        conn.commit()
    finally:
        conn.close()


def _header_map(headers: tuple) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(headers):
        key = str(cell or "").strip().lower()
        if key:
            mapping[key] = idx
    return mapping


def _col(mapping: dict[str, int], *names: str) -> int:
    for name in names:
        if name in mapping:
            return mapping[name]
    raise ValueError(f"лист «{SHEET_NAME}»: нет колонки {names[0]!r}")


def parse_composite_pile_kit_rows_from_xlsx(xlsx_path: str) -> list[CompositePileAssemblyRow]:
    """Parse sheet «Комплектация» into assembly rows (canon marks)."""
    path = Path(xlsx_path)
    if not path.is_file():
        raise FileNotFoundError(xlsx_path)

    wb = load_workbook(path, read_only=True, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"нет листа «{SHEET_NAME}» в {path.name}")
    ws = wb[SHEET_NAME]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration as exc:
        raise ValueError(f"пустой лист «{SHEET_NAME}»") from exc

    mapping = _header_map(tuple(header))
    i_joint = _col(mapping, "тип стыка")
    i_pile = _col(mapping, "марка сваи")
    i_upper = _col(mapping, "марка верхней секции")
    i_lower = _col(mapping, "марка нижней секции")

    result: list[CompositePileAssemblyRow] = []
    for raw in rows_iter:
        if not raw:
            continue
        pile_raw = raw[i_pile] if i_pile < len(raw) else None
        if pile_raw is None or str(pile_raw).strip() == "":
            continue
        pile_mark = canonicalize_composite_pile_mark(str(pile_raw))
        upper_mark = canonicalize_composite_pile_mark(str(raw[i_upper] or ""))
        lower_mark = canonicalize_composite_pile_mark(str(raw[i_lower] or ""))
        if not pile_mark or not upper_mark or not lower_mark:
            raise ValueError(f"кривая строка комплектации: {raw!r}")
        joint = str(raw[i_joint] or "").strip()
        result.append(
            CompositePileAssemblyRow(
                pile_mark=pile_mark,
                joint_type=joint,
                upper_mark=upper_mark,
                lower_mark=lower_mark,
            )
        )
    wb.close()
    return result


def import_composite_pile_kit_from_xlsx(
    xlsx_path: str,
    db_path: str = DEFAULT_DB,
) -> int:
    """Replace composite_pile_assemblies with rows from xlsx. Returns row count."""
    rows = parse_composite_pile_kit_rows_from_xlsx(xlsx_path)
    init_composite_pile_assemblies_schema(db_path)
    conn = _connect(db_path)
    try:
        conn.execute("DELETE FROM composite_pile_assemblies")
        conn.executemany(
            """
            INSERT INTO composite_pile_assemblies
                (pile_mark, joint_type, upper_mark, lower_mark)
            VALUES (?, ?, ?, ?)
            """,
            [(r.pile_mark, r.joint_type, r.upper_mark, r.lower_mark) for r in rows],
        )
        conn.commit()
    finally:
        conn.close()
    return len(rows)


def get_composite_pile_assembly(
    pile_mark: str,
    db_path: str = DEFAULT_DB,
) -> Optional[AssemblyPair]:
    """Lookup whole-pile mark → (upper_mark, lower_mark). None if unknown."""
    canon = canonicalize_composite_pile_mark(pile_mark)
    if not canon:
        return None
    init_composite_pile_assemblies_schema(db_path)
    conn = _connect(db_path)
    try:
        cur = conn.execute(
            """
            SELECT upper_mark, lower_mark
            FROM composite_pile_assemblies
            WHERE pile_mark = ?
            """,
            (canon,),
        )
        row = cur.fetchone()
        if not row:
            return None
        return str(row[0]), str(row[1])
    finally:
        conn.close()


def list_section_marks(db_path: str = DEFAULT_DB) -> set[str]:
    """All upper/lower section marks present in the kit."""
    init_composite_pile_assemblies_schema(db_path)
    conn = _connect(db_path)
    try:
        cur = conn.execute(
            "SELECT upper_mark, lower_mark FROM composite_pile_assemblies"
        )
        marks: set[str] = set()
        for upper, lower in cur.fetchall():
            marks.add(str(upper))
            marks.add(str(lower))
        return marks
    finally:
        conn.close()
