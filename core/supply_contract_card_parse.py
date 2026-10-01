"""Разбор текстовой карточки покупателя в поля договора. Номер не выдаёт."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from io import BytesIO
from typing import Any
from zipfile import ZipFile

from docx import Document
from openpyxl import load_workbook
from pdfminer.high_level import extract_text

from core.supply_contract import normalize_digits, normalize_email

FIELD_NAMES = (
    "legal_form",
    "full_name",
    "short_name",
    "signatory_position",
    "signatory_name",
    "signatory_verb",
    "authority_basis",
    "poa_number",
    "poa_date",
    "inn",
    "kpp",
    "ogrn",
    "legal_address",
    "postal_address",
    "phone",
    "email",
    "bank_name",
    "account",
    "corr_account",
    "bik",
    "edo_operator",
    "edo_id",
    "okved",
)

_DIGIT_LENGTHS = {
    "inn": (10, 12),
    "kpp": (9,),
    "ogrn": (13, 15),
    "account": (20,),
    "corr_account": (20,),
    "bik": (9,),
}

_INN10_WEIGHTS = (2, 4, 10, 3, 5, 9, 4, 6, 8)
_INN12_WEIGHTS_11 = (7, 2, 4, 10, 3, 5, 9, 4, 6, 8)
_INN12_WEIGHTS_12 = (3, 7, 2, 4, 10, 3, 5, 9, 4, 6, 8)
_ACCOUNT_WEIGHTS = (7, 1, 3)
_SOURCE_TEXT_LIMIT = 20_000
_BANK_FIELD_KEYS = ("bank_name", "account", "corr_account", "bik")
_LABEL_TAILS = frozenset({
    "организации",
    "покупателя",
    "контрагента",
    "получателя",
    "банка",
    "в соответствии с учредительными документами",
})
_HOLE_FIELDS = (
    "legal_form",
    "signatory_position",
    "signatory_name",
    "inn",
    "kpp",
    "ogrn",
    "legal_address",
    "email",
    "bank_name",
    "account",
    "corr_account",
    "bik",
)

_REQUIRED = (
    "legal_form",
    "full_name",
    "short_name",
    "signatory_name",
    "inn",
    "ogrn",
    "legal_address",
    "email",
    "bank_name",
    "account",
    "corr_account",
    "bik",
)

_FORM_MARKERS = (
    ("ooo", ("общество с ограниченной ответственностью", "ооо")),
    ("ao", ("акционерное общество", "ао")),
    ("kfh", ("крестьянское (фермерское) хозяйство", "крестьянское фермерское хозяйство", "кфх")),
    ("ip", ("индивидуальный предприниматель", "ип")),
)


@dataclass
class CardParseResult:
    fields: dict[str, Any] = field(default_factory=dict)
    doubtful: list[str] = field(default_factory=list)
    accounts: list[str] = field(default_factory=list)
    banks: list[dict[str, Any]] = field(default_factory=list)
    source_text: str = ""
    verify_failed: bool = False
    inn_disputed: bool = False
    bank_unpaired: bool = False

    def as_dict(self) -> dict[str, Any]:
        return {
            "fields": {name: self.fields.get(name) for name in FIELD_NAMES},
            "doubtful": list(self.doubtful),
            "accounts": list(self.accounts),
            "banks": list(self.banks),
            "source_text": self.source_text[:_SOURCE_TEXT_LIMIT],
            "verify_failed": self.verify_failed,
        }


def empty_fields() -> dict[str, Any]:
    return {name: None for name in FIELD_NAMES}


def parse_card_text(text: str) -> CardParseResult:
    """Поля и сомнительные. Цифры и почта нормализуются как при подтверждении."""
    source = text or ""
    fields = empty_fields()
    doubtful: list[str] = []
    lowered = source.lower()

    fields["legal_form"] = _legal_form(lowered)
    short_name, full_name = _names(source, fields["legal_form"])
    fields["short_name"] = short_name
    fields["full_name"] = full_name
    position, signatory = _signatory(source)
    fields["signatory_position"] = position
    fields["signatory_name"] = signatory
    fields["signatory_verb"] = _verb(lowered)
    basis, poa_number, poa_date = _authority(source, signatory)
    fields["authority_basis"] = basis
    fields["poa_number"] = poa_number
    fields["poa_date"] = poa_date

    inn_values, kpp_values = _inn_kpp_values(source)
    inn_disputed = _store_digit(fields, doubtful, "inn", inn_values, accept_inn)
    _store_digit(fields, doubtful, "kpp", kpp_values, accept_kpp)
    _store_digit(fields, doubtful, "ogrn", _labeled_digits(source, r"огрн(?:ип)?"), accept_ogrn)
    _store_digit(fields, doubtful, "bik", _labeled_digits(source, r"бик"), accept_bik)
    accounts = _store_account(fields, doubtful, _account_numbers(source, corr=False))
    _store_digit(
        fields,
        doubtful,
        "corr_account",
        _account_numbers(source, corr=True),
        accept_corr_account,
    )

    fields["legal_address"] = _labeled_line(source, r"юридическ\w*\s+адрес", address=True)
    fields["postal_address"] = _labeled_line(source, r"почтов\w*\s+адрес", address=True)
    fields["phone"] = _phone(source)
    fields["email"] = _email(source)
    fields["bank_name"] = _bank_name(source)
    fields["edo_operator"] = _labeled_line(source, r"оператор\s+эдо")
    fields["edo_id"] = _edo_id(source)
    okved, okved_without_code = _primary_okved(source)
    fields["okved"] = okved
    if okved_without_code:
        _mark_doubtful(doubtful, "okved")
    banks, bank_unpaired = _apply_bank_outcome(fields, source)

    for key, allowed in _DIGIT_LENGTHS.items():
        value = fields.get(key)
        if value and len(str(value)) not in allowed and key not in doubtful:
            doubtful.append(key)
    for key in _REQUIRED:
        if not fields.get(key) and key not in doubtful:
            doubtful.append(key)
    if fields["legal_form"] in {"ooo", "ao"} and not fields.get("kpp"):
        if "kpp" not in doubtful:
            doubtful.append("kpp")
    if fields["legal_form"] in {"ooo", "ao"} and not fields.get("signatory_position"):
        if "signatory_position" not in doubtful:
            doubtful.append("signatory_position")
    return CardParseResult(
        fields=fields,
        doubtful=doubtful,
        accounts=accounts,
        banks=banks,
        source_text=(source or "")[:_SOURCE_TEXT_LIMIT],
        inn_disputed=inn_disputed,
        bank_unpaired=bank_unpaired,
    )


def extract_docx_text(data: bytes) -> str:
    document = Document(BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(part for part in parts if part and part.strip())


def extract_pdf_text(data: bytes) -> str:
    """Текст текстового слоя PDF. Пустая строка — слой не найден или файл не читается."""
    if not data:
        return ""
    try:
        return extract_text(BytesIO(data)) or ""
    except Exception:
        return ""


def has_text_layer(text: str) -> bool:
    return bool(re.search(r"[A-Za-zА-Яа-яЁё0-9]", text or ""))


def extract_xlsx_text(data: bytes) -> str:
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        parts: list[str] = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for cell in row:
                    if cell is None:
                        continue
                    text = str(cell).strip()
                    if text:
                        parts.append(text)
        return "\n".join(parts)
    finally:
        workbook.close()


def is_xlsx(data: bytes) -> bool:
    if not data.startswith(b"PK"):
        return False
    try:
        with ZipFile(BytesIO(data)) as archive:
            names = set(archive.namelist())
    except Exception:
        return False
    return "xl/workbook.xml" in names and "word/document.xml" not in names


def is_docx(data: bytes) -> bool:
    if not data.startswith(b"PK"):
        return False
    try:
        with ZipFile(BytesIO(data)) as archive:
            return "word/document.xml" in archive.namelist()
    except Exception:
        return False


def is_pdf(data: bytes) -> bool:
    return data.startswith(b"%PDF")


def _legal_form(lowered: str) -> str | None:
    for code, markers in _FORM_MARKERS:
        for marker in markers:
            if re.search(rf"(?<![а-яёa-z]){re.escape(marker)}(?![а-яёa-z])", lowered):
                return code
    return None


def _names(text: str, legal_form: str | None) -> tuple[str | None, str | None]:
    quoted = re.findall(r"«([^»]{2,120})»", text)
    quoted = [item.strip() for item in quoted if item.strip()]
    short = quoted[0] if quoted else None
    full = None
    long_form = re.search(
        r"(Общество с ограниченной ответственностью\s+«[^»]+»"
        r"|Акционерное общество\s+«[^»]+»"
        r"|Индивидуальный предприниматель\s+[А-ЯЁ][^,\n]{2,80}"
        r"|Крестьянское \(фермерское\) хозяйство\s+«[^»]+»)",
        text,
    )
    if long_form:
        full = re.sub(r"\s+", " ", long_form.group(1)).strip()
    if full is None and short and legal_form == "ooo":
        full = f"Общество с ограниченной ответственностью «{short}»"
    if full is None and short and legal_form == "ao":
        full = f"Акционерное общество «{short}»"
    if full is None and short and legal_form == "kfh":
        full = f"Крестьянское (фермерское) хозяйство «{short}»"
    if short is None and full:
        short = full
    if full is None:
        full = short
    return short, full


def _signatory(text: str) -> tuple[str | None, str | None]:
    match = re.search(
        r"(Генеральн\w+\s+директор\w*"
        r"|Индивидуальн\w+\s+предпринимател\w*"
        r"|Глава\s+КФХ"
        r"|Директор\w*"
        r"|Президент\w*)"
        r"\s+([А-ЯЁ][а-яё]+(?:\s+[А-ЯЁ][а-яё]+){1,2})",
        text,
    )
    if not match:
        return None, None
    return match.group(1).strip(), match.group(2).strip()


def _verb(lowered: str) -> str | None:
    if "действующей" in lowered:
        return "действующей"
    if "действующего" in lowered:
        return "действующего"
    return None


def _authority(text: str, signatory_name: str | None) -> tuple[str | None, str | None, str | None]:
    window = _signatory_window(text, signatory_name)
    source = window if window is not None else text
    lowered = source.lower()
    charter_at = lowered.find("устав")
    poa_at = lowered.find("доверен")
    if charter_at >= 0 and (poa_at < 0 or charter_at < poa_at):
        return "устав", None, None
    if poa_at >= 0:
        number = re.search(r"доверенност\w*[^0-9\n]{0,40}№\s*([0-9A-Za-zА-Яа-я/-]+)", source, re.I)
        issued = re.search(r"от\s+(\d{2}\.\d{2}\.\d{4})", source)
        return "доверенность", (number.group(1) if number else None), (issued.group(1) if issued else None)
    if window is not None and "устав" in text.lower():
        return "устав", None, None
    return None, None, None


def _signatory_window(text: str, signatory_name: str | None) -> str | None:
    if not signatory_name:
        return None
    lowered = text.lower()
    index = lowered.find(signatory_name.lower())
    if index < 0:
        return None
    return text[index : index + 180]


def _inn_kpp_values(text: str) -> tuple[list[str], list[str]]:
    inn: list[str] = []
    kpp: list[str] = []
    slash = re.compile(
        r"инн\s*/\s*кпп\s*[:№]?\s*([0-9][0-9 \u00a0]{4,20})\s*/\s*([0-9][0-9 \u00a0]{4,20})",
        re.I,
    )
    spans: list[tuple[int, int]] = []
    for match in slash.finditer(text):
        inn.append(normalize_digits(match.group(1)))
        kpp.append(normalize_digits(match.group(2)))
        spans.append(match.span())
    cleaned = _blank_spans(text, spans)
    inn.extend(_labeled_digits(cleaned, r"инн"))
    kpp.extend(_labeled_digits(cleaned, r"кпп"))
    return inn, kpp


def _blank_spans(text: str, spans: list[tuple[int, int]]) -> str:
    if not spans:
        return text
    chars = list(text)
    for start, end in spans:
        for index in range(start, end):
            chars[index] = " "
    return "".join(chars)


def _labeled_digits(text: str, label: str) -> list[str]:
    found: list[str] = []
    for raw in _labeled_value_strings(text, _pattern_for_label(label)):
        found.extend(_numbers_of_length(raw, _lengths_for_label(label)))
    return found


def _pattern_for_label(label: str) -> re.Pattern[str]:
    body = r"бик(?:\s+банка)?" if label == "бик" else label
    return re.compile(rf"(?<![а-яёa-z]){body}(?![а-яёa-z])\s*[:№]?", re.I)


def _lengths_for_label(label: str) -> tuple[int, ...]:
    if "инн" in label:
        return (10, 12)
    if "кпп" in label:
        return (9,)
    if "огрн" in label:
        return (13, 15)
    if label == "бик":
        return (9,)
    return (20,)


def _numbers_of_length(value: str, lengths: tuple[int, ...]) -> list[str]:
    cleaned = re.sub(r"\d{2}\.\d{2}\.\d{4}", " ", value or "")
    chunks = re.findall(r"\d+", cleaned)
    direct = [chunk for chunk in chunks if len(chunk) in lengths]
    if direct:
        return direct
    if not chunks:
        return []
    joined = "".join(chunks)
    if joined:
        return [joined]
    return []


def _nonempty_lines(text: str) -> list[str]:
    return [line.strip() for line in (text or "").splitlines() if line.strip()]


_LINE_IS_LABEL = re.compile(
    r"^(?:инн|кпп|огрн|бик|банк|юридическ\w*\s+адрес|почтов\w*\s+адрес|"
    r"р\s*/\s*с|расч|к\s*/\s*с|корр|e-?mail|электронн\w*\s+почт|тел)",
    re.I,
)


def _line_is_label(line: str) -> bool:
    return bool(_LINE_IS_LABEL.match(line.strip()))


_NEXT_LABEL = re.compile(
    r"(?<![а-яёa-z])(?:"
    r"инн|кпп|огрн(?:ип)?|бик(?:\s+банка)?|"
    r"р\s*/\s*с|расч[её]тн\w*\s+сч[её]т|"
    r"к\s*/\s*с|кор(?:респондентск\w*)?\s*сч[её]т"
    r")(?![а-яёa-z])",
    re.I,
)


def _rest_after_label(line: str, match: re.Match[str]) -> str:
    rest = re.sub(r"\([^)]*\)", " ", line[match.end() :])
    rest = rest.strip(" \t:.;,-–—")
    if rest.lower().replace("ё", "е") in _LABEL_TAILS:
        return ""
    return rest


def _cut_at_next_label(rest: str) -> str:
    """Хвост подписи кончается там, где на той же строке начинается следующая."""
    match = _NEXT_LABEL.search(rest)
    if match is None:
        return rest
    return rest[: match.start()].strip(" \t:.;,|-–—")


def _labeled_value_strings(text: str, pattern: re.Pattern[str]) -> list[str]:
    lines = _nonempty_lines(text)
    found: list[str] = []
    index = 0
    while index < len(lines):
        match = pattern.search(lines[index])
        if not match:
            index += 1
            continue
        rest = _rest_after_label(lines[index], match)
        step = 1
        if not rest and index + 1 < len(lines) and not _line_is_label(lines[index + 1]):
            rest = lines[index + 1].strip()
            step = 2
        if rest:
            rest = _cut_at_next_label(rest)
        if rest:
            found.append(rest)
        index += step
    return found


def _account_numbers(text: str, *, corr: bool) -> list[str]:
    if corr:
        label = r"(?:к\s*/\s*с|кор(?:респондентск\w*)?\s*сч[её]т)"
    else:
        label = r"(?:р\s*/\s*с|расч[её]тн\w*\s+сч[её]т)"
    return _labeled_digits(text, label)


def accept_inn(raw: str | None) -> str | None:
    digits = normalize_digits(raw)
    if len(digits) == 10 and _inn_control(digits[:9], _INN10_WEIGHTS) == int(digits[9]):
        return digits
    if len(digits) == 12:
        eleventh = _inn_control(digits[:10], _INN12_WEIGHTS_11)
        twelfth = _inn_control(digits[:11], _INN12_WEIGHTS_12)
        if eleventh == int(digits[10]) and twelfth == int(digits[11]):
            return digits
    return None


def accept_kpp(raw: str | None) -> str | None:
    digits = normalize_digits(raw)
    if len(digits) != 9:
        return None
    return digits


def accept_ogrn(raw: str | None) -> str | None:
    digits = normalize_digits(raw)
    if len(digits) == 13 and int(digits[:12]) % 11 % 10 == int(digits[12]):
        return digits
    if len(digits) == 15 and int(digits[:14]) % 13 % 10 == int(digits[14]):
        return digits
    return None


def accept_bik(raw: str | None) -> str | None:
    digits = normalize_digits(raw)
    if len(digits) != 9:
        return None
    return digits


def accept_corr_account(raw: str | None) -> str | None:
    digits = normalize_digits(raw)
    if len(digits) != 20:
        return None
    return digits


def accept_account(raw: str | None, bik: str | None) -> str | None:
    """20 цифр. С принятым БИК — ключ 579-П, без БИК длина достаточна."""
    digits = normalize_digits(raw)
    if len(digits) != 20:
        return None
    accepted_bik = accept_bik(bik)
    if accepted_bik is None:
        return digits
    if _account_key_ok(digits, accepted_bik):
        return digits
    return None


def _inn_control(body: str, weights: tuple[int, ...]) -> int:
    total = sum(int(digit) * weight for digit, weight in zip(body, weights, strict=True))
    return (total % 11) % 10


def _account_key_ok(account: str, bik: str) -> bool:
    composite = bik[-3:] + account
    total = sum(int(digit) * _ACCOUNT_WEIGHTS[index % 3] for index, digit in enumerate(composite))
    return total % 10 == 0


def accept_bundle(
    bank_name: str | None,
    account: str,
    corr: str | None,
    bik: str,
) -> dict[str, Any] | None:
    accepted_bik = accept_bik(bik)
    accepted_account = accept_account(account, accepted_bik)
    if accepted_account is None or accepted_bik is None:
        return None
    accepted_corr = accept_corr_account(corr) if corr else None
    if corr and accepted_corr is None:
        return None
    name = (bank_name or "").strip() or None
    if name and _suspicious_bank(name):
        name = None
    return {
        "bank_name": name,
        "account": accepted_account,
        "corr_account": accepted_corr,
        "bik": accepted_bik,
    }


def _empty_block() -> dict[str, str | None]:
    return {"bank_name": None, "account": None, "corr_account": None, "bik": None}


def _bank_blocks(text: str) -> list[dict[str, str | None]]:
    lines = _nonempty_lines(text)
    blocks: list[dict[str, str | None]] = []
    current = _empty_block()
    kinds = (
        ("account", _RE_ACCOUNT, (20,)),
        ("corr_account", _RE_CORR, (20,)),
        ("bik", _RE_BIK, (9,)),
        ("bank_name", _RE_BANK, ()),
    )
    index = 0
    while index < len(lines):
        line = lines[index]
        matched: tuple[str, re.Match[str], tuple[int, ...]] | None = None
        for kind, pattern, lengths in kinds:
            match = pattern.search(line)
            if match:
                matched = (kind, match, lengths)
                break
        if matched is None:
            index += 1
            continue
        kind, match, lengths = matched
        rest = _rest_after_label(line, match)
        step = 1
        if not rest and index + 1 < len(lines) and not _line_is_label(lines[index + 1]):
            rest = lines[index + 1].strip()
            step = 2
        if current.get(kind):
            blocks.append(current)
            current = _empty_block()
        if kind == "bank_name":
            current[kind] = rest or None
        else:
            numbers = _numbers_of_length(rest, lengths)
            current[kind] = numbers[0] if numbers else None
        index += step
    if any(current.values()):
        blocks.append(current)
    return blocks


def _bundle_from_block(block: dict[str, str | None]) -> dict[str, Any] | None:
    account = block.get("account")
    bik = block.get("bik")
    if not account or not bik:
        return None
    return accept_bundle(block.get("bank_name"), account, block.get("corr_account"), bik)


def _apply_bank_outcome(fields: dict[str, Any], text: str) -> tuple[list[dict[str, Any]], bool]:
    bundles = [item for item in (_bundle_from_block(block) for block in _bank_blocks(text)) if item]
    if len(bundles) >= 2:
        for key in _BANK_FIELD_KEYS:
            fields[key] = None
        return bundles, False
    raw_accounts = _account_numbers(text, corr=False)
    raw_biks = _labeled_digits(text, "бик")
    raw_corrs = _account_numbers(text, corr=True)
    raw_banks = _labeled_value_strings(text, _RE_BANK)
    unpaired = False
    if len(raw_corrs) >= 2:
        fields["corr_account"] = None
        unpaired = True
    if len(raw_accounts) >= 2 and len(raw_biks) == 1:
        fields["account"] = None
        unpaired = True
    if len(raw_biks) >= 2:
        fields["bik"] = None
        unpaired = True
    if len(raw_banks) >= 2:
        fields["bank_name"] = None
        unpaired = True
    return [], unpaired


def card_holes(result: CardParseResult) -> list[str]:
    """Пустые обязательные поля, спор ИНН и несобранные банковские числа."""
    frozen = set(frozen_fields(result))
    skip_bank = len(result.banks) >= 2 or result.bank_unpaired
    holes: list[str] = []
    if result.bank_unpaired:
        holes.append("banks")
    for key in _HOLE_FIELDS:
        if key in _BANK_FIELD_KEYS and skip_bank:
            continue
        if key == "account" and len(result.accounts) > 1:
            continue
        if key in frozen:
            continue
        if key in {"signatory_position", "kpp"}:
            if result.fields.get(key) or key not in result.doubtful:
                continue
        elif result.fields.get(key):
            continue
        holes.append(key)
    if result.inn_disputed and "inn" not in holes and "inn" not in frozen:
        holes.append("inn")
    return holes


def frozen_fields(result: CardParseResult) -> dict[str, str]:
    """Цифры, которые прошли контроль и не спорят со второй такой же."""
    frozen: dict[str, str] = {}
    if result.fields.get("inn") and not result.inn_disputed:
        frozen["inn"] = str(result.fields["inn"])
    for key in ("kpp", "ogrn", "bik", "corr_account", "account"):
        value = result.fields.get(key)
        if value:
            frozen[key] = str(value)
    return frozen


def _store_digit(
    fields: dict[str, Any],
    doubtful: list[str],
    key: str,
    values: list[str],
    accept,
) -> bool:
    unique = list(dict.fromkeys(values))
    accepted = [item for item in (accept(value) for value in unique) if item]
    accepted = list(dict.fromkeys(accepted))
    if not accepted:
        fields[key] = None
        if unique:
            _mark_doubtful(doubtful, key)
        return False
    fields[key] = accepted[0]
    if len(accepted) > 1:
        _mark_doubtful(doubtful, key)
        return True
    return False


def _store_account(fields: dict[str, Any], doubtful: list[str], values: list[str]) -> list[str]:
    unique = list(dict.fromkeys(values))
    bik = fields.get("bik")
    accepted = [item for item in (accept_account(value, bik) for value in unique) if item]
    accepted = list(dict.fromkeys(accepted))
    if len(accepted) == 1:
        fields["account"] = accepted[0]
    else:
        fields["account"] = None
    if not accepted and unique:
        _mark_doubtful(doubtful, "account")
    if len(accepted) != 1 or accept_bik(bik) is None:
        if accepted or unique:
            _mark_doubtful(doubtful, "account")
    return accepted


def _mark_doubtful(doubtful: list[str], key: str) -> None:
    if key not in doubtful:
        doubtful.append(key)


def _labeled_line(text: str, label: str, *, address: bool = False) -> str | None:
    pattern = re.compile(
        rf"(?<![а-яёa-z]){label}(?:\s+(?:организации|покупателя|контрагента|получателя))?\s*[:\-–]?",
        re.I,
    )
    values = _labeled_value_strings(text, pattern)
    if not values:
        return None
    value = values[0].strip(" \t,;")
    if not value:
        return None
    if address and _suspicious_address(value):
        return None
    return value


def _suspicious_address(value: str) -> bool:
    text = value.strip()
    if re.fullmatch(r"\([^)]*\)", text):
        return True
    plain = re.sub(r"\([^)]*\)", " ", text).strip(" \t:.;,-–—")
    if not plain:
        return True
    return plain.lower().replace("ё", "е") in _LABEL_TAILS


def _suspicious_bank(value: str) -> bool:
    lowered = value.lower().replace("ё", "е")
    if re.search(r"р\s*/\s*с", lowered) or re.search(r"к\s*/\s*с", lowered):
        return True
    if re.search(r"(?<![а-яёa-z])бик(?![а-яёa-z])", lowered):
        return True
    return bool(re.search(r"\d{20}", re.sub(r"\D", "", value)))


def _phone(text: str) -> str | None:
    match = re.search(r"(?:тел(?:ефон)?\.?)\s*[:\-]?\s*(\+?[0-9][0-9()\-\s]{6,20})", text, re.I)
    if not match:
        return None
    return re.sub(r"\s+", " ", match.group(1)).strip()


_EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")


def _email(text: str) -> str | None:
    found: list[str] = []
    suspicious = False
    for line in _nonempty_lines(text):
        matches = _EMAIL_RE.findall(normalize_email(line))
        if len(matches) > 1:
            suspicious = True
        for email in matches:
            local = email.split("@", 1)[0]
            if re.search(r"\d{7,}", local):
                suspicious = True
            else:
                found.append(email)
    if suspicious or len(found) != 1:
        return None
    return found[0]


_RE_ACCOUNT = re.compile(
    r"(?<![а-яёa-z])(?:р\s*/\s*с|расч[её]тн\w*\s+сч[её]т)(?![а-яёa-z])\s*[:№]?",
    re.I,
)
_RE_CORR = re.compile(
    r"(?<![а-яёa-z])(?:к\s*/\s*с|кор(?:респондентск\w*)?\s*сч[её]т)(?![а-яёa-z])\s*[:№]?",
    re.I,
)
_RE_BIK = re.compile(r"(?<![а-яёa-z])бик(?:\s+банка)?(?![а-яёa-z])\s*[:№]?", re.I)
_RE_BANK = re.compile(
    r"(?<![а-яёa-z])банк(?![а-яёa-z])(?:\s+(?:получателя|организации))?\s*(?::|\s*$)",
    re.I,
)


def _bank_name(text: str) -> str | None:
    for value in _labeled_value_strings(text, _RE_BANK):
        cleaned = value.strip(" \t,;")
        if not cleaned:
            continue
        if _suspicious_bank(cleaned):
            return None
        return cleaned
    quoted_bank = re.search(r"(ПАО\s+«[^»]+»|АО\s+«[^»]+»)", text)
    if quoted_bank and "банк" in text.lower():
        name = quoted_bank.group(1).strip()
        if not _suspicious_bank(name):
            return name
    branch = re.search(r"([^\n]*отделени[ея][^\n]*)", text, re.I)
    if branch:
        name = branch.group(1).strip(" \t,;")
        if name and not _suspicious_bank(name):
            return name
    return None


def _edo_id(text: str) -> str | None:
    match = re.search(r"(?:идентификатор\s+(?:участника\s+)?эдо)\s*[:\-–]?\s*([A-Za-z0-9]+)", text, re.I)
    if not match:
        return None
    return match.group(1)


_OKVED_CODE_RE = re.compile(r"\d{2}(?:\.\d{1,2}){1,3}")
_OKVED_LABEL_RE = re.compile(r"(?<![а-яёa-z])оквэд(?![а-яёa-z])", re.I)
_OKVED_EXTRA_RE = re.compile(r"(?<![а-яёa-z])доп\s*\.", re.I)


def okved_code_from_value(value: str | None) -> str | None:
    """Первый код в уже выделенном значении. Хвост «Доп.» не читается."""
    text = str(value or "").strip()
    if not text:
        return None
    labeled = re.match(r"оквэд(?![а-яёa-z])\s*", text, re.I)
    if labeled:
        text = text[labeled.end() :]
    extra = _OKVED_EXTRA_RE.search(text)
    window = text[: extra.start()] if extra else text
    found = _OKVED_CODE_RE.search(window)
    if not found:
        return None
    return found.group(0)


def _primary_okved(text: str) -> tuple[str | None, bool]:
    """Код и флаг «подпись есть, кода нет». Нет подписи — пусто и не сомнительно."""
    label = _OKVED_LABEL_RE.search(text or "")
    if not label:
        return None, False
    code = okved_code_from_value(text[label.end() :])
    if not code:
        return None, True
    return code, False
