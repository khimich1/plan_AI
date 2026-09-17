#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Парсер выгрузки 1С «Универсальный отчёт» (.xlsx) с УИД номенклатуры.

Шапка плавающая: строка с «Номенклатура» + «УИД». Колонки по именам
(якоря объединённых ячеек). Строка «Отбор:» над шапкой → partial=True.
«Итого» пропускается. Количества и цены в компоновке нет.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any, Optional

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

_FORMAT_MSG = "Похоже, это не универсальный отчёт 1С"


class UniversalReport1CFormatError(ValueError):
    """Файл не является универсальным отчётом 1С с УИД."""


@dataclass(frozen=True, slots=True)
class UniversalReportRow:
    guid: str
    code: str
    name: str
    qty: float | None
    weight: float | None
    volume: float | None
    price: float | None
    source_file: str
    row_index: int


@dataclass(frozen=True, slots=True)
class UniversalReport:
    rows: tuple[UniversalReportRow, ...]
    partial: bool
    source_file: str


def parse_universal_report_xlsx(path: Path | str) -> UniversalReport:
    """Прочитать .xlsx «Универсального отчёта» 1С.

    Raises:
        FileNotFoundError: путь не существует.
        UniversalReport1CFormatError: нет шапки «Номенклатура»+«УИД».
    """
    file_path = Path(path)
    if not file_path.is_file():
        raise FileNotFoundError(str(file_path))
    try:
        workbook = load_workbook(filename=str(file_path), data_only=True, read_only=False)
    except Exception as exc:
        raise UniversalReport1CFormatError(f"{_FORMAT_MSG}: не удалось прочитать файл") from exc
    try:
        sheet = workbook[workbook.sheetnames[0]]
        return _parse_sheet(sheet, file_path.name)
    finally:
        workbook.close()


def _parse_sheet(sheet: Worksheet, source_file: str) -> UniversalReport:
    header_row = _find_header_row(sheet)
    if header_row is None:
        raise UniversalReport1CFormatError(_FORMAT_MSG)
    columns = _map_columns(sheet, header_row)
    if "name" not in columns or "guid" not in columns:
        raise UniversalReport1CFormatError(_FORMAT_MSG)

    partial = _has_filter_above(sheet, header_row)
    rows: list[UniversalReportRow] = []
    for excel_row in range(header_row + 1, sheet.max_row + 1):
        name = _cell_str(sheet, excel_row, columns["name"])
        if not name:
            continue
        if _is_total_row(name):
            continue
        guid = _normalize_guid(_cell_str(sheet, excel_row, columns["guid"]))
        if not guid:
            continue
        code_col = columns.get("code")
        weight_col = columns.get("weight")
        volume_col = columns.get("volume")
        rows.append(
            UniversalReportRow(
                guid=guid,
                code=_cell_code(sheet, excel_row, code_col) if code_col else "",
                name=name,
                qty=None,
                weight=_cell_float(sheet, excel_row, weight_col) if weight_col else None,
                volume=_cell_float(sheet, excel_row, volume_col) if volume_col else None,
                price=None,
                source_file=source_file,
                row_index=excel_row,
            )
        )
    return UniversalReport(rows=tuple(rows), partial=partial, source_file=source_file)


def _find_header_row(sheet: Worksheet) -> Optional[int]:
    for row_idx in range(1, sheet.max_row + 1):
        texts = [_norm_header(_cell_value(sheet, row_idx, col)) for col in range(1, sheet.max_column + 1)]
        has_name = any(text == "номенклатура" or text.startswith("номенклатура.") for text in texts)
        has_uid = any(text == "уид" or text.startswith("уид ") for text in texts)
        if has_name and has_uid:
            return row_idx
    return None


def _map_columns(sheet: Worksheet, header_row: int) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for col in range(1, sheet.max_column + 1):
        text = _norm_header(_cell_value(sheet, header_row, col))
        if not text:
            continue
        if text == "номенклатура" or text.startswith("номенклатура."):
            mapping.setdefault("name", col)
        elif text == "уид" or text.startswith("уид "):
            mapping.setdefault("guid", col)
        elif text == "код" or text.startswith("код "):
            mapping.setdefault("code", col)
        elif text.startswith("вес") and "знаменател" not in text:
            mapping.setdefault("weight", col)
        elif text.startswith("объем") and "знаменател" not in text:
            mapping.setdefault("volume", col)
    return mapping


def _has_filter_above(sheet: Worksheet, header_row: int) -> bool:
    for row_idx in range(1, header_row):
        for col in range(1, sheet.max_column + 1):
            text = _cell_str(sheet, row_idx, col)
            if text.lower().replace("ё", "е").startswith("отбор"):
                return True
    return False


def _is_total_row(name: str) -> bool:
    return name.strip().lower().replace("ё", "е").startswith("итого")


def _normalize_guid(value: str) -> str:
    text = (value or "").strip()
    if text.startswith("{") and text.endswith("}") and len(text) >= 2:
        text = text[1:-1].strip()
    # Как pricelist_1c_parser: без скобок, lowercase.
    return text.lower()


def _cell_value(sheet: Worksheet, row: int, col: int) -> Any:
    return sheet.cell(row, col).value


def _cell_str(sheet: Worksheet, row: int, col: int) -> str:
    return _as_str(_cell_value(sheet, row, col))


def _cell_code(sheet: Worksheet, row: int, col: int) -> str:
    value = _cell_value(sheet, row, col)
    if value is None:
        return ""
    if isinstance(value, bool):
        return ""
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def _cell_float(sheet: Worksheet, row: int, col: int) -> Optional[float]:
    value = _cell_value(sheet, row, col)
    if value in ("", None):
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = _as_str(value).replace(" ", "").replace(",", ".")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _as_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return str(value).strip()


def _norm_header(value: Any) -> str:
    text = _as_str(value).replace("ё", "е").replace("Ё", "Е").lower()
    return " ".join(text.split())


__all__ = [
    "UniversalReport",
    "UniversalReport1CFormatError",
    "UniversalReportRow",
    "parse_universal_report_xlsx",
]
