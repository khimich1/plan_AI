#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Импорт GUID составных свай из листа «Сваи составные — только в 1С»."""

from __future__ import annotations

import sqlite3
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook

from core.composite_pile_kit import list_section_marks
from core.composite_pile_line_parser import is_composite_section_mark
from core.composite_pile_price_db import list_composite_pile_price_marks
from core.composite_pile_text_normalizer import canonicalize_composite_pile_mark
from core.nomenclature_guid import ensure_schema, upsert
from core.price_db import DEFAULT_DB, _connect

SHEET_NAME = "Сваи составные — только в 1С"
FORBIDDEN_SOURCE = "прайс сваи составные.xls"


@dataclass(frozen=True)
class CompositePileGuidImportReport:
    written_auto: int
    written_ambiguous: int
    skipped: int


def _known_section_marks(db_path: str) -> set[str]:
    return list_section_marks(db_path) | list_composite_pile_price_marks(db_path)


def _refuse_forbidden_source(xlsx_path: str) -> None:
    name = Path(xlsx_path).name.strip().lower().replace("ё", "е")
    if name == FORBIDDEN_SOURCE or "прайс сваи составные" in name:
        raise ValueError(
            "файл «Прайс сваи составные.xls» не используется для GUID составных"
        )


def parse_composite_pile_guid_rows_from_xlsx(
    xlsx_path: str,
    *,
    sheet_name: str = SHEET_NAME,
) -> list[tuple[str, str]]:
    """Return list of (guid, raw_name) from the 1C sheet."""
    _refuse_forbidden_source(xlsx_path)
    path = Path(xlsx_path)
    if not path.is_file():
        raise FileNotFoundError(xlsx_path)
    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"нет листа «{sheet_name}»")
    ws = wb[sheet_name]
    rows_iter = ws.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if header is None:
        raise ValueError("пустой лист GUID")
    header_l = [str(c or "").strip().lower() for c in header]
    try:
        i_guid = next(i for i, h in enumerate(header_l) if "guid" in h)
        i_name = next(i for i, h in enumerate(header_l) if "наименован" in h)
    except StopIteration as exc:
        raise ValueError("ожидались колонки GUID и Наименование") from exc

    out: list[tuple[str, str]] = []
    for raw in rows_iter:
        if not raw:
            continue
        guid = str(raw[i_guid] or "").strip() if i_guid < len(raw) else ""
        name = str(raw[i_name] or "").strip() if i_name < len(raw) else ""
        if not guid or not name:
            continue
        out.append((guid, name))
    wb.close()
    return out


def import_composite_pile_guids_from_xlsx(
    xlsx_path: str,
    db_path: str = DEFAULT_DB,
    *,
    sheet_name: str = SHEET_NAME,
) -> CompositePileGuidImportReport:
    """Import exact section canon matches as auto; duplicates → ambiguous.

    Names that do not canonicalize to a known series section are skipped.
    Never writes guid_1c_u.
    """
    raw_rows = parse_composite_pile_guid_rows_from_xlsx(
        xlsx_path, sheet_name=sheet_name
    )
    known = _known_section_marks(db_path)
    if not known:
        raise ValueError(
            "сначала импортируйте комплектацию или прайс составных "
            "(нужен whitelist секций серии)"
        )

    by_canon: dict[str, set[str]] = defaultdict(set)
    skipped = 0
    for guid, name in raw_rows:
        canon = canonicalize_composite_pile_mark(name)
        if not is_composite_section_mark(canon) or canon not in known:
            skipped += 1
            continue
        by_canon[canon].add(guid)

    conn = _connect(db_path)
    try:
        ensure_schema(conn)
        written_auto = 0
        written_ambiguous = 0
        for canon, guids in by_canon.items():
            if len(guids) == 1:
                guid = next(iter(guids))
                upsert(
                    conn,
                    "composite_pile",
                    canon,
                    guid_1c=guid,
                    guid_1c_u=None,
                    match_status="auto",
                    match_note=None,
                )
                written_auto += 1
            else:
                upsert(
                    conn,
                    "composite_pile",
                    canon,
                    guid_1c=None,
                    guid_1c_u=None,
                    match_status="ambiguous",
                    match_note=f"несколько GUID: {', '.join(sorted(guids))}",
                )
                written_ambiguous += 1
        conn.commit()
    finally:
        conn.close()

    return CompositePileGuidImportReport(
        written_auto=written_auto,
        written_ambiguous=written_ambiguous,
        skipped=skipped,
    )
