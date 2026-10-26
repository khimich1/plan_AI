# -*- coding: utf-8 -*-
"""Пул дорожек КЗ 2026: пары продольных резов и дельты перерасхода.

Лист «по 1» — каждая дорожка отдельно через оптимизатор («сильные первыми»).
Дальше те же дорожки пачками 5, 10, 21, 30, 60 без перекрытия.

Дельты листа «по 1» считаются к фактической карте. Дельты пула — к листу «по 1».
Отрицательная дельта дешевле базы. Цена плит из остатка — масштаб, не экономия.
"""
from __future__ import annotations

import argparse
import contextlib
import io
import json
import logging
import math
import sqlite3
import sys
import time
import traceback
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from core.price_db import length_m_to_price_length_dm  # noqa: E402
from core.project_paths import PRICE_DB_PATH  # noqa: E402
from scripts.build_kz_overrun_xlsx_2026 import CACHE  # noqa: E402
from scripts.build_kz_overrun_xlsx_2026 import grade_of  # noqa: E402
from scripts.pool_overrun_experiment_2026 import (  # noqa: E402
    MIN_USEFUL_MM,
    money_and_waste,
    pack_slabs,
    per_m,
    pieces_of,
    series_load,
    signature,
)
from scripts.build_kz_overrun_xlsx_2026 import load_wires, make_lookup  # noqa: E402

OUT = Path("/home/roman/Рабочий стол/с производства/Пул_пары_перерасход_КЗ_2026.xlsx")
JSON_OUT = ROOT / "data" / "kz_pool_pair_2026.json"
POOL_SIZES = (5, 10, 21, 30, 60)
PLATE_WIDTH_MM = 1200

DETAIL_HEADERS = [
    "Уровень", "Пачка", "Вход", "№", "Дата", "Дорожка",
    "Метраж, м", "Плит, шт", "Состав дорожки",
    "Перебетонир.: метраж, м", "% метража перебетонир.", "Цена перебетонир., руб",
    "Среднее переармир., проволок", "Макс переармир., проволок",
    "Лишняя проволока, м", "Цена переармир., руб", "Итого перерасход, руб",
    "Без пары, шт", "Остаток, м²", "Обрезок, м²",
    "База дельты", "Δ метраж, м", "Δ переарм., руб", "Δ перебетон., руб",
    "Δ переарм., руб/м", "Δ перебетон., руб/м",
    "Цена плит из остатка, руб", "Себестоимость плит из остатка, руб", "Набор совпал", "Примечание",
]

SUMMARY_HEADERS = [
    "Пул", "Дорожек на входе", "Дорожек на выходе", "Дорожек в хвосте", "Пачек",
    "Пачек с тем же набором", "Метраж, м",
    "Δ метраж к карте, м", "Δ метраж к «по 1», м",
    "Переармирование, руб", "Переармирование, руб/м",
    "Δ переарм. к карте, руб", "Δ переарм. к «по 1», руб", "Δ переарм. к «по 1», руб/м",
    "Перебетонирование, руб", "Перебетонирование, руб/м",
    "Δ перебетон к карте, руб", "Δ перебетон к «по 1», руб", "Δ перебетон к «по 1», руб/м",
    "Резов без пары, шт", "Остаток резов, м²", "Обрезок вторичный, м²",
    "Цена плит из остатка, руб", "Себестоимость плит из остатка, руб", "Секунды", "Примечание",
]


def rest_key(length: float, rest_mm: int) -> tuple[float, int]:
    return (round(float(length), 2), int(rest_mm))


def split_batches(n: int, size: int) -> list[tuple[int, int]]:
    if size <= 0:
        raise ValueError("размер пачки должен быть больше нуля")
    out: list[tuple[int, int]] = []
    start = 0
    while start < n:
        end = min(n, start + size)
        out.append((start, end))
        start = end
    return out


def tail_tracks(n: int, size: int) -> int:
    if n <= 0 or size <= 0:
        return 0
    return n % size


def price_load(load: float | int) -> int:
    code = int(math.floor(float(load)))
    return 12 if code >= 13 else code


def cost_load(load: float | int) -> float:
    """В справочнике себестоимости только 8п и 12,5п."""
    return 8.0 if price_load(load) <= 8 else 12.5


def plate_cost(
    costs: dict[tuple[int, float], float],
    length_m: float,
    width_mm: int,
    load: float | int,
) -> float | None:
    """Себестоимость заготовки 12 дм той же длины и нагрузки.

    Узкая плита отдельно всё равно льётся на 1200 мм, остаток ширины — отход.
    """
    if width_mm <= 0 or length_m <= 0:
        return None
    length_dm = int(round(float(length_m) * 10))
    load_code = cost_load(load)

    def raw_at(dm: int) -> float | None:
        value = costs.get((dm, load_code))
        if value is not None and value > 0:
            return float(value)
        return None

    return raw_at(length_dm) or raw_at(length_dm - 1) or raw_at(length_dm + 1)


def catalog_sale_price(
    prices: dict[tuple[int, int], float],
    length_m: float,
    width_mm: int,
    load: float | int,
) -> float | None:
    """Прайс по длине и нагрузке, умноженный на долю ширины от 1200 мм."""
    if width_mm <= 0 or length_m <= 0:
        return None
    length_dm = length_m_to_price_length_dm(length_m)
    load_code = price_load(load)

    def raw_at(dm: int) -> float | None:
        value = prices.get((dm, load_code))
        if value is not None and value > 0:
            return float(value)
        return None

    raw = raw_at(length_dm) or raw_at(length_dm - 1) or raw_at(length_dm + 1)
    if raw is None:
        return None
    share = min(1.0, width_mm / float(PLATE_WIDTH_MM))
    return raw * share


def cut_accounts(primary_cuts: list[dict] | None, secondary_cuts: list[dict] | None) -> dict:
    """Несъеденные первичные остатки и обрезок вторичных резов."""
    parents: set[str] = set()
    for sec in secondary_cuts or []:
        parent = sec.get("parent_instance_id")
        if parent:
            parents.add(parent)
    keys: Counter = Counter()
    rest_m2 = 0.0
    unpaired = 0
    for cut in primary_cuts or []:
        rest = int(cut.get("rest") or 0)
        if rest <= 0:
            continue
        if cut.get("primary_instance_id") in parents:
            continue
        length = float((cut.get("lengths") or [0])[0] or 0)
        unpaired += 1
        keys[rest_key(length, rest)] += 1
        rest_m2 += (rest / 1000.0) * length

    scrap_m2 = 0.0
    seen: set[str] = set()
    for index, sec in enumerate(secondary_cuts or []):
        parent = sec.get("parent_instance_id") or f"orphan-{index}"
        if parent in seen:
            continue
        seen.add(parent)
        waste = float(sec.get("waste") or 0)
        source_length = float((sec.get("source_lengths") or [0])[0] or 0)
        scrap_m2 += (waste / 1000.0) * source_length
        if sec.get("type") in ("transverse", "multiple_transverse"):
            out_len = float((sec.get("lengths") or [0])[0] or 0)
            source_width_m = float(sec.get("source") or 0) / 1000.0
            scrap_m2 += max(0.0, source_length - out_len) * source_width_m
    return {
        "unpaired": unpaired,
        "keys": keys,
        "rest_m2": rest_m2,
        "scrap_m2": scrap_m2,
    }


def rescued_secondaries(available: Counter, secondary_cuts: list[dict] | None) -> list[dict]:
    """Вторичные плиты с остатка, который на листе «по 1» был несъеден.

    Один остаток списывается один раз. В список входят все плиты с этого остатка.
    """
    left = Counter(available)
    grouped: dict[str, list[dict]] = {}
    order: list[str] = []
    for index, sec in enumerate(secondary_cuts or []):
        parent = sec.get("parent_instance_id") or f"orphan-{index}"
        if parent not in grouped:
            grouped[parent] = []
            order.append(parent)
        grouped[parent].append(sec)
    rescued: list[dict] = []
    for parent in order:
        group = grouped[parent]
        head = group[0]
        length = float((head.get("source_lengths") or [0])[0] or 0)
        rest = int(head.get("source") or 0)
        key = rest_key(length, rest)
        if left[key] <= 0:
            continue
        left[key] -= 1
        for sec in group:
            width = int((sec.get("cuts") or [0])[0] or 0)
            rescued.append({
                "length": float((sec.get("lengths") or [0])[0] or 0),
                "width_mm": width,
                "load": sec.get("load_code") or 8,
            })
    return rescued


def rub_delta(current: float, base: float) -> int:
    return int(round(current - base))


def per_m_delta(current_rub: float, current_m: float, base_rub: float, base_m: float) -> float:
    return round(per_m(current_rub, current_m) - per_m(base_rub, base_m), 2)


def _order_piece(length: float, width_m: float, load, wires, lookup, name: str | None) -> dict | None:
    if length <= 0 or width_m <= 0:
        return None
    width_mm = int(round(float(width_m) * 1000))
    lc = series_load(load if load is not None else 8)
    ld = int(round(length * 10))
    if wires is None:
        wires = lookup(ld, lc)
    if wires is None:
        return None
    width_label = int(width_mm / 100) if abs(width_mm / 100 - round(width_mm / 100)) < 0.05 else round(width_mm / 100, 1)
    return {
        "length": length,
        "width_mm": width_mm,
        "load": lc,
        "wires": float(wires),
        "grade": grade_of(ld, lc),
        "ld": ld,
        "name": name or f"ПБ {ld}-{width_label}-{lc}п",
    }


def order_pieces_from_layout(track: dict, lookup) -> tuple[list[dict], float]:
    """Плиты дорожки в заказанной ширине. У продольного реза ширина — main_w, не 1200 мм."""
    out: list[dict] = []
    extra_scrap = 0.0
    for item in track.get("items") or []:
        mode = item.get("mode")
        length = float(item.get("length") or 0)
        load = item.get("load_code") or 8
        wires = item.get("reinforcement")
        if mode == "transverse":
            width_m = float(item.get("width") or 1.2)
            remainder = float(item.get("remainder") or 0)
            if remainder > 0.05:
                extra_scrap += remainder * width_m
            length = float(item.get("target_length") or length)
            piece = _order_piece(length, width_m, load, wires, lookup, item.get("plate_name"))
            if piece:
                out.append(piece)
            continue
        if mode == "split" or item.get("main_w"):
            piece = _order_piece(
                length, float(item.get("main_w") or 0), load, wires, lookup,
                item.get("plate_name") or item.get("label_main"),
            )
            if piece:
                out.append(piece)
            for sec in item.get("secondary_cuts") or []:
                sec_w = float(sec.get("width") or 0)
                sec_len = float(sec.get("target_length") or length)
                remainder = float(sec.get("remainder") or 0)
                if remainder > 0.05:
                    extra_scrap += remainder * sec_w
                sec_piece = _order_piece(
                    sec_len, sec_w, sec.get("load_code") or load, sec.get("reinforcement"),
                    lookup, sec.get("label"),
                )
                if sec_piece:
                    out.append(sec_piece)
            continue
        width_m = float(item.get("width") or 1.2)
        piece = _order_piece(length, width_m, load, wires, lookup, item.get("plate_name"))
        if piece:
            out.append(piece)
    return out, extra_scrap


def layout_unpaired(pieces: list[dict]) -> int:
    count = 0
    for slab in pack_slabs(pieces):
        gap = max(0, PLATE_WIDTH_MM - slab["used_mm"])
        if gap >= MIN_USEFUL_MM:
            count += 1
    return count


def load_price_table() -> dict[tuple[int, int], float]:
    con = sqlite3.connect(f"file:{PRICE_DB_PATH}?mode=ro", uri=True)
    try:
        rows = con.execute("SELECT length_dm, load_code, price FROM prices")
        return {(int(length), int(load)): float(price) for length, load, price in rows}
    finally:
        con.close()


def run_opt(pieces: list[dict], *, top_up: bool = False) -> tuple[dict | None, list, str]:
    """Оптимизатор и раскладка «сильные первыми». Возвращает план резов и дорожки.

    При ``top_up`` после нарезки хвост каждой дорожки добирается целыми плитами
    со следующих дорожек до 101 м — тот же шаг, что в производственном плане.
    """
    from core.domain.plate_order import PlateOrder
    from core.optimization import optimize_with_cascading_longitudinal_cuts
    from core.optimization.layout_runtime_snapshot import (
        build_layout_runtime_snapshot_from_plate_order_context,
    )
    from core.optimization.result_contract import is_optimization_success
    from core.plate_order_context import PlateOrderContext
    from core.visualization import LayoutIntegrityError, TrackLayoutInvariantError, split_sequence_into_tracks
    from viz_modules.layout_sequence.builder import build_layout_sequence

    if not pieces:
        return None, [], "нет плит с нормой проволок"
    grouped: dict[tuple, int] = defaultdict(int)
    for piece in pieces:
        grouped[(round(piece["length"], 3), piece["width_mm"], piece["load"])] += 1
    orders_2d = []
    for (length, width_mm, load), qty in grouped.items():
        orders_2d.append({
            "length": length,
            "width": width_mm,
            "qty": qty,
            "load_code": load,
            "plate_name": f"ПБ {int(round(length * 10))}-{width_mm / 100:.2g}-{load}п",
        })

    logging.disable(logging.CRITICAL)
    optimization_result = None
    buf = io.StringIO()
    try:
        with contextlib.redirect_stdout(buf):
            core_order = PlateOrder.from_orders_2d(orders_2d)
            plate_ctx = PlateOrderContext.fresh_empty()
            plate_ctx.hydrate_from_order(core_order)
            with plate_ctx.bound():
                optimization_result = optimize_with_cascading_longitudinal_cuts(orders_2d=orders_2d)
            if not is_optimization_success(optimization_result) or not optimization_result.get("total_plates"):
                code = optimization_result.get("_opt_error_code") or optimization_result.get("_opt_status")
                return optimization_result, [], f"optimize:{code}"
            loads = sorted({int(float(item.get("load_code", 8))) for item in orders_2d})
            plate_ctx.load_optimization_snapshot(
                optimization_result=optimization_result,
                plan_by_load={"all": optimization_result},
                load_to_reinforcement_map={load: ["all"] for load in loads},
            )
            with plate_ctx.bound():
                layout_rt = build_layout_runtime_snapshot_from_plate_order_context(
                    plate_ctx,
                    layout_reinforcement_order="desc",
                )
                seq = build_layout_sequence(runtime=layout_rt)
                try:
                    tracks = split_sequence_into_tracks(seq, strict_layout_integrity=False) or []
                except (LayoutIntegrityError, TrackLayoutInvariantError) as exc:
                    return optimization_result, [], f"split:{exc}"
            if top_up and tracks:
                from core.track_top_up import top_up_tracks_from_following

                top_up_tracks_from_following(tracks)
        if not tracks:
            return optimization_result, [], "раскладка пустая"
        return optimization_result, tracks, ""
    except Exception as exc:
        return optimization_result, [], f"{type(exc).__name__}: {exc}"
    finally:
        logging.disable(logging.NOTSET)


def _money_of_layout(laid: list[dict], lookup) -> tuple[list[dict], list[dict]]:
    rows = []
    got: list[dict] = []
    for index, track in enumerate(laid, 1):
        pieces, extra_scrap = order_pieces_from_layout(track, lookup)
        got.extend(pieces)
        stats = money_and_waste(pieces, track.get("length"))
        stats["scrap_m2"] += extra_scrap
        poured = track.get("length") or stats["occupied_m"]
        rows.append({
            "idx": index,
            "label": track.get("label") or "",
            "meters": round(float(poured or 0), 2),
            "qty": stats["qty"],
            "composition": stats["composition"],
            "over_bc_m": round(stats["over_bc_m"], 1),
            "plate_m": stats["plate_m"],
            "bc_pct": stats["bc_pct"],
            "concrete_rub": round(stats["concrete_rub"]),
            "delta_avg": round(stats["delta_avg"], 1),
            "delta_max": int(stats["delta_max"]) if stats["qty"] else 0,
            "over_wire_m": round(stats["over_wire_m"]),
            "wire_rub": round(stats["wire_rub"]),
            "total_rub": round(stats["wire_rub"] + stats["concrete_rub"]),
            "unpaired": layout_unpaired(pieces),
            "strip_m2": round(stats["strip_m2"], 2),
            "scrap_m2": round(stats["scrap_m2"] + stats["end_gap_m2"], 2),
        })
    return rows, got


def evaluate_pool(pieces: list[dict], lookup, prices: dict, *, top_up: bool = False) -> dict:
    started = time.perf_counter()
    opt, laid, err = run_opt(pieces, top_up=top_up)
    elapsed = time.perf_counter() - started
    accounts = cut_accounts(
        (opt or {}).get("primary_cuts"),
        (opt or {}).get("secondary_cuts"),
    )
    track_rows, got = _money_of_layout(laid, lookup) if laid else ([], [])
    match = bool(track_rows) and signature(got) == signature(pieces)
    note = err
    if track_rows and not match:
        note = (note + "; " if note else "") + "набор плит отличается от входа"
    return {
        "ok": bool(track_rows),
        "seconds": round(elapsed, 1),
        "note": note,
        "match": match,
        "tracks": track_rows,
        "meters": round(sum(row["meters"] for row in track_rows), 2),
        "qty": sum(row["qty"] for row in track_rows),
        "plate_m": sum(row["plate_m"] for row in track_rows),
        "over_bc_m": round(sum(row["over_bc_m"] for row in track_rows), 1),
        "concrete_rub": sum(row["concrete_rub"] for row in track_rows),
        "wire_rub": sum(row["wire_rub"] for row in track_rows),
        "over_wire_m": sum(row["over_wire_m"] for row in track_rows),
        "delta_max": max((row["delta_max"] for row in track_rows), default=0),
        "unpaired": accounts["unpaired"],
        "rest_m2": round(accounts["rest_m2"], 2),
        "scrap_m2": round(accounts["scrap_m2"], 2),
        "keys": accounts["keys"],
        "secondary": (opt or {}).get("secondary_cuts") or [],
        "prices": prices,
    }


def rescued_rub(batch: dict, available: Counter, costs: dict[tuple[int, float], float]) -> tuple[int, int, int, int]:
    plates = rescued_secondaries(available, batch.get("secondary"))
    sale = cost = 0.0
    missing_sale = missing_cost = 0
    for plate in plates:
        price = catalog_sale_price(batch["prices"], plate["length"], plate["width_mm"], plate["load"])
        if price is None:
            missing_sale += 1
        else:
            sale += price
        unit = plate_cost(costs, plate["length"], plate["width_mm"], plate["load"])
        if unit is None:
            missing_cost += 1
        else:
            cost += unit
    return int(round(sale)), int(round(cost)), missing_sale, missing_cost


def _delta_cells(batch: dict, base: dict | None, base_name: str, rescued: int | None, missing: int) -> dict:
    if not batch["ok"] or base is None:
        return {
            "base": base_name,
            "d_m": None,
            "d_wire": None,
            "d_concrete": None,
            "d_wire_m": None,
            "d_concrete_m": None,
            "rescued": rescued,
            "missing": missing,
        }
    return {
        "base": base_name,
        "d_m": round(batch["meters"] - base["meters"], 2),
        "d_wire": rub_delta(batch["wire_rub"], base["wire_rub"]),
        "d_concrete": rub_delta(batch["concrete_rub"], base["concrete_rub"]),
        "d_wire_m": per_m_delta(batch["wire_rub"], batch["meters"], base["wire_rub"], base["meters"]),
        "d_concrete_m": per_m_delta(
            batch["concrete_rub"], batch["meters"], base["concrete_rub"], base["meters"],
        ),
        "rescued": rescued,
        "missing": missing,
    }


def _note(batch: dict, delta: dict) -> str:
    parts = []
    if batch.get("note"):
        parts.append(batch["note"])
    if delta.get("missing"):
        parts.append(f"без цены: {delta['missing']} плит из остатка")
    if delta.get("missing_cost"):
        parts.append(f"без себестоимости: {delta['missing_cost']} плит из остатка")
    return "; ".join(parts)


def detail_rows(pool_label: str, batches: list[dict], bases: list[dict | None], base_name: str, rescued_flags: list[tuple]) -> list[list]:
    rows: list[list] = []
    for number, (batch, base, rescued_pair) in enumerate(zip(batches, bases, rescued_flags), 1):
        rescued, cost, missing, missing_cost = rescued_pair
        delta = _delta_cells(batch, base, base_name, rescued, missing)
        delta["missing_cost"] = missing_cost
        pct = (batch["over_bc_m"] / batch["plate_m"]) if batch["plate_m"] else 0.0
        rows.append([
            "пачка", number, batch.get("input_label") or "", "", batch.get("date") or "", pool_label,
            batch["meters"] if batch["ok"] else None,
            batch["qty"] if batch["ok"] else None,
            "",
            batch["over_bc_m"] if batch["ok"] else None,
            pct if batch["ok"] else None,
            batch["concrete_rub"] if batch["ok"] else None,
            None,
            batch["delta_max"] if batch["ok"] else None,
            batch["over_wire_m"] if batch["ok"] else None,
            batch["wire_rub"] if batch["ok"] else None,
            (batch["wire_rub"] + batch["concrete_rub"]) if batch["ok"] else None,
            batch["unpaired"] if batch["ok"] else None,
            batch["rest_m2"] if batch["ok"] else None,
            batch["scrap_m2"] if batch["ok"] else None,
            delta["base"], delta["d_m"], delta["d_wire"], delta["d_concrete"],
            delta["d_wire_m"], delta["d_concrete_m"],
            delta["rescued"], cost,
            "да" if batch["match"] else "нет",
            _note(batch, delta),
        ])
        for track in batch["tracks"]:
            rows.append([
                "дорожка", number, "", track["idx"], "", track["label"],
                track["meters"], track["qty"], track["composition"],
                track["over_bc_m"], track["bc_pct"], track["concrete_rub"],
                track["delta_avg"], track["delta_max"], track["over_wire_m"],
                track["wire_rub"], track["total_rub"],
                track["unpaired"], track["strip_m2"], track["scrap_m2"],
                None, None, None, None, None, None, None, None, None, None,
            ])
    return rows


def _against(batches: list[dict], base_field: str) -> tuple:
    paired = [batch for batch in batches if batch["ok"] and batch.get(base_field)]
    if not paired:
        return (None, None, None, None, None)
    meters = sum(batch["meters"] for batch in paired)
    wire = sum(batch["wire_rub"] for batch in paired)
    concrete = sum(batch["concrete_rub"] for batch in paired)
    base_m = sum(batch[base_field]["meters"] for batch in paired)
    base_w = sum(batch[base_field]["wire_rub"] for batch in paired)
    base_c = sum(batch[base_field]["concrete_rub"] for batch in paired)
    return (
        round(meters - base_m, 1),
        rub_delta(wire, base_w),
        per_m_delta(wire, meters, base_w, base_m),
        rub_delta(concrete, base_c),
        per_m_delta(concrete, meters, base_c, base_m),
    )


def summarize(label: str, batches: list[dict], note: str, *, tail: int) -> dict:
    ok = [batch for batch in batches if batch["ok"]]
    meters = sum(batch["meters"] for batch in ok)
    wire = sum(batch["wire_rub"] for batch in ok)
    concrete = sum(batch["concrete_rub"] for batch in ok)
    failed_batches = [batch for batch in batches if not batch["ok"] and batch.get("base_map")]
    meters += sum(batch["base_map"]["meters"] for batch in failed_batches)
    wire += sum(batch["base_map"]["wire_rub"] for batch in failed_batches)
    concrete += sum(batch["base_map"]["concrete_rub"] for batch in failed_batches)
    text = note
    if failed_batches:
        text = (text + "; " if text else "") + (
            f"{len(failed_batches)} пачек не разложились по одной и в итоге оставлены как карта"
        )
    map_pack = _against(batches, "base_map")
    single_pack = _against(batches, "base_single")
    return {
        "label": label,
        "input_tracks": sum(batch.get("input_count") or 0 for batch in batches),
        "output_tracks": sum(len(batch["tracks"]) for batch in ok),
        "tail": tail,
        "batches": len(batches),
        "matched": sum(1 for batch in ok if batch["match"]),
        "meters": round(meters, 1),
        "d_m_map": map_pack[0],
        "d_m_single": single_pack[0],
        "wire_rub": wire,
        "wire_per_m": round(per_m(wire, meters), 2),
        "d_wire_map": map_pack[1],
        "d_wire_single": single_pack[1],
        "d_wire_m_single": single_pack[2],
        "concrete_rub": concrete,
        "concrete_per_m": round(per_m(concrete, meters), 2),
        "d_concrete_map": map_pack[3],
        "d_concrete_single": single_pack[3],
        "d_concrete_m_single": single_pack[4],
        "unpaired": sum(batch["unpaired"] for batch in ok),
        "rest_m2": round(sum(batch["rest_m2"] for batch in ok), 2),
        "scrap_m2": round(sum(batch["scrap_m2"] for batch in ok), 2),
        "rescued": sum(int(batch.get("rescued_rub") or 0) for batch in batches),
        "rescued_cost": sum(int(batch.get("rescued_cost") or 0) for batch in batches),
        "seconds": round(sum(batch["seconds"] for batch in batches), 1),
        "note": text,
    }


def input_label(tracks: list[dict], start: int, end: int) -> str:
    first, last = tracks[start], tracks[end - 1]
    span = f"{start + 1}–{end}"
    if start + 1 == end:
        return f"{span} · {first['date']} {first['sheet']}"
    return f"{span} · {first['date']} {first['sheet']} … {last['date']} {last['sheet']}"


def map_batch(track: dict, pieces: list[dict]) -> dict:
    stats = money_and_waste(pieces, track.get("total_len"))
    meters = float(track.get("total_len") or stats["occupied_m"] or 0)
    return {
        "ok": True,
        "meters": round(meters, 2),
        "wire_rub": round(stats["wire_rub"]),
        "concrete_rub": round(stats["concrete_rub"]),
        "input_count": 1,
    }


def save_book(summary_rows: list[dict], sheets: dict[str, list[list]], out: Path, *, top_up: bool = False) -> None:
    hdr_fill = PatternFill("solid", fgColor="1F4E79")
    hdr_font = Font(color="FFFFFF", bold=True, size=10)
    batch_fill = PatternFill("solid", fgColor="D6E3F0")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    money = "#,##0"
    pct = "0.0%"

    def paint_header(ws, ncols: int) -> None:
        for col in range(1, ncols + 1):
            cell = ws.cell(row=1, column=col)
            cell.fill = hdr_fill
            cell.font = hdr_font
            cell.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
            cell.border = border

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Сводка"
    ws.append(SUMMARY_HEADERS)
    paint_header(ws, len(SUMMARY_HEADERS))
    for row in summary_rows:
        ws.append([
            row["label"], row["input_tracks"], row["output_tracks"], row["tail"], row["batches"],
            row["matched"], row["meters"],
            row["d_m_map"], row["d_m_single"],
            row["wire_rub"], row["wire_per_m"],
            row["d_wire_map"], row["d_wire_single"], row["d_wire_m_single"],
            row["concrete_rub"], row["concrete_per_m"],
            row["d_concrete_map"], row["d_concrete_single"], row["d_concrete_m_single"],
            row["unpaired"], row["rest_m2"], row["scrap_m2"],
            row["rescued"], row["rescued_cost"], row["seconds"], row["note"],
        ])
    for col in range(1, len(SUMMARY_HEADERS) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["Z"].width = 55
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:Z{ws.max_row}"
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=10, max_col=24):
        for cell in row:
            if isinstance(cell.value, (int, float)) and cell.column not in (11, 14, 16, 19, 21, 22):
                cell.number_format = money

    for title, table in sheets.items():
        ws = wb.create_sheet(title[:31])
        ws.append(DETAIL_HEADERS)
        paint_header(ws, len(DETAIL_HEADERS))
        for values in table:
            ws.append(values)
            if values and values[0] == "пачка":
                for col in range(1, len(DETAIL_HEADERS) + 1):
                    ws.cell(row=ws.max_row, column=col).fill = batch_fill
        if ws.max_row > 1:
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=11, max_col=11):
                for cell in row:
                    if isinstance(cell.value, float):
                        cell.number_format = pct
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=12, max_col=28):
                for cell in row:
                    if isinstance(cell.value, (int, float)) and cell.column not in (13, 14, 19, 20, 22, 25, 26):
                        cell.number_format = money
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(DETAIL_HEADERS))}{max(ws.max_row, 1)}"
        ws.column_dimensions["C"].width = 42
        ws.column_dimensions["I"].width = 55
        ws.column_dimensions["AD"].width = 42
        for col in range(1, len(DETAIL_HEADERS) + 1):
            if get_column_letter(col) not in ("C", "I", "AD"):
                ws.column_dimensions[get_column_letter(col)].width = 14

    ws = wb.create_sheet("Методика")
    lines = [
        "Гипотеза программы: лист «по 1» дешевле строки «карта» по переармированию и перебетонированию.",
        "Гипотеза пула: размер 5, 10, 21, 30 или 60 дешевле листа «по 1» по этим двум рублям и по метражу заливки.",
        "Отрицательная дельта — дешевле базы. Эффект может быть и положительным: тогда программа или пул дороже.",
        "Цена плит из остатка — продажная цена плит, севших в остаток, который на листе «по 1» никто не забрал.",
        "Эта цена не складывается с дельтами: плита уже была в заказе, новой выручки нет.",
        "Прайс pb.db не знает ширину. Цена длины и нагрузки умножается на ширину / 1200. Нагрузка 13 считается как 12.",
        "Себестоимость плит из остатка — factory_plate_costs, заготовка 12 дм той же длины, целиком. Нагрузки 6 и 8 берутся как 8п, остальные как 12,5п. Отдельно эту плиту лили бы на 1200 мм, несъеденный остаток ширины — отход.",
        "Без пары на строке пачки — первичный рез решателя, чей остаток не стал родителем вторичного реза.",
        "Без пары на строке дорожки — полоса раскладки с зазором ширины от 200 мм.",
        "Пул режется подряд без перекрытия. Хвост года, не кратный размеру, идёт короткой последней пачкой.",
        "Режим раскладки: сильные первыми, жадное склеивание выключено.",
        (
            "После нарезки хвост дорожки добирается целыми плитами со следующих дорожек до 101 м. "
            "Первую плиту чужой дорожки и плиту с резом не переносят. Проволоки донора не выше максимума получателя."
            if top_up else
            "Дозаполнение хвоста до 101 м выключено: дорожка закрывается на последней плите, которая влезла в 101 м."
        ),
        "Пачка, где набор плит не совпал с входом, для выбора размера на формовке не годится.",
    ]
    for line in lines:
        ws.append([line])
    ws.column_dimensions["A"].width = 120

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def _sum_maps(maps: list[dict]) -> dict:
    return {
        "meters": round(sum(item["meters"] for item in maps), 2),
        "wire_rub": sum(item["wire_rub"] for item in maps),
        "concrete_rub": sum(item["concrete_rub"] for item in maps),
    }


def load_plate_costs() -> dict[tuple[int, float], float]:
    con = sqlite3.connect(f"file:{PRICE_DB_PATH}?mode=ro", uri=True)
    try:
        rows = con.execute(
            "SELECT length_dm, load_code, full_cost_with_kef FROM factory_plate_costs WHERE width_dm = 12"
        )
        return {(int(length), float(load)): float(cost) for length, load, cost in rows}
    finally:
        con.close()


def build(tracks: list[dict], lookup, prices: dict, costs: dict[tuple[int, float], float], sizes: tuple[int, ...], *, top_up: bool = False) -> tuple[list[dict], dict[str, list[list]]]:
    maps = []
    singles = []
    print(f"дорожек: {len(tracks)}", flush=True)
    for index, track in enumerate(tracks, 1):
        pieces = pieces_of(track, lookup)
        mapped = map_batch(track, pieces)
        mapped_eval = evaluate_pool(pieces, lookup, prices, top_up=top_up)
        mapped_eval["input_count"] = 1
        mapped_eval["input_label"] = input_label(tracks, index - 1, index)
        mapped_eval["date"] = track["date"]
        maps.append(mapped)
        singles.append(mapped_eval)
        if index == 1 or index % 20 == 0 or index == len(tracks):
            print(f"  по 1: {index}/{len(tracks)}  {track['date']}  {mapped_eval['seconds']} с", flush=True)

    map_totals = _sum_maps(maps)
    sheets: dict[str, list[list]] = {}
    summary: list[dict] = []

    summary.append({
        "label": "карта",
        "input_tracks": len(tracks),
        "output_tracks": len(tracks),
        "tail": 0,
        "batches": len(tracks),
        "matched": "",
        "meters": round(map_totals["meters"], 1),
        "d_m_map": None,
        "d_m_single": None,
        "wire_rub": map_totals["wire_rub"],
        "wire_per_m": round(per_m(map_totals["wire_rub"], map_totals["meters"]), 2),
        "d_wire_map": None,
        "d_wire_single": None,
        "d_wire_m_single": None,
        "concrete_rub": map_totals["concrete_rub"],
        "concrete_per_m": round(per_m(map_totals["concrete_rub"], map_totals["meters"]), 2),
        "d_concrete_map": None,
        "d_concrete_single": None,
        "d_concrete_m_single": None,
        "unpaired": None,
        "rest_m2": None,
        "scrap_m2": None,
        "rescued": None,
        "rescued_cost": None,
        "seconds": 0,
        "note": "фактические дорожки КЗ, без оптимизатора. База эффекта программы.",
    })

    bases = []
    flags = [(None, None, 0, 0) for _ in singles]
    for batch, mapped in zip(singles, maps):
        batch["rescued_rub"] = 0
        batch["rescued_cost"] = 0
        batch["base_map"] = mapped
        bases.append(mapped if batch["ok"] else None)
    sheets["по 1"] = detail_rows("по 1", singles, bases, "карта", flags)
    summary.append(summarize(
        "по 1", singles,
        "эффект программы — дельты к карте. Отрицательное значит дешевле карты.",
        tail=0,
    ))
    _dump(summary, sheets)

    for size in sizes:
        ranges = split_batches(len(tracks), size)
        batches = []
        bases = []
        flags = []
        print(f"пул {size}: пачек {len(ranges)}", flush=True)
        for start, end in ranges:
            pieces = [piece for track in tracks[start:end] for piece in pieces_of(track, lookup)]
            batch = evaluate_pool(pieces, lookup, prices, top_up=top_up)
            batch["input_count"] = end - start
            batch["input_label"] = input_label(tracks, start, end)
            batch["date"] = tracks[start]["date"]
            available: Counter = Counter()
            base_parts = []
            fallback = 0
            for i in range(start, end):
                if singles[i]["ok"]:
                    base_parts.append(singles[i])
                    available.update(singles[i]["keys"])
                else:
                    base_parts.append(maps[i])
                    fallback += 1
            rescued, cost, missing, missing_cost = (
                rescued_rub(batch, available, costs) if batch["ok"] else (None, None, 0, 0)
            )
            batch["rescued_rub"] = rescued or 0
            batch["rescued_cost"] = cost or 0
            batch["base_single"] = _sum_maps(base_parts)
            batch["base_map"] = _sum_maps(maps[start:end])
            if fallback:
                extra = f"в базе {fallback} дор. как карта: по одной не разложились"
                batch["note"] = f"{batch['note']}; {extra}" if batch.get("note") else extra
            batches.append(batch)
            bases.append(batch["base_single"])
            flags.append((rescued, cost, missing, missing_cost))
            print(
                f"  {start + 1}–{end}: дорожек {len(batch['tracks'])}  "
                f"{batch['seconds']} с  совпал={batch['match']}"
                + (f"  {batch['note']}" if batch["note"] else ""),
                flush=True,
            )
        sheets[f"по {size}"] = detail_rows(f"по {size}", batches, bases, "по 1", flags)
        summary.append(summarize(
            f"по {size}", batches,
            "эффект пула — дельты к «по 1». Цена плит из остатка в экономию не входит.",
            tail=tail_tracks(len(tracks), size),
        ))
        _dump(summary, sheets)
    return summary, sheets


def _dump(summary: list[dict], sheets: dict[str, list[list]]) -> None:
    payload = {"summary": summary, "sheets": {name: rows for name, rows in sheets.items()}}
    JSON_OUT.parent.mkdir(parents=True, exist_ok=True)
    JSON_OUT.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")


def _cli() -> None:
    parser = argparse.ArgumentParser(description="Пул пар и перерасход дорожек КЗ")
    parser.add_argument("--limit", type=int, default=None)
    parser.add_argument("--sizes", type=str, default=",".join(str(size) for size in POOL_SIZES))
    parser.add_argument("--out", type=Path, default=None)
    parser.add_argument(
        "--top-up",
        action="store_true",
        help="добирать хвост дорожки целыми плитами со следующих до 101 м",
    )
    global JSON_OUT
    args = parser.parse_args()
    sizes = tuple(int(part) for part in args.sizes.split(",") if part.strip())
    if args.top_up:
        JSON_OUT = JSON_OUT.with_name("kz_pool_pair_2026_topup.json")
    out = args.out or (
        OUT.with_name("Пул_пары_перерасход_КЗ_2026_дозаполнение.xlsx") if args.top_up else OUT
    )
    lookup = make_lookup(load_wires())
    tracks = json.load(open(CACHE))
    if args.limit:
        tracks = tracks[: args.limit]
    prices = load_price_table()
    costs = load_plate_costs()
    try:
        summary, sheets = build(tracks, lookup, prices, costs, sizes, top_up=args.top_up)
    except Exception:
        traceback.print_exc()
        raise
    save_book(summary, sheets, out, top_up=args.top_up)
    print(f"[OK] {out}", flush=True)


if __name__ == "__main__":
    _cli()
