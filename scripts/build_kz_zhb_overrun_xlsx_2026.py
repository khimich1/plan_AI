# -*- coding: utf-8 -*-
"""Excel-отчёт перебетонирования сборного ЖБ по КЗ за 2026 год.

Источник: «02 КЗ на ЖБ» (листы «ФС.КЗ.День»): марка в заявке vs марка в формовке.
Цены м³ по классам — рецептура pb.db (pile_cost_recipe) × цены 08.10.2026
(цемент 7.54, песок 0.66, щебень 1.8), как в файле-образце пересчёта.

Листы:
  1. «Сводка по номенклатуре»
  2. «Строки КЗ 2026»
  3. «Нормативы»
"""
from __future__ import annotations

import datetime
import glob
import json
import os
import re
import sqlite3
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DB = Path("/home/roman/project/Шишов/pb.db")
KZ_DIR = Path("/home/roman/Рабочий стол/с производства/02 КЗ на ЖБ")
CACHE = Path("/home/roman/project/Шишов/data/kz_zhb_rows_cache.json")
OUT = Path("/home/roman/Рабочий стол/с производства/Перебетонирование_ЖБ_КЗ_2026.xlsx")

CEMENT, SAND, CRUSHED = 7.54, 0.66, 1.8   # цены 08.10.2026, без НДС

GRADE_RE = re.compile(r"^[ВB]\s*(\d{2}(?:[.,]\d)?)$")
NOM_RE = re.compile(r"Сваи|марш|ступен|ФБС|КРОССБЛОК|КБ\s|плита|ПБ|перемыч|ригель|прогон|колонн|балк|панель|блок", re.I)


def grade_num(s: str) -> float | None:
    m = GRADE_RE.match(str(s).strip())
    if not m:
        return None
    return float(m.group(1).replace(",", "."))


def m3_prices() -> dict[float, float]:
    """Цена м³ по классам: рецептура × цены 08.10.2026."""
    con = sqlite3.connect(f"file:{DB}?immutable=1", uri=True)
    prices = {}
    for grade, cem, sand, crush in con.execute(
        "SELECT concrete_grade, cement_kg_m3, sand_kg_m3, crushed_kg_m3 FROM pile_cost_recipe"
    ):
        key = grade.replace("B", "").replace("_", ".")
        try:
            g = float(key.replace("granite", "").rstrip("."))
        except ValueError:
            g = {"B30_granite": 30.0}.get(grade)
        if g is None:
            continue
        prices[g] = cem * CEMENT + sand * SAND + crush * CRUSHED
    con.close()
    return prices


def parse_kz() -> list[dict]:
    if CACHE.exists():
        return json.load(open(CACHE))
    files = sorted(glob.glob(str(KZ_DIR / "**" / "*.xlsx"), recursive=True))
    rows = []
    cutoff = datetime.date(2026, 1, 1)
    for f in files:
        try:
            wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        except Exception:
            continue
        for sn in wb.sheetnames:
            if not re.match(r"ФС\.КЗ\.День", sn):
                continue
            date_obj = None
            for row in wb[sn].iter_rows(values_only=True):
                vals = [c for c in row]
                # дата формовки
                if date_obj is None:
                    for i, cell in enumerate(vals):
                        if isinstance(cell, str) and "Дата формовки" in cell:
                            for c2 in vals[i + 1:]:
                                if isinstance(c2, datetime.datetime):
                                    date_obj = c2.date()
                                    break
                # строка номенклатуры: ищем пару марок подряд в колонках
                for i in range(len(vals) - 1):
                    g1 = grade_num(vals[i]) if isinstance(vals[i], str) else None
                    g2 = grade_num(vals[i + 1]) if isinstance(vals[i + 1], str) else None
                    if g1 is not None and g2 is not None:
                        name = ""
                        order = ""
                        for j in range(i - 1, -1, -1):
                            if isinstance(vals[j], str) and NOM_RE.search(vals[j] or ""):
                                name = re.sub(r"\s+", " ", vals[j]).strip()
                            elif vals[j] not in (None, "") and not name:
                                order = str(vals[j]).strip()
                        qty, vol = None, None
                        for c2 in vals[i + 2:]:
                            if isinstance(c2, (int, float)) and qty is None and 0 < c2 < 10000:
                                qty = c2
                            elif isinstance(c2, (int, float)) and qty is not None and vol is None and 0 < c2 < 10000:
                                vol = float(c2)
                                break
                        if name and vol:
                            rows.append({
                                "date": str(date_obj) if date_obj else "",
                                "order": order, "name": name,
                                "g_order": g1, "g_fact": g2,
                                "qty": qty, "vol": vol,
                                "file": os.path.basename(f),
                            })
                        break
        wb.close()
    rows = [r for r in rows if r["date"] >= "2026-01-01"]
    rows.sort(key=lambda r: r["date"])
    CACHE.parent.mkdir(parents=True, exist_ok=True)
    json.dump(rows, open(CACHE, "w"), ensure_ascii=False)
    return rows


def main() -> None:
    prices = m3_prices()
    print("цены м³ по классам:", {g: round(p) for g, p in sorted(prices.items())})
    rows = parse_kz()
    print(f"строк КЗ 2026: {len(rows)}")

    # ---------- расчёт ----------
    agg = defaultdict(lambda: {"vol": 0.0, "qty": 0.0, "over_vol": 0.0, "over_rub": 0.0,
                               "under_vol": 0.0, "under_rub": 0.0, "n_over": 0, "n_under": 0})
    detail = []
    skipped_no_price = 0
    for r in rows:
        p_o = prices.get(r["g_order"])
        p_f = prices.get(r["g_fact"])
        if p_o is None or p_f is None:
            skipped_no_price += 1
            continue
        delta_m3 = p_f - p_o
        delta_rub = delta_m3 * r["vol"]
        a = agg[r["name"]]
        a["vol"] += r["vol"]
        a["qty"] += r["qty"] or 0
        if delta_m3 > 0.01:
            a["over_vol"] += r["vol"]
            a["over_rub"] += delta_rub
            a["n_over"] += 1
        elif delta_m3 < -0.01:
            a["under_vol"] += r["vol"]
            a["under_rub"] += delta_rub
            a["n_under"] += 1
        detail.append({**r, "p_o": p_o, "p_f": p_f,
                       "cost_o": p_o * r["vol"], "cost_f": p_f * r["vol"],
                       "delta": delta_rub})
    print(f"без цены (класс вне таблицы): {skipped_no_price}")

    # ---------- стили ----------
    hdr_fill = PatternFill("solid", fgColor="1F4E79")
    hdr_font = Font(color="FFFFFF", bold=True, size=10)
    tot_fill = PatternFill("solid", fgColor="D9E1F2")
    warn_font = Font(color="9C0006", bold=True)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    money = '#,##0'

    def style_header(ws, ncols, row=1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=row, column=c)
            cell.fill = hdr_fill
            cell.font = hdr_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border

    def autowidth(ws, widths):
        for i, w_ in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w_

    wb = openpyxl.Workbook()

    # ======== лист 1: сводка ========
    ws1 = wb.active
    ws1.title = "Сводка по номенклатуре"
    h1 = ["Наименование", "Кол-во", "Объём, м³",
          "Перебетонир.: раз", "Перебетонир.: м³", "Перебетонир.: руб",
          "Ниже заявки: раз", "Ниже заявки: м³", "Ниже заявки: руб",
          "Итого дельта, руб"]
    ws1.append(h1)
    style_header(ws1, len(h1))
    items = sorted(agg.items(), key=lambda kv: -(kv[1]["over_rub"]))
    for name, a in items:
        ws1.append([name, a["qty"], round(a["vol"], 2), a["n_over"], round(a["over_vol"], 2),
                    round(a["over_rub"]), a["n_under"], round(a["under_vol"], 2),
                    round(a["under_rub"]), round(a["over_rub"] + a["under_rub"])])
    tr = ws1.max_row + 1
    ws1.cell(row=tr, column=1, value="ИТОГО")
    for col, key in [(2, "qty"), (3, "vol"), (4, "n_over"), (5, "over_vol"), (6, "over_rub"),
                     (7, "n_under"), (8, "under_vol"), (9, "under_rub")]:
        v = sum(a[key] for _, a in items)
        ws1.cell(row=tr, column=col, value=round(v, 2) if isinstance(v, float) else v)
    ws1.cell(row=tr, column=10, value=round(sum(a["over_rub"] + a["under_rub"] for _, a in items)))
    for c in range(1, len(h1) + 1):
        cell = ws1.cell(row=tr, column=c)
        cell.font = Font(bold=True)
        cell.fill = tot_fill
        cell.border = border
        if c in (6, 9, 10):
            cell.number_format = money
    for r in range(2, tr):
        for c in range(1, len(h1) + 1):
            cell = ws1.cell(row=r, column=c)
            cell.border = border
            if c in (6, 9, 10):
                cell.number_format = money
    autowidth(ws1, [44, 9, 11, 12, 12, 14, 12, 12, 13, 14])
    ws1.freeze_panes = "A2"
    ws1.auto_filter.ref = f"A1:J{tr - 1}"

    # ======== лист 2: строки КЗ ========
    ws2 = wb.create_sheet("Строки КЗ 2026")
    h2 = ["Дата", "Заказ", "Наименование", "Марка в заявке", "Марка в формовке",
          "Кол-во", "Объём, м³", "Цена м³ (заявка), руб", "Цена м³ (факт), руб",
          "Стоимость по заявке, руб", "Стоимость по факту, руб", "Дельта, руб"]
    ws2.append(h2)
    style_header(ws2, len(h2))
    for d in detail:
        g_o = f"В{d['g_order']:g}".replace(".", ",")
        g_f = f"В{d['g_fact']:g}".replace(".", ",")
        ws2.append([d["date"], d["order"], d["name"], g_o, g_f, d["qty"], round(d["vol"], 2),
                    round(d["p_o"]), round(d["p_f"]), round(d["cost_o"]), round(d["cost_f"]),
                    round(d["delta"])])
    tr2 = ws2.max_row + 1
    ws2.cell(row=tr2, column=1, value="ИТОГО")
    ws2.cell(row=tr2, column=7, value=round(sum(d["vol"] for d in detail), 2))
    for col, key in [(10, "cost_o"), (11, "cost_f"), (12, "delta")]:
        ws2.cell(row=tr2, column=col, value=round(sum(d[key] for d in detail)))
    for c in range(1, len(h2) + 1):
        cell = ws2.cell(row=tr2, column=c)
        cell.font = Font(bold=True)
        cell.fill = tot_fill
        cell.border = border
        if c >= 8:
            cell.number_format = money
    for r in range(2, tr2):
        for c in range(1, len(h2) + 1):
            cell = ws2.cell(row=r, column=c)
            cell.border = border
            if c >= 8:
                cell.number_format = money
        # подсветка строк с несовпадением марок
        if ws2.cell(row=r, column=4).value != ws2.cell(row=r, column=5).value:
            for c in range(1, len(h2) + 1):
                ws2.cell(row=r, column=c).font = warn_font
    autowidth(ws2, [11, 10, 40, 12, 13, 8, 10, 14, 14, 15, 15, 12])
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:L{tr2 - 1}"

    # ======== лист 3: нормативы ========
    ws3 = wb.create_sheet("Нормативы")
    rows3 = [
        ["ЦЕНЫ 08.10.2026 (без НДС) — «Расчет новых цен на ПБ 08.10.2026.xls»", "", ""],
        ["цемент, руб/кг", CEMENT, ""],
        ["песок, руб/кг", SAND, ""],
        ["щебень фр. 5-20, руб/кг", CRUSHED, ""],
        ["", "", ""],
        ["РЕЦЕПТУРА (pb.db, pile_cost_recipe) и цена м³", "", ""],
        ["класс", "цемент/песок/щебень, кг", "цена м³, руб"],
    ]
    for g, p in sorted(prices.items()):
        rows3.append([f"В{g:g}".replace(".", ","), "", round(p)])
    rows3 += [
        ["", "", ""],
        ["МЕТОДИКА", "", ""],
        ["Перебетонирование", "марка в формовке выше марки в заявке; дельта = (цена факт − цена заявка) × объём", ""],
        ["«Ниже заявки»", "марка в формовке НИЖЕ заявки — недобор марки, риск качества (дельта отрицательная)", ""],
        ["Строки без цены", f"класс вне рецептуры pb.db — пропущено: {skipped_no_price}", ""],
        ["", "", ""],
        ["Период", "01.01.2026 – 15.10.2026", ""],
        ["Строк КЗ в отчёте", len(detail), ""],
    ]
    for row in rows3:
        ws3.append(row)
    ws3.cell(row=1, column=1).font = Font(bold=True, size=12)
    ws3.cell(row=6, column=1).font = Font(bold=True, size=12)
    for c in range(1, 4):
        cell = ws3.cell(row=7, column=c)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.border = border
    autowidth(ws3, [46, 40, 14])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"[OK] сохранено: {OUT}")
    print(f"итого перебетонирование: {sum(a['over_rub'] for _, a in items):,.0f} руб")
    print(f"итого ниже заявки: {sum(a['under_rub'] for _, a in items):,.0f} руб")


if __name__ == "__main__":
    main()
