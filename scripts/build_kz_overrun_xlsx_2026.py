# -*- coding: utf-8 -*-
"""Excel-отчёт перерасхода дорожек ПБ по КЗ за календарный год.

По умолчанию — 2026. Другой год: ``python scripts/build_kz_overrun_xlsx_2026.py --year 2022``.

Обновлено 15.10.2026 под нормы файла «Расчет новых цен на ПБ 08.10.2026.xls»:
  проволока Вр 1400 — 100.33 руб/кг (0.17 кг/м);
  дельта B40−B30 = +20 кг цемента − 30 кг песка = 131.0 руб/м³;
  марки: нагрузки 6/8 — М400 до ПБ 66, М500 с ПБ 67;
         нагрузки 10/12.5/16/21 — всегда М500 (по «Расход номинал»).

Листы:
  1. «Сводка по номенклатуре»
  2. «Дорожки 2026»
  3. «Нормативы»
"""
from __future__ import annotations

import argparse
import datetime
import glob
import json
import re
import sqlite3
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DB = Path("/home/roman/project/Шишов/pb.db")
KZ_DIR = Path("/home/roman/Рабочий стол/с производства/03 КЗ на ПБ")
CACHE = Path("/home/roman/project/Шишов/data/kz_tracks_cache.json")
OUT = Path("/home/roman/Рабочий стол/с производства/Перерасход_дорожек_КЗ_2026.xlsx")

# --- нормативы 08.10.2026 ---
WIRE_KG_M = 0.17
WIRE_PRICE_KG = 100.33
WIRE_PRICE_M = WIRE_KG_M * WIRE_PRICE_KG          # 17.06 руб/м
VOL_M_PER_M_12 = 0.17                              # м³/м для ширины 12 дм
CEMENT, SAND = 7.54, 0.66
DELTA_B40_B30 = 20 * CEMENT - 30 * SAND            # 131.0 руб/м³
GRADE_RANK = {"М400": 400, "M400": 400, "M500": 500, "М500": 500}

PLATE_RE = re.compile(r"ПБ\s*(\d+(?:[.,]\d+)?)\s*-\s*(\d+(?:[.,]\d+)?)\s*-\s*(\d+(?:[.,]\d+)?)\s*п", re.I)


def grade_of(l_dm: int, load: int) -> str:
    """Марка по правилам файла 08.10.2026 («Расход номинал»)."""
    if load in (6, 8):
        return "М400" if l_dm <= 66 else "М500"
    return "М500"


def parse_kz(year: int = 2026, kz_dir: Path | None = None, cache: Path | None = None) -> list[dict]:
    """Парсинг КЗ ПБ за календарный год (с кэшем в data/)."""
    kz_dir = Path(kz_dir) if kz_dir else KZ_DIR
    cache = Path(cache) if cache else CACHE
    start = datetime.date(year, 1, 1)
    end = datetime.date(year + 1, 1, 1)
    start_s, end_s = start.isoformat(), end.isoformat()
    if cache.exists():
        data = json.load(open(cache))
        if data and str(data[0].get("date", "")).startswith(str(year)):
            return [t for t in data if start_s <= t["date"] < end_s]
    files = sorted(glob.glob(str(kz_dir / "**" / "*.xlsx"), recursive=True))
    tracks = []
    for f in files:
        try:
            wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        except Exception:
            continue
        for sn in wb.sheetnames:
            if not re.match(r"ФП\d\.КЗ\.День", sn):
                continue
            date_obj, plates, total_len = None, [], None
            for row in wb[sn].iter_rows(values_only=True):
                for i, cell in enumerate(row):
                    if cell is None:
                        continue
                    if isinstance(cell, str):
                        if date_obj is None and "Дата формовки" in cell:
                            for c2 in row[i + 1:]:
                                if isinstance(c2, datetime.datetime):
                                    date_obj = c2.date()
                                    break
                        if "Общий метраж" in cell:
                            for c2 in row[i + 1:]:
                                if isinstance(c2, (int, float)) and 10 < c2 < 200:
                                    total_len = float(c2)
                                    break
                        m = PLATE_RE.search(cell)
                        if m:
                            qty = None
                            for c2 in row[i + 1:]:
                                if isinstance(c2, (int, float)) and 0 < c2 < 500:
                                    qty = int(c2)
                                    break
                            if qty:
                                plates.append([
                                    float(m.group(1).replace(",", ".")),
                                    float(m.group(2).replace(",", ".")),
                                    float(m.group(3).replace(",", ".")),
                                    qty,
                                ])
            if date_obj and plates and start <= date_obj < end:
                tracks.append({
                    "date": str(date_obj), "sheet": sn,
                    "total_len": total_len, "plates": plates,
                })
        wb.close()
    tracks.sort(key=lambda t: (t["date"], t["sheet"]))
    cache.parent.mkdir(parents=True, exist_ok=True)
    json.dump(tracks, open(cache, "w"), ensure_ascii=False)
    return tracks


def load_wires() -> dict:
    con = sqlite3.connect(f"file:{DB}?immutable=1", uri=True)
    wires = {}
    for l_dm, load, w in con.execute(
        "SELECT length_dm, load_code, reinforcement_value FROM pb_reinforcement_series"
    ):
        wires[(int(l_dm), int(load))] = float(w)
    con.close()
    return wires


def make_lookup(wires):
    def lookup(l_dm: int, load: int):
        for key in ((l_dm, load), (l_dm, 13) if load == 12 else None,
                    (l_dm - 1, load), (l_dm + 1, load)):
            if key and key in wires:
                return wires[key]
        return None
    return lookup


def main(year: int = 2026, kz_dir: Path | None = None, cache: Path | None = None,
         out: Path | None = None) -> None:
    wires = load_wires()
    lookup = make_lookup(wires)
    tracks = parse_kz(year, kz_dir=kz_dir, cache=cache)
    print(f"дорожек {year}: {len(tracks)}")
    out = Path(out) if out else OUT.with_name(f"Перерасход_дорожек_КЗ_{year}.xlsx")

    rows_tracks = []
    agg = defaultdict(lambda: {
        "qty": 0, "meters": 0.0, "over_bc_m": 0.0, "bc_rub": 0.0,
        "over_wire_m": 0.0, "wire_rub": 0.0, "delta_w_qty": 0.0,
    })

    for t in tracks:
        plates = []
        for l_dm, width_dm, load, qty in t["plates"]:
            ld, lc = int(round(l_dm)), int(round(load))
            w = lookup(ld, lc)
            if w is None:
                continue
            plates.append({
                "name": f"ПБ {ld}-{int(round(width_dm))}-{lc}п",
                "L": l_dm / 10, "width": width_dm, "w": w,
                "g": grade_of(ld, lc), "qty": qty,
            })
        if not plates:
            continue

        tmax_w = max(p["w"] for p in plates)
        gmax = max(GRADE_RANK[p["g"]] for p in plates)
        tot_m = sum(p["L"] * p["qty"] for p in plates)
        total_len = t["total_len"] if t["total_len"] else tot_m

        over_bc_m = over_bc_vol = over_wire_m = delta_sum = 0.0
        over_qty = 0
        for p in plates:
            m = p["L"] * p["qty"]
            bc_rub_plate = 0.0
            if GRADE_RANK[p["g"]] < gmax:
                over_bc_m += m
                vol = VOL_M_PER_M_12 * (p["width"] / 12) * m
                over_bc_vol += vol
                bc_rub_plate = vol * DELTA_B40_B30
            wire_rub_plate = 0.0
            d = tmax_w - p["w"]
            if d > 0.01:
                wm = d * p["L"] * p["qty"]
                over_wire_m += wm
                wire_rub_plate = wm * WIRE_PRICE_M
                delta_sum += d * p["qty"]
                over_qty += p["qty"]
            a = agg[p["name"]]
            a["qty"] += p["qty"]
            a["meters"] += m
            a["over_bc_m"] += m if bc_rub_plate else 0.0
            a["bc_rub"] += bc_rub_plate
            a["over_wire_m"] += (d * m) if d > 0.01 else 0.0
            a["wire_rub"] += wire_rub_plate
            a["delta_w_qty"] += d * p["qty"] if d > 0.01 else 0.0

        bc_rub = over_bc_vol * DELTA_B40_B30
        wire_rub = over_wire_m * WIRE_PRICE_M
        rows_tracks.append({
            "date": t["date"], "sheet": t["sheet"], "total_len": total_len,
            "composition": "; ".join(f"{p['name']}×{p['qty']}" for p in plates),
            "qty": sum(p["qty"] for p in plates),
            "over_bc_m": over_bc_m, "bc_pct": over_bc_m / tot_m if tot_m else 0.0,
            "bc_rub": bc_rub,
            "delta_avg": delta_sum / over_qty if over_qty else 0.0,
            "delta_max": max((tmax_w - p["w"]) for p in plates),
            "over_wire_m": over_wire_m, "wire_rub": wire_rub,
            "total_rub": bc_rub + wire_rub,
        })

    # ---------- стили ----------
    hdr_fill = PatternFill("solid", fgColor="1F4E79")
    hdr_font = Font(color="FFFFFF", bold=True, size=10)
    tot_fill = PatternFill("solid", fgColor="D9E1F2")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    money = '#,##0'
    pct = '0.0%'

    def style_header(ws, ncols, row=1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=row, column=c)
            cell.fill = hdr_fill
            cell.font = hdr_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border

    def autowidth(ws, widths):
        for i, w_ in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w_

    wb = openpyxl.Workbook()

    # ======== лист 1: сводка по номенклатуре ========
    ws1 = wb.active
    ws1.title = "Сводка по номенклатуре"
    h1 = ["Номенклатура", "Плит, шт", "Метраж, м",
          "Перебетонир.: метраж, м", "% метража перебетонир.", "Цена перебетонир., руб",
          "Среднее переармир., проволок", "Лишняя проволока, м", "Цена переармир., руб",
          "Итого перерасход, руб"]
    ws1.append(h1)
    style_header(ws1, len(h1))
    items = sorted(agg.items(), key=lambda kv: -(kv[1]["bc_rub"] + kv[1]["wire_rub"]))
    for name, a in items:
        ws1.append([name, a["qty"], round(a["meters"], 1), round(a["over_bc_m"], 1),
                    a["over_bc_m"] / a["meters"] if a["meters"] else 0.0,
                    round(a["bc_rub"]),
                    round(a["delta_w_qty"] / a["qty"], 1) if a["qty"] else 0.0,
                    round(a["over_wire_m"]), round(a["wire_rub"]),
                    round(a["bc_rub"] + a["wire_rub"])])
    tr = ws1.max_row + 1
    ws1.cell(row=tr, column=1, value="ИТОГО")
    ws1.cell(row=tr, column=2, value=sum(a["qty"] for _, a in items))
    ws1.cell(row=tr, column=3, value=round(sum(a["meters"] for _, a in items), 1))
    ws1.cell(row=tr, column=4, value=round(sum(a["over_bc_m"] for _, a in items), 1))
    tm = sum(a["meters"] for _, a in items)
    tobc = sum(a["over_bc_m"] for _, a in items)
    ws1.cell(row=tr, column=5, value=tobc / tm if tm else 0.0)
    ws1.cell(row=tr, column=6, value=round(sum(a["bc_rub"] for _, a in items)))
    ws1.cell(row=tr, column=8, value=round(sum(a["over_wire_m"] for _, a in items)))
    ws1.cell(row=tr, column=9, value=round(sum(a["wire_rub"] for _, a in items)))
    ws1.cell(row=tr, column=10, value=round(sum(a["bc_rub"] + a["wire_rub"] for _, a in items)))
    for c in range(1, len(h1) + 1):
        cell = ws1.cell(row=tr, column=c)
        cell.font = Font(bold=True)
        cell.fill = tot_fill
        cell.border = border
        if c in (6, 9, 10):
            cell.number_format = money
        elif c == 5:
            cell.number_format = pct
    for r in range(2, tr):
        for c in range(1, len(h1) + 1):
            cell = ws1.cell(row=r, column=c)
            cell.border = border
            if c in (6, 9, 10):
                cell.number_format = money
            elif c == 5:
                cell.number_format = pct
    autowidth(ws1, [16, 9, 11, 14, 13, 15, 14, 14, 15, 15])
    ws1.freeze_panes = "A2"
    ws1.auto_filter.ref = f"A1:J{tr - 1}"

    # ======== лист 2: дорожки ========
    ws2 = wb.create_sheet(f"Дорожки {year}")
    h2 = ["Дата", "Дорожка", "Метраж, м", "Плит, шт", "Состав дорожки",
          "Перебетонир.: метраж, м", "% метража перебетонир.", "Цена перебетонир., руб",
          "Среднее переармир., проволок", "Макс переармир., проволок",
          "Лишняя проволока, м", "Цена переармир., руб", "Итого перерасход, руб"]
    ws2.append(h2)
    style_header(ws2, len(h2))
    for r_ in rows_tracks:
        ws2.append([r_["date"], r_["sheet"], round(r_["total_len"], 1), r_["qty"],
                    r_["composition"], round(r_["over_bc_m"], 1), r_["bc_pct"],
                    round(r_["bc_rub"]), round(r_["delta_avg"], 1), int(r_["delta_max"]),
                    round(r_["over_wire_m"]), round(r_["wire_rub"]), round(r_["total_rub"])])
    tr2 = ws2.max_row + 1
    ws2.cell(row=tr2, column=1, value="ИТОГО")
    ws2.cell(row=tr2, column=3, value=round(sum(r_["total_len"] for r_ in rows_tracks), 1))
    ws2.cell(row=tr2, column=4, value=sum(r_["qty"] for r_ in rows_tracks))
    ws2.cell(row=tr2, column=6, value=round(sum(r_["over_bc_m"] for r_ in rows_tracks), 1))
    tm2 = sum(r_["total_len"] for r_ in rows_tracks)
    ws2.cell(row=tr2, column=7, value=sum(r_["over_bc_m"] for r_ in rows_tracks) / tm2 if tm2 else 0.0)
    for col, key in [(8, "bc_rub"), (11, "over_wire_m"), (12, "wire_rub"), (13, "total_rub")]:
        ws2.cell(row=tr2, column=col, value=round(sum(r_[key] for r_ in rows_tracks)))
    for c in range(1, len(h2) + 1):
        cell = ws2.cell(row=tr2, column=c)
        cell.font = Font(bold=True)
        cell.fill = tot_fill
        cell.border = border
        if c in (8, 12, 13):
            cell.number_format = money
        elif c == 7:
            cell.number_format = pct
    for r in range(2, tr2):
        for c in range(1, len(h2) + 1):
            cell = ws2.cell(row=r, column=c)
            cell.border = border
            if c in (8, 12, 13):
                cell.number_format = money
            elif c == 7:
                cell.number_format = pct
    autowidth(ws2, [11, 13, 9, 8, 60, 14, 13, 15, 14, 14, 13, 15, 15])
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:M{tr2 - 1}"

    # ======== лист 3: нормативы ========
    ws3 = wb.create_sheet("Нормативы")
    rows3 = [
        ["НОРМАТИВЫ РАСЧЁТА — обновлено 08.10.2026", "", ""],
        ["", "", ""],
        ["Параметр", "Значение", "Источник"],
        ["Проволока Ø5 Вр-II, кг/м", WIRE_KG_M, "«Расчет новых цен на ПБ 08.10.2026.xls», «Справочные данные»"],
        ["Цена проволоки, руб/кг", WIRE_PRICE_KG, "там же (цены на 08.10.2026, Вр 1400 Северсталь)"],
        ["Цена метра проволоки, руб/м", round(WIRE_PRICE_M, 2), "= 0.17 × 100.33"],
        ["Объём бетона дорожки, м³/м (шир. 1190)", VOL_M_PER_M_12, "там же: 0.165–0.17 м³ в погонном метре"],
        ["Рецепт B30: цемент/песок/щебень, кг/м³", "360/900/1100", "там же, «Состав бетона»"],
        ["Рецепт B40: цемент/песок/щебень, кг/м³", "380/870/1100", "там же"],
        ["Дельта B40−B30, руб/м³", round(DELTA_B40_B30, 1), "= +20 кг цемента × 7.54 − 30 кг песка × 0.66"],
        ["Граница марок: нагрузки 6/8", "М400 до ПБ 66, М500 с ПБ 67", "там же, «Расход номинал»"],
        ["Граница марок: нагрузки 10/12.5/16/21", "всегда М500", "там же"],
        ["Требуемые проволоки по плите", "pb_reinforcement_series", "pb.db (совпадает с файлом 08.10.2026)"],
        ["", "", ""],
        ["МЕТОДИКА", "", ""],
        ["Переармирование плиты", "(макс. проволок дорожки − норма плиты) × метраж", "дорожка льётся по максимуму"],
        ["Цена переармирования", "лишняя проволока, м × 17.06 руб/м", ""],
        ["Перебетонирование плиты", "марка плиты ниже макс. марки дорожки", ""],
        ["Цена перебетонирования", "объём плиты, м³ × 131 руб/м³", ""],
        ["% метража перебетонировано", "метраж ниже-марочных плит / метраж дорожки (номенклатуры)", ""],
        ["", "", ""],
        ["ИЗМЕНЕНИЯ против версии 14.10.2026", "", ""],
        ["Проволока", "80 → 100.33 руб/кг (+25%)", ""],
        ["Дельта марок бетона", "669.6 → 131 руб/м³ (новая рецептура)"],
        ["Граница М400→М500", "ПБ 60 → ПБ 67 (нагр. 6/8); нагр. 10+ — всегда М500", ""],
        ["", "", ""],
        ["Период", f"01.01.{year} – 31.12.{year}", str(kz_dir or KZ_DIR)],
        ["Дорожек в отчёте", len(rows_tracks), ""],
        ["Даты в картах", f"{rows_tracks[0]['date']} – {rows_tracks[-1]['date']}" if rows_tracks else "нет", "по полю «Дата формовки»"],
    ]
    for row in rows3:
        ws3.append(row)
    for r in (1, 15, 21):
        ws3.cell(row=r, column=1).font = Font(bold=True, size=12)
    for c in range(1, 4):
        cell = ws3.cell(row=3, column=c)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.border = border
    autowidth(ws3, [42, 34, 60])

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"[OK] сохранено: {out}")
    print(f"дорожек: {len(rows_tracks)}, номенклатур: {len(items)}")
    print(f"итого перерасход {year}: {sum(r_['total_rub'] for r_ in rows_tracks):,.0f} руб")


def _cli() -> None:
    p = argparse.ArgumentParser(description="Отчёт перерасхода дорожек ПБ по КЗ")
    p.add_argument("--year", type=int, default=2026)
    p.add_argument("--kz-dir", type=Path, default=None)
    p.add_argument("--cache", type=Path, default=None)
    p.add_argument("--out", type=Path, default=None)
    args = p.parse_args()
    cache = args.cache
    if cache is None and args.year != 2026:
        cache = CACHE.with_name(f"kz_tracks_cache_{args.year}.json")
    kz_dir = args.kz_dir
    if kz_dir is None and args.year != 2026:
        year_dir = KZ_DIR / str(args.year)
        if year_dir.is_dir():
            kz_dir = year_dir
    main(args.year, kz_dir=kz_dir, cache=cache, out=args.out)


if __name__ == "__main__":
    _cli()
