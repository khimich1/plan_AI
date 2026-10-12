#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Один прогон xlsx 06.10.2026 → core/pile_cost_norm.json.

Читает формулы листа «Стоимость»: в норму попадает только ячейка расхода,
которую формула умножает на цену. Скрипт не вызывается из API.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections.abc import Iterator, Mapping
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from core.pile_cost import (  # noqa: E402
    Overhead,
    PriceBook,
    Recipe,
    NormLine,
    buyer_price,
)
from core.pile_price_db import GRADE_CODES  # noqa: E402

EXCLUDED_MARK = "С130.35-13А**"
KEPT_STAR_MARK = "С130.35-13А*"
NORM_PATH = PROJECT_ROOT / "core" / "pile_cost_norm.json"
PRICE_TOLERANCE = 0.05

_SKIP_MARKS = frozenset({"", "-", "Марка", "Марка изделия", "Наименование"})
_RECIPE_ARTICLES = frozenset({"cement", "sand", "crushed_20"})
# F — цемент, G — щебень рецепта, I — песок. Колонка H («Щ.5») в lines остаётся.
_RECIPE_CONSUMPTION_COLUMNS = frozenset({"F", "G", "I"})
_COST_BOOKS = (
    "Себестоимость 30х30",
    "Себестоимость 35х35",
    "Себестоимость 40х40",
)
_HEADER_GRADE = (
    ("35", None),
    ("15", "B15"),
    ("20", "B20"),
    ("22", "B22_5"),
    ("25", "B25"),
    ("30", "B30_granite"),
)
# Строка листа «Цены» → код статьи. Подпись проверяется при чтении файла.
_ARTICLE_BY_PRICE_ROW: dict[int, tuple[str, str]] = {
    4: ("cement", "Цемент"),
    5: ("crushed_20", "Щебень 20"),
    6: ("crushed_5", "Щебень "),
    7: ("sand", "Песок"),
    8: ("additive", "Добавка"),
    14: ("a500c_10", "Арматура А500С 10"),
    15: ("a500c_12", "Арматура А500С 12"),
    16: ("a500c_14", "Арматура А500С 14"),
    17: ("a500c_16", "Арматура А500С 16"),
    18: ("a500c_18", "Арматура А500С 18"),
    19: ("a500c_20", "Арматура А500С 20"),
    20: ("a500c_22", "Арматура А500С 22"),
    21: ("a500c_25", "Арматура А500С 25"),
    22: ("a500c_28", "Арматура А500С 28"),
    24: ("a500c_6", "Арматура А500С 6"),
    25: ("a500c_8", "Арматура А500С 8"),
    27: ("ai_10", "Арматура АI 10"),
    28: ("ai_12", "Арматура АI 12"),
    29: ("ai_14", "Арматура АI 14"),
    30: ("ai_16", "Арматура АI 16"),
    31: ("ai_18", "Арматура АI 18"),
    32: ("ai_20", "Арматура АI 20"),
    33: ("ai_22", "Арматура АI 22"),
    35: ("ai_25", "Арматура АI 25"),
    36: ("ai_6", "Арматура АI 6"),
    37: ("ai_8", "Арматура АI 8"),
    38: ("aiii_10", "Арматура АIII 10"),
    49: ("wire_3", "Проволока Вр 1 - 3"),
    50: ("wire_4_8", "Проволока Вр 1 - 4,8"),
    51: ("wire_5", "Проволока Вр 1 - 5,0"),
}
_PRICE_COLUMNS = (
    (3, "B15"),
    (4, "B20"),
    (5, "B22_5"),
    (6, "B25"),
    (7, "B30_granite"),
)
_RECIPES = {
    "B15": Recipe(320, 990, 950),
    "B20": Recipe(340, 950, 950),
    "B22_5": Recipe(355, 940, 950),
    "B25": Recipe(370, 930, 950),
    "B30_granite": Recipe(400, 920, 950),
}
_OVERHEAD = Overhead(general=3186, payroll=3400, forming=40, other=606)

_VLOOKUP = re.compile(
    r"VLOOKUP\([^,]+,\s*(?:'(?P<q>[^']+)'|(?P<sheet>[^!]+))!"
    r"\$(?P<c1>[A-Z]+)\$(?P<r1>\d+):\$(?P<c2>[A-Z]+)\$(?P<r2>\d+),(?P<idx>\d+),"
)
_REMOTE = re.compile(
    r"^(?:'(?P<q>[^']+)'|(?P<sheet>[^!]+))!(?P<cell>\$?[A-Z]{1,3}\$?\d+)$"
)
_LOCAL = re.compile(r"^(?P<cell>\$?[A-Z]{1,3}\$?\d+)$")
_PRICE_REF = re.compile(r"Цены!\$?(?P<col>[A-Z]{1,3})\$?(?P<row>\d+)")
_DEFAULT_ADDITIVE = re.compile(r"^=\$?[A-Z]{1,3}\$?\d+\*0[.,]006$")
_SECTION = re.compile(r"(\d+)\s*СЕЧЕНИЕ", re.IGNORECASE)
_ARITH_TOKEN = re.compile(r"\d+(?:\.\d+)?|[A-Z]{1,3}\d+|[()+\-*/]")


class ExtractError(ValueError):
    """Формула листа «Стоимость» не разобрана."""


def _grade_from_header(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip().lower().replace(" ", "").replace(",", ".")
    for prefix, grade in _HEADER_GRADE:
        if text.startswith(prefix):
            return grade
    return None


def _cell_addr(raw: str) -> str:
    return raw.replace("$", "")


def _remote_target(formula_body: str) -> tuple[str, str] | None:
    match = _REMOTE.match(formula_body)
    if match is None:
        return None
    sheet = (match.group("q") or match.group("sheet") or "").strip()
    return sheet, _cell_addr(match.group("cell"))


def _tokenize_arithmetic(expr: str) -> list[str]:
    compact = expr.replace("$", "").replace(" ", "").replace(",", ".")
    tokens = _ARITH_TOKEN.findall(compact)
    if "".join(tokens) != compact:
        raise ExtractError(f"не разобраны килограммы: {expr}")
    return tokens


def _eval_arithmetic(wb: Workbook, sheet: str, expr: str, depth: int) -> float:
    """Килограммы вида «ячейка / 6 * 5» и суммы чисел. Не движок Excel."""
    tokens = _tokenize_arithmetic(expr)
    pos = 0

    def peek() -> str | None:
        return tokens[pos] if pos < len(tokens) else None

    def eat() -> str:
        nonlocal pos
        token = tokens[pos]
        pos += 1
        return token

    def factor() -> float:
        token = eat()
        if token == "(":
            value = expr_sum()
            if peek() != ")":
                raise ExtractError(f"нет скобки в {expr}")
            eat()
            return value
        if re.fullmatch(r"\d+(?:\.\d+)?", token):
            return float(token)
        if re.fullmatch(r"[A-Z]{1,3}\d+", token):
            value = _follow_number(wb, sheet, token, depth + 1)
            if value is None:
                raise ExtractError(f"пустая ячейка {sheet}!{token}")
            return value
        raise ExtractError(f"не разобраны килограммы: {expr}")

    def term() -> float:
        value = factor()
        while peek() in {"*", "/"}:
            op = eat()
            rhs = factor()
            value = value * rhs if op == "*" else value / rhs
        return value

    def expr_sum() -> float:
        value = term()
        while peek() in {"+", "-"}:
            op = eat()
            rhs = term()
            value = value + rhs if op == "+" else value - rhs
        return value

    result = expr_sum()
    if pos != len(tokens):
        raise ExtractError(f"хвост формулы килограммов: {expr}")
    return result


def _follow_number(wb: Workbook, sheet: str, cell: str, depth: int = 0) -> float | None:
    if depth > 6:
        return None
    value = wb[sheet][cell].value
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if not isinstance(value, str) or not value.startswith("=") or "#REF" in value:
        return None
    body = value[1:]
    remote = _remote_target(body)
    if remote is not None:
        return _follow_number(wb, remote[0], remote[1], depth + 1)
    local = _LOCAL.match(body)
    if local is not None:
        return _follow_number(wb, sheet, _cell_addr(local.group("cell")), depth + 1)
    if _DEFAULT_ADDITIVE.match(value.replace(" ", "")):
        return None
    return _eval_arithmetic(wb, sheet, body, depth + 1)


def _follow_text(wb: Workbook, raw: str) -> str | None:
    if not isinstance(raw, str):
        return None
    if not raw.startswith("="):
        return raw
    if "#REF" in raw:
        return None
    remote = _remote_target(raw[1:])
    if remote is None:
        return None
    inner = wb[remote[0]][remote[1]].value
    if isinstance(inner, str) and not inner.startswith("="):
        return inner
    return None


def _follow_volume(wb: Workbook, sheet: str, raw: object) -> float | None:
    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
        return float(raw)
    if not isinstance(raw, str) or not raw.startswith("=") or "#REF" in raw:
        return None
    remote = _remote_target(raw[1:])
    if remote is not None:
        return _follow_number(wb, remote[0], remote[1])
    local = _LOCAL.match(raw[1:])
    if local is not None:
        return _follow_number(wb, sheet, _cell_addr(local.group("cell")))
    return None


def _load_articles(wb: Workbook) -> dict[int, str]:
    prices = wb["Цены"]
    articles: dict[int, str] = {}
    for row, (article, label) in _ARTICLE_BY_PRICE_ROW.items():
        actual = prices.cell(row, 2).value
        if not isinstance(actual, str) or not actual.startswith(label):
            raise ExtractError(f"лист «Цены» строка {row}: {actual!r}, ждали префикс {label!r}")
        articles[row] = article
    return articles


def _load_unit_prices(wb: Workbook, articles: Mapping[int, str]) -> dict[str, float]:
    prices = wb["Цены"]
    book: dict[str, float] = {}
    for row, article in articles.items():
        value = prices.cell(row, 3).value
        if not isinstance(value, (int, float)) or isinstance(value, bool):
            raise ExtractError(f"нет числа цены в строке {row}")
        book[article] = float(value)
    return book


def _price_catalog(ws: Worksheet) -> tuple[list[str], dict[str, int]]:
    section: int | None = None
    marks: list[str] = []
    sections: dict[str, int] = {}
    for row in range(1, ws.max_row + 1):
        label = ws.cell(row, 2).value
        price = ws.cell(row, 3).value
        if not isinstance(label, str):
            continue
        section_match = _SECTION.search(label)
        if section_match is not None:
            section = int(section_match.group(1))
            continue
        if isinstance(price, str) and price.startswith("="):
            if section is None:
                raise ExtractError(f"марка {label!r} выше первого сечения")
            if label not in sections:
                marks.append(label)
                sections[label] = section
    return marks, sections


def _cost_blocks(wb: Workbook) -> list[tuple[str, int, int, str]]:
    found: list[tuple[str, int, int, str]] = []
    seen: set[tuple[str, int, int, str]] = set()
    for book_name in _COST_BOOKS:
        sheet = wb[book_name]
        for row in sheet.iter_rows():
            for cell in row:
                raw = cell.value
                if not isinstance(raw, str) or "VLOOKUP" not in raw:
                    continue
                match = _VLOOKUP.search(raw)
                if match is None or int(match.group("idx")) != 19:
                    continue
                grade = _grade_from_header(sheet.cell(4, cell.column).value)
                if grade is None:
                    continue
                cost_sheet = (match.group("q") or match.group("sheet") or "").strip()
                block = (cost_sheet, int(match.group("r1")), int(match.group("r2")), grade)
                if block not in seen:
                    seen.add(block)
                    found.append(block)
    return found


def _terms(formula: object) -> Iterator[tuple[str | None, str, int]]:
    if not isinstance(formula, str) or "Цены!" not in formula or "#REF" in formula:
        return
    body = formula[1:] if formula.startswith("=") else formula
    for part in body.split("+"):
        piece = part.strip().strip("()")
        if "Цены!" not in piece:
            continue
        price = _PRICE_REF.search(piece)
        if price is None:
            raise ExtractError(f"не разобрана цена: {formula}")
        if price.group("col").upper() != "C":
            raise ExtractError(f"цена не в колонке C: {formula}")
        left = piece[: price.start()].replace("*", "").strip().strip("()")
        remote = _remote_target(left)
        if remote is not None:
            yield remote[0], remote[1], int(price.group("row"))
            continue
        local = _LOCAL.match(left)
        if local is None:
            raise ExtractError(f"не разобрана ячейка расхода: {formula}")
        yield None, _cell_addr(local.group("cell")), int(price.group("row"))


def _column_letters(cell: str) -> str:
    return "".join(char for char in cell if char.isalpha())


def _consumption_column(
    wb: Workbook,
    cost_sheet: str,
    source_sheet: str | None,
    cell: str,
) -> str | None:
    if source_sheet is not None:
        return _column_letters(cell)
    value = wb[cost_sheet][cell].value
    if isinstance(value, str) and value.startswith("=") and "#REF" not in value:
        remote = _remote_target(value[1:])
        if remote is not None:
            return _column_letters(remote[1])
    return None


def _entered_additive(wb: Workbook, cost_sheet: str, cell: str) -> float | None:
    """Число килограммов, если ячейка добавки не формула «цемент × 0,006»."""
    value = wb[cost_sheet][cell].value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    if not isinstance(value, str) or not value.startswith("="):
        return None
    remote = _remote_target(value[1:])
    if remote is None:
        return None
    source = wb[remote[0]][remote[1]].value
    if isinstance(source, (int, float)) and not isinstance(source, bool):
        return float(source)
    if isinstance(source, str) and _DEFAULT_ADDITIVE.match(source.replace(" ", "")):
        return None
    return None


def _lines_for_row(
    wb: Workbook,
    cost_sheet: str,
    row: int,
    articles: Mapping[int, str],
    sheet: Worksheet,
) -> tuple[list[dict[str, float | str]], float | None, int]:
    lines: list[dict[str, float | str]] = []
    additive: float | None = None
    unknown_prices = 0
    for column in range(1, sheet.max_column + 1):
        for source_sheet, cell, price_row in _terms(sheet.cell(row, column).value):
            article = articles.get(price_row)
            if article is None:
                unknown_prices += 1
                continue
            if article == "additive":
                kg_sheet = cost_sheet if source_sheet is None else source_sheet
                entered = _entered_additive(wb, kg_sheet, cell)
                if entered is not None:
                    additive = entered
                continue
            if article in _RECIPE_ARTICLES or _consumption_column(
                wb, cost_sheet, source_sheet, cell
            ) in _RECIPE_CONSUMPTION_COLUMNS:
                continue
            kg = (
                _follow_number(wb, cost_sheet, cell)
                if source_sheet is None
                else _follow_number(wb, source_sheet, cell)
            )
            if kg is None or kg <= 1e-9:
                continue
            lines.append({"article": article, "kg": kg})
    return lines, additive, unknown_prices


def _row_record(
    wb: Workbook,
    sheet: Worksheet,
    row: int,
    grade: str,
    articles: Mapping[int, str],
) -> tuple[dict[str, object] | None, int]:
    mark = _follow_text(wb, sheet.cell(row, 3).value)
    if mark is None or mark.strip() in _SKIP_MARKS:
        return None, 0
    volume = _follow_volume(wb, sheet.title, sheet.cell(row, 4).value)
    if volume is None:
        return None, 0
    lines, additive, unknown = _lines_for_row(wb, sheet.title, row, articles, sheet)
    record: dict[str, object] = {
        "mark": mark,
        "concrete_grade": grade,
        "volume_m3": volume,
        "lines": lines,
    }
    if additive is not None:
        record["additive_kg"] = additive
    return record, unknown


def extract_norm(wb: Workbook) -> tuple[list[dict[str, object]], dict[str, int]]:
    articles = _load_articles(wb)
    marks, sections = _price_catalog(wb["Прайс"])
    found: dict[tuple[str, str], dict[str, object]] = {}
    stats = {"excluded": 0, "unknown_price_refs": 0, "not_on_price_sheet": 0}
    for cost_sheet, start, end, grade in _cost_blocks(wb):
        sheet = wb[cost_sheet]
        seen_in_block: set[str] = set()
        for row in range(start, end + 1):
            record, unknown = _row_record(wb, sheet, row, grade, articles)
            stats["unknown_price_refs"] += unknown
            if record is None:
                continue
            mark = str(record["mark"])
            if mark in seen_in_block:
                continue
            seen_in_block.add(mark)
            if mark == EXCLUDED_MARK:
                stats["excluded"] += 1
                continue
            if mark not in sections:
                stats["not_on_price_sheet"] += 1
                continue
            key = (mark, grade)
            if key in found:
                continue
            record["section_cm"] = sections[mark]
            found[key] = record

    rows: list[dict[str, object]] = []
    for mark in marks:
        if mark == EXCLUDED_MARK:
            continue
        for grade in GRADE_CODES:
            record = found.get((mark, grade))
            if record is not None:
                rows.append(record)
    stats["price_marks"] = len(marks)
    stats["missing_keys"] = sum(
        1
        for mark in marks
        if mark != EXCLUDED_MARK
        for grade in GRADE_CODES
        if (mark, grade) not in found
    )
    return rows, stats


def _buyer(record: Mapping[str, object], prices: PriceBook) -> float:
    lines = tuple(
        NormLine(str(line["article"]), float(line["kg"]))
        for line in record["lines"]  # type: ignore[union-attr]
    )
    additive = record.get("additive_kg")
    result = buyer_price(
        section_cm=int(record["section_cm"]),  # type: ignore[arg-type]
        volume_m3=float(record["volume_m3"]),  # type: ignore[arg-type]
        concrete_grade=str(record["concrete_grade"]),
        recipe=_RECIPES[str(record["concrete_grade"])],
        lines=lines,
        prices=prices,
        overhead=_OVERHEAD,
        additive_kg=None if additive is None else float(additive),  # type: ignore[arg-type]
    )
    return result.buyer_price


def compare_to_price_sheet(
    rows: list[dict[str, object]],
    wb_values: Workbook,
    articles: Mapping[int, str],
) -> list[tuple[str, str, float, float]]:
    by_key = {(str(row["mark"]), str(row["concrete_grade"])): row for row in rows}
    prices = PriceBook(_load_unit_prices(wb_values, articles))
    sheet = wb_values["Прайс"]
    mismatches: list[tuple[str, str, float, float]] = []
    for row in range(1, sheet.max_row + 1):
        mark = sheet.cell(row, 2).value
        index = sheet.cell(row, 1).value
        if not isinstance(mark, str) or mark == EXCLUDED_MARK:
            continue
        if isinstance(index, bool) or not isinstance(index, (int, float)):
            continue
        if not isinstance(sheet.cell(row, 3).value, (int, float)):
            continue
        for column, grade in _PRICE_COLUMNS:
            expected = sheet.cell(row, column).value
            if not isinstance(expected, (int, float)) or isinstance(expected, bool):
                continue
            record = by_key.get((mark, grade))
            if record is None:
                mismatches.append((mark, grade, float(expected), float("nan")))
                continue
            got = _buyer(record, prices)
            if abs(got - float(expected)) > PRICE_TOLERANCE:
                mismatches.append((mark, grade, float(expected), got))
    return mismatches


def _guard_snapshot(rows: list[dict[str, object]]) -> None:
    marks = {str(row["mark"]) for row in rows}
    if EXCLUDED_MARK in marks or any("**" in mark for mark in marks):
        raise ExtractError("в снимке осталась марка с двумя звёздочками")
    if KEPT_STAR_MARK not in marks:
        raise ExtractError(f"в снимке нет {KEPT_STAR_MARK}")
    sample = next(
        row
        for row in rows
        if row["mark"] == "С30.20-3" and row["concrete_grade"] == "B15"
    )
    kg = {str(line["article"]): float(line["kg"]) for line in sample["lines"]}  # type: ignore[union-attr]
    if abs(float(sample["volume_m3"]) - 0.13) > 1e-9:  # type: ignore[arg-type]
        raise ExtractError(f"объём С30.20-3 / B15 = {sample['volume_m3']}")
    expected_kg = {"wire_5": 3.1, "ai_10": 3.15, "a500c_10": 8.0}
    if set(kg) != set(expected_kg):
        raise ExtractError(f"строки С30.20-3 / B15: {kg}")
    for article, value in expected_kg.items():
        if abs(kg[article] - value) > 1e-9:
            raise ExtractError(f"{article} = {kg[article]}")


def write_norm(path: Path, rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {"rows": rows}
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Собрать снимок нормы цельных свай из xlsx.")
    parser.add_argument("xlsx_path", type=Path, help="Локальный xlsx от 06.10.2026")
    parser.add_argument("--out", type=Path, default=NORM_PATH, help="Куда записать JSON")
    args = parser.parse_args(argv)
    if not args.xlsx_path.is_file():
        print(f"Файл не найден: {args.xlsx_path}", file=sys.stderr)
        return 1

    formulas = load_workbook(args.xlsx_path, data_only=False)
    values = load_workbook(args.xlsx_path, data_only=True)
    try:
        rows, stats = extract_norm(formulas)
        _guard_snapshot(rows)
        articles = _load_articles(formulas)
        mismatches = compare_to_price_sheet(rows, values, articles)
    finally:
        formulas.close()
        values.close()

    marks = {str(row["mark"]) for row in rows}
    print(f"Строк снимка: {len(rows)}")
    print(f"Марок: {len(marks)}")
    print(f"Исключено {EXCLUDED_MARK}: {stats['excluded']}")
    print(f"Нет на листе «Прайс»: {stats['not_on_price_sheet']}")
    print(f"Ключей без нормы: {stats['missing_keys']}")
    print(f"Ссылок на пустую цену: {stats['unknown_price_refs']}")
    print(f"Расхождений с «Прайс» > {PRICE_TOLERANCE}: {len(mismatches)}")
    for mark, grade, expected, got in mismatches[:30]:
        print(f"  {mark} {grade}: лист={expected} снимок={got}")

    if mismatches or stats["missing_keys"]:
        print("Снимок не записан: цены листа «Прайс» или полнота марок не сошлись.", file=sys.stderr)
        return 1
    write_norm(args.out, rows)
    print(f"Записано: {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
