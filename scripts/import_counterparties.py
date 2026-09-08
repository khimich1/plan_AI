#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Импорт справочника контрагентов 1С из xlsx в plita.db.

Формат: один лист, колонки Наименование / Клиент / Код / ИНН / КПП.
Upsert по code_1c. Исчезнувшие из файла записи не удаляются.
"""

from __future__ import annotations

import argparse
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Sequence

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import load_workbook

from app.repositories.counterparties_repository import (
    CounterpartiesRepository,
    UpsertStats,
)
from core.kp_db_common import DEFAULT_DB

_EXPECTED_HEADERS = ("наименование", "клиент", "код", "инн", "кпп")
_FALSE_CLIENT = {"нет", "no", "0", "false"}


@dataclass
class ImportReport:
    added: int = 0
    updated: int = 0
    unchanged: int = 0
    missing_in_file: int = 0
    inn_conflicts: list[tuple[str, list[str]]] = field(default_factory=list)


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def _is_client(value: Any) -> bool:
    return _cell_text(value).casefold() not in _FALSE_CLIENT


def parse_counterparties_xlsx(path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration as exc:
            raise ValueError("Файл выгрузки пуст") from exc
        labels = [_cell_text(cell).casefold() for cell in header]
        if len([x for x in labels if x]) < 5 or labels[:5] != list(_EXPECTED_HEADERS):
            raise ValueError(
                "Ожидается выгрузка из 5 колонок: Наименование, Клиент, Код, ИНН, КПП"
            )
        records: list[dict[str, Any]] = []
        for raw in rows:
            if raw is None:
                continue
            cells = list(raw) + [None] * (5 - len(raw))
            name = _cell_text(cells[0])
            code_1c = _cell_text(cells[2])
            if not name and not code_1c:
                continue
            if not name or not code_1c:
                continue
            records.append(
                {
                    "name": name,
                    "code_1c": code_1c,
                    "is_client": _is_client(cells[1]),
                    "inn": _cell_text(cells[3]) or None,
                    "kpp": _cell_text(cells[4]) or None,
                    "source": "import",
                }
            )
        return records
    finally:
        wb.close()


def _inn_conflicts(
    records: Sequence[dict[str, Any]],
    repo: CounterpartiesRepository,
) -> list[tuple[str, list[str]]]:
    by_inn: dict[str, set[str]] = defaultdict(set)
    for rec in records:
        inn = rec.get("inn")
        if inn:
            by_inn[str(inn)].add(str(rec["code_1c"]))
    for inn, codes in by_inn.items():
        for existing in repo.find_by_inn(inn):
            codes.add(str(existing["code_1c"]))
    return sorted(
        ((inn, sorted(codes)) for inn, codes in by_inn.items() if len(codes) > 1),
        key=lambda item: item[0],
    )


def _same_as_existing(existing: dict[str, Any], rec: dict[str, Any]) -> bool:
    return (
        existing["name"] == rec["name"]
        and (existing["inn"] or None) == (rec.get("inn") or None)
        and (existing["kpp"] or None) == (rec.get("kpp") or None)
        and int(existing["is_client"]) == int(bool(rec.get("is_client", True)))
    )


def _preview_stats(
    records: Iterable[dict[str, Any]],
    repo: CounterpartiesRepository,
) -> UpsertStats:
    added = updated = unchanged = 0
    for rec in records:
        existing = repo.get_by_code_1c(rec["code_1c"])
        if existing is None:
            added += 1
        elif _same_as_existing(existing, rec):
            unchanged += 1
        else:
            updated += 1
    return UpsertStats(added=added, updated=updated, unchanged=unchanged)


def import_counterparties(
    xlsx_path: Path | str,
    db_path: str,
    *,
    dry_run: bool = False,
) -> ImportReport:
    path = Path(xlsx_path)
    if not path.is_file():
        raise FileNotFoundError(f"Файл не найден: {path}")
    records = parse_counterparties_xlsx(path)
    repo = CounterpartiesRepository(db_path=db_path)
    file_codes = {str(rec["code_1c"]) for rec in records}
    missing = len(repo.list_codes() - file_codes)
    conflicts = _inn_conflicts(records, repo)
    if dry_run:
        stats = _preview_stats(records, repo)
    else:
        stats = repo.upsert(records)
    return ImportReport(
        added=stats.added,
        updated=stats.updated,
        unchanged=stats.unchanged,
        missing_in_file=missing,
        inn_conflicts=conflicts,
    )


def render_report(report: ImportReport, *, dry_run: bool = False) -> str:
    prefix = "DRY-RUN " if dry_run else ""
    lines = [
        f"{prefix}added={report.added} updated={report.updated} "
        f"unchanged={report.unchanged} missing_in_file={report.missing_in_file} "
        f"inn_conflicts={len(report.inn_conflicts)}"
    ]
    for inn, codes in report.inn_conflicts:
        lines.append(
            f"WARNING: ИНН {inn} встречается у кодов {', '.join(codes)}"
        )
    return "\n".join(lines)


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Импорт контрагентов 1С (xlsx) в таблицу counterparties."
    )
    parser.add_argument("file", help="Путь к выгрузке xlsx")
    parser.add_argument(
        "--db",
        default=DEFAULT_DB,
        help=f"Путь к plita.db (по умолчанию: {DEFAULT_DB})",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Показать дельту без записи в БД",
    )
    args = parser.parse_args(argv)
    path = Path(args.file)
    if not path.is_file():
        print(f"❌ Файл не найден: {path}")
        return 1
    try:
        report = import_counterparties(path, args.db, dry_run=args.dry_run)
    except ValueError as exc:
        print(f"❌ {exc}")
        return 1
    print(render_report(report, dry_run=args.dry_run))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
