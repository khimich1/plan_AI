"""Подвал КП: физическая масса всех изделий с известным весом, не только плит."""

from __future__ import annotations

import sqlite3
import subprocess
import tempfile

from openpyxl import load_workbook

from core.cargo_delivery_pricing import document_cargo_weight_kg
from core.commercial_offer import generate_commercial_offer_pdf
from core.commercial_offer_xlsx import generate_commercial_offer_xlsx
from core.kp_plate_weight import resolve_kp_line_weight_kg
from core.pile_catalog import PileCatalogEntry, upsert_pile_catalog
from core.product_weight_catalog import ProductWeightRecord, normalize_product_mark, upsert_product_weights


def _catalog(tmp_path) -> str:
    db_path = str(tmp_path / "plita.db")
    sqlite3.connect(db_path).close()
    upsert_pile_catalog(
        db_path,
        [PileCatalogEntry("С60.30", 6.0, 300, 0.55, 1380.0, 14)],
    )
    mark = "ФБС24.4.6"
    upsert_product_weights(
        db_path,
        [
            ProductWeightRecord(
                mark=mark,
                mark_norm=normalize_product_mark(mark),
                product_type="fbs",
                weight_kg=500.0,
                volume_m3=0.2,
                display_name=mark,
            )
        ],
    )
    return db_path


def _pile(qty: int, mark: str = "С60.30-6") -> dict:
    return {
        "product_type": "piles",
        "product_kind": "pile",
        "mark": mark,
        "name": mark,
        "concrete_grade": "B25",
        "qty": qty,
        "unit_price": 100.0,
    }


def test_plate_footer_mass_matches_existing_formula() -> None:
    plate = {"name": "ПБ 60-12-8", "length_m": 6.0, "width_m": 1.2, "qty": 2, "load_class": 800}
    _, line_kg = resolve_kp_line_weight_kg(plate)
    assert document_cargo_weight_kg([plate]) == line_kg


def test_one_pile_uses_catalog_weight_not_zero(tmp_path) -> None:
    db_path = _catalog(tmp_path)
    assert document_cargo_weight_kg([_pile(1)], pile_catalog_db_path=db_path) == 1380.0


def test_full_truck_of_piles_sums_every_piece(tmp_path) -> None:
    """14 штук С60.30 — ровно одна машина, остаток рейса 0 кг, масса заказа 19 320 кг."""
    db_path = _catalog(tmp_path)
    assert document_cargo_weight_kg([_pile(14)], pile_catalog_db_path=db_path) == 19320.0


def test_unknown_pile_and_composite_do_not_add_mass(tmp_path) -> None:
    db_path = _catalog(tmp_path)
    order = [
        _pile(1, mark="С999.99-1"),
        {
            "product_type": "composite_piles",
            "product_kind": "composite_pile",
            "mark": "С60.30-ВС.1",
            "name": "С60.30-ВС.1",
            "qty": 3,
            "unit_price": 100.0,
        },
    ]
    assert document_cargo_weight_kg(order, pile_catalog_db_path=db_path) == 0.0


def test_fbs_catalog_weight_and_mixed_sum(tmp_path) -> None:
    db_path = _catalog(tmp_path)
    plate = {
        "product_type": "plates",
        "name": "ПБ 60-12-8",
        "length_m": 6.0,
        "width_m": 1.2,
        "qty": 1,
        "load_class": 800,
        "unit_price": 100.0,
    }
    fbs = {
        "product_type": "fbs",
        "product_kind": "fbs",
        "mark": "ФБС24.4.6",
        "name": "ФБС24.4.6",
        "qty": 2,
        "unit_price": 100.0,
    }
    _, plate_kg = resolve_kp_line_weight_kg(plate)
    assert document_cargo_weight_kg(
        [plate, _pile(1), fbs],
        pile_catalog_db_path=db_path,
        weight_catalog_db_path=db_path,
    ) == plate_kg + 1380.0 + 1000.0


def _pdf_text(data: bytes) -> str:
    with tempfile.NamedTemporaryFile(suffix=".pdf") as handle:
        handle.write(data)
        handle.flush()
        result = subprocess.run(
            ["pdftotext", "-layout", handle.name, "-"],
            check=True,
            capture_output=True,
        )
    return result.stdout.decode("utf-8", errors="replace")


def test_pdf_and_xlsx_footer_show_pile_mass(tmp_path) -> None:
    db_path = _catalog(tmp_path)
    common = dict(
        order_data=[_pile(1)],
        offer_number="W1",
        offer_date="14.10.2026",
        customer_name="ООО Строитель",
        pile_catalog_db_path=db_path,
        weight_catalog_db_path=db_path,
    )
    pdf_text = _pdf_text(generate_commercial_offer_pdf(**common).getvalue())
    assert "1 380,000" in pdf_text

    book = load_workbook(generate_commercial_offer_xlsx(**common), data_only=False)
    sheet = book.active
    footer = " ".join(
        str(cell.value)
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and "общим весом" in cell.value
    )
    assert "1,380.000" in footer
