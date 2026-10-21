"""Разовый импорт карточек из трёх колонок (CDC-006). Файл собирается в tmp_path."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook

from app.services.supply_contract_service import (
    SupplyContractService,
    SupplyContractValidationError,
)
from core import kp_db_schema
from core.kp_db_common import _connect

_SHEET = "ИНН без сомнений"
_HEADERS = ("Номер договора", "Имя контрагента", "ИНН")
_SHARED_INN = "440117229007"
_ADMIN = {"id": 1, "role": "admin", "username": "admin"}
_MANAGER = {"id": 7, "role": "manager", "username": "manager"}

# Восемь боевых исключений плюс 176 обычных номеров. 0001/10/26 в файле нет.
_SPECIAL = (
    ("0793/02/25", "ИП Секлюцкий", _SHARED_INN, "2025-02-01"),
    ("0794/02/25", "ИП Секлюцкий второй", _SHARED_INN, "2025-02-01"),
    ("0933/03/26", "ИП Секлюцкий третий", _SHARED_INN, "2026-03-01"),
    ("0786/1/25", "ИП Ручной", "5001000001", "2025-01-01"),
    ("0844-06-25", "ООО Дефис", "5001000002", "2025-06-01"),
    ("966/06/26КП", "ООО С хвостом КП", "5001000003", "2026-06-01"),
    ("ТБ-3/08/26Б", "ООО Тендер", "5001000004", "2026-08-01"),
    ("1033/09/26", "ООО Максимум", "5001000005", "2026-09-01"),
)


def _service(tmp_path: Path) -> SupplyContractService:
    db_path = str(tmp_path / "cards.db")
    kp_db_schema._schema_ready.clear()
    kp_db_schema.ensure_schema(db_path)
    return SupplyContractService(db_path=db_path, today=date(2026, 10, 20))


def _count(db_path: str, table: str) -> int:
    with _connect(db_path) as conn:
        row = conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()
    return int(row[0])


def _fixture_rows() -> list[tuple[str, str, str]]:
    rows = [(number, name, inn) for number, name, inn, _date in _SPECIAL]
    for index in range(1, 177):
        rows.append((f"{index:04d}/03/24", f"Контрагент {index:04d}", f"77{index:010d}"))
    return rows


def _write_xlsx(
    tmp_path: Path,
    rows: list[tuple[str, str, str]],
    *,
    filename: str = "cards.xlsx",
    sheet_name: str = _SHEET,
    headers: tuple[str, str, str] = _HEADERS,
    active_title: str | None = None,
) -> bytes:
    path = tmp_path / filename
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    if active_title is not None:
        sheet.title = active_title
        sheet.append(("не то", "совсем", "иное"))
        sheet.append(("мусор", "мусор", "мусор"))
        sheet = book.create_sheet(sheet_name)
    else:
        sheet.title = sheet_name
    sheet.append(list(headers))
    for number, name, inn in rows:
        sheet.append([number, name, inn])
    book.save(path)
    return path.read_bytes()


def _seed_counterparty(db_path: str) -> None:
    with _connect(db_path) as conn:
        conn.execute(
            """
            INSERT INTO counterparties (code_1c, name, name_normalized, inn)
            VALUES ('00-0001', 'ООО Ромашка', 'ооо ромашка', '7701000001')
            """
        )


def _seed_kept_contract(db_path: str) -> None:
    with _connect(db_path) as conn:
        conn.execute(
            """
            INSERT INTO supply_contract (
                counterparty_id, imported_name, full_name, short_name, inn, number,
                contract_date, manager_name, status, legal_form,
                signatory_name, signatory_verb, authority_basis,
                ogrn, legal_address, email, bank_name, account, corr_account, bik,
                created_at, created_by_user_id
            ) VALUES (
                NULL, NULL, 'ООО Ромашка', 'Ромашка', '7701000001', '0001/10/26',
                '2026-10-15', 'Иван', 'оригинал в бухгалтерии', 'ooo',
                'Иванов Иван', 'действующего', 'устав',
                '1027700000001', 'Москва', 'a@b.ru', 'Банк',
                '40702810000000000001', '30101810000000000001', '044525225',
                '2026-10-01T09:00:00', 3
            )
            """
        )


def test_fixture_creates_184_incomplete_cards_and_leaves_counterparties(
    tmp_path: Path,
) -> None:
    service = _service(tmp_path)
    _seed_counterparty(service.repository.db_path)
    parties_before = _count(service.repository.db_path, "counterparties")
    rows = _fixture_rows()
    assert len(rows) == 184
    assert "0001/10/26" not in {number for number, _name, _inn in rows}

    result = service.import_contract_cards(
        _write_xlsx(tmp_path, rows),
        user=_ADMIN,
    )

    assert result == {"created": 184, "skipped": 0}
    assert _count(service.repository.db_path, "supply_contract") == 184
    assert _count(service.repository.db_path, "counterparties") == parties_before
    assert _count(service.repository.db_path, "supply_contract_event") == 0
    for number, name, inn, contract_date in _SPECIAL:
        stored = service.repository.get_by_number(number)
        assert stored is not None
        assert stored["number"] == number
        assert stored["imported_name"] == name
        assert stored["inn"] == inn
        assert stored["contract_date"] == contract_date
        assert stored["status"] == "нет"
        assert stored["counterparty_id"] is None
        assert stored["full_name"] is None
        assert stored["short_name"] is None
        assert stored["manager_name"] is None
        assert stored["legal_form"] is None
        assert stored["created_by_user_id"] == 1
    shared = [
        row
        for row in service.repository.list_for_search()
        if row["inn"] == _SHARED_INN
    ]
    assert {row["number"] for row in shared} == {
        "0793/02/25",
        "0794/02/25",
        "0933/03/26",
    }
    filler = service.repository.get_by_number("0001/03/24")
    assert filler is not None
    assert filler["imported_name"] == "Контрагент 0001"
    assert filler["number"] != "0001/10/26"


def test_existing_0001_keeps_name_inn_requisites_date_and_status(tmp_path: Path) -> None:
    service = _service(tmp_path)
    _seed_counterparty(service.repository.db_path)
    _seed_kept_contract(service.repository.db_path)
    before = service.repository.get_by_number("0001/10/26")
    parties_before = _count(service.repository.db_path, "counterparties")

    result = service.import_contract_cards(
        _write_xlsx(tmp_path, _fixture_rows(), filename="registry.xlsx"),
        user=_ADMIN,
    )

    after = service.repository.get_by_number("0001/10/26")
    assert result == {"created": 184, "skipped": 0}
    assert after == before
    assert after["full_name"] == "ООО Ромашка"
    assert after["inn"] == "7701000001"
    assert after["legal_form"] == "ooo"
    assert after["contract_date"] == "2026-10-15"
    assert after["status"] == "оригинал в бухгалтерии"
    assert after["counterparty_id"] is None
    assert _count(service.repository.db_path, "counterparties") == parties_before
    assert _count(service.repository.db_path, "supply_contract") == 185


def test_second_import_does_not_overwrite_a_manually_corrected_name(
    tmp_path: Path,
) -> None:
    service = _service(tmp_path)
    payload = _write_xlsx(tmp_path, _fixture_rows(), filename="twice.xlsx")
    service.import_contract_cards(payload, user=_ADMIN)
    stored = service.repository.get_by_number("0793/02/25")
    assert stored is not None
    service.update_card(
        int(stored["id"]),
        user=_ADMIN,
        name="Исправлено вручную",
    )

    again = service.import_contract_cards(payload, user=_ADMIN)

    kept = service.repository.get_by_number("0793/02/25")
    assert again == {"created": 0, "skipped": 184}
    assert kept is not None
    assert kept["imported_name"] == "Исправлено вручную"
    assert kept["full_name"] is None
    assert kept["inn"] == _SHARED_INN
    assert kept["contract_date"] == "2025-02-01"
    assert kept["status"] == "нет"
    assert _count(service.repository.db_path, "supply_contract") == 184


def test_create_without_number_after_import_uses_next_seq(tmp_path: Path) -> None:
    service = _service(tmp_path)
    service.import_contract_cards(
        _write_xlsx(tmp_path, _fixture_rows(), filename="seq.xlsx"),
        user=_ADMIN,
    )

    created = service.create_card(name="Новый после импорта", user=_ADMIN)

    assert created["number"] == "1034/10/26"
    assert created["contract_date"] == "2026-10-20"
    assert created["imported_name"] == "Новый после импорта"
    assert created["status"] == "нет"
    assert created["counterparty_id"] is None


def test_existing_number_is_skipped_and_its_name_stays(tmp_path: Path) -> None:
    service = _service(tmp_path)
    service.repository.insert_card(
        imported_name="Уже исправлено",
        number="0793/02/25",
        user_id=3,
        inn="111111111111",
        contract_date=date(2024, 5, 5),
    )

    result = service.import_contract_cards(
        _write_xlsx(
            tmp_path,
            [
                ("0793 / 02 / 25", "Из файла", _SHARED_INN),
                ("0786/1/25", "Новый", "5001000001"),
            ],
            filename="skip.xlsx",
        ),
        user=_ADMIN,
    )

    kept = service.repository.get_by_number("0793/02/25")
    fresh = service.repository.get_by_number("0786/1/25")
    assert result == {"created": 1, "skipped": 1}
    assert kept is not None
    assert kept["imported_name"] == "Уже исправлено"
    assert kept["inn"] == "111111111111"
    assert kept["contract_date"] == "2024-05-05"
    assert kept["created_by_user_id"] == 3
    assert fresh is not None
    assert fresh["imported_name"] == "Новый"
    assert fresh["full_name"] is None
    assert fresh["counterparty_id"] is None
    assert _count(service.repository.db_path, "supply_contract") == 2


def test_duplicate_number_inside_the_file_is_skipped_not_rolled_back(
    tmp_path: Path,
) -> None:
    service = _service(tmp_path)

    result = service.import_contract_cards(
        _write_xlsx(
            tmp_path,
            [
                ("0793/02/25", "Первое имя", _SHARED_INN),
                ("0793/02/25", "Второе имя", "999999999999"),
                ("0786/1/25", "Другой", "5001000001"),
            ],
            filename="dup.xlsx",
        ),
        user=_ADMIN,
    )

    stored = service.repository.get_by_number("0793/02/25")
    assert result == {"created": 2, "skipped": 1}
    assert stored is not None
    assert stored["imported_name"] == "Первое имя"
    assert stored["inn"] == _SHARED_INN
    assert service.repository.get_by_number("0786/1/25") is not None
    assert _count(service.repository.db_path, "supply_contract") == 2


def test_empty_name_rejects_the_whole_file(tmp_path: Path) -> None:
    service = _service(tmp_path)
    _seed_counterparty(service.repository.db_path)

    with pytest.raises(SupplyContractValidationError, match="пустое имя контрагента"):
        service.import_contract_cards(
            _write_xlsx(
                tmp_path,
                [
                    ("0793/02/25", "ИП Секлюцкий", _SHARED_INN),
                    ("0786/1/25", "   ", "5001000001"),
                ],
                filename="empty-name.xlsx",
            ),
            user=_ADMIN,
        )

    assert _count(service.repository.db_path, "supply_contract") == 0
    assert _count(service.repository.db_path, "counterparties") == 1
    assert service.repository.get_by_number("0793/02/25") is None


def test_number_that_fails_normalization_rejects_the_whole_file(tmp_path: Path) -> None:
    service = _service(tmp_path)

    with pytest.raises(SupplyContractValidationError, match="хотя бы одна цифра"):
        service.import_contract_cards(
            _write_xlsx(
                tmp_path,
                [
                    ("0793/02/25", "ИП Секлюцкий", _SHARED_INN),
                    ("без цифр", "ООО Пусто", "5001000001"),
                ],
                filename="bad-number.xlsx",
            ),
            user=_ADMIN,
        )

    assert service.repository.get_by_number("0793/02/25") is None
    assert _count(service.repository.db_path, "supply_contract") == 0


def test_manager_cannot_import_cards(tmp_path: Path) -> None:
    service = _service(tmp_path)

    with pytest.raises(SupplyContractValidationError, match="только администратору"):
        service.import_contract_cards(
            _write_xlsx(
                tmp_path,
                [("0786/1/25", "ИП Ручной", "5001000001")],
                filename="manager.xlsx",
            ),
            user=_MANAGER,
        )

    assert _count(service.repository.db_path, "supply_contract") == 0


def test_named_sheet_is_used_when_the_active_sheet_is_something_else(
    tmp_path: Path,
) -> None:
    service = _service(tmp_path)

    result = service.import_contract_cards(
        _write_xlsx(
            tmp_path,
            [("0786/1/25", "ИП Ручной", "5001000001")],
            filename="named.xlsx",
            active_title="Спорные ИНН",
        ),
        user=_ADMIN,
    )

    stored = service.repository.get_by_number("0786/1/25")
    assert result == {"created": 1, "skipped": 0}
    assert stored is not None
    assert stored["imported_name"] == "ИП Ручной"
    assert stored["contract_date"] == "2025-01-01"


def test_active_sheet_is_used_when_the_named_sheet_is_absent(tmp_path: Path) -> None:
    service = _service(tmp_path)

    result = service.import_contract_cards(
        _write_xlsx(
            tmp_path,
            [("0844-06-25", "ООО Дефис", "")],
            filename="active.xlsx",
            sheet_name="Лист1",
            headers=("НОМЕР ДОГОВОРА", "ИМЯ КОНТРАГЕНТА", "инн"),
        ),
        user=_ADMIN,
    )

    stored = service.repository.get_by_number("0844-06-25")
    assert result == {"created": 1, "skipped": 0}
    assert stored is not None
    assert stored["imported_name"] == "ООО Дефис"
    assert stored["inn"] is None
    assert stored["contract_date"] == "2025-06-01"
    assert stored["full_name"] is None


def test_missing_columns_do_not_insert_rows(tmp_path: Path) -> None:
    service = _service(tmp_path)
    path = tmp_path / "lawyer.xlsx"
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = _SHEET
    sheet.append(["Номер договора", "Дата", "Контрагент"])
    sheet.append(["1027/09/26", "25.09.2026", "ООО Ромашка"])
    book.save(path)

    with pytest.raises(SupplyContractValidationError, match="имя контрагента"):
        service.import_contract_cards(path.read_bytes(), user=_ADMIN)

    assert _count(service.repository.db_path, "supply_contract") == 0


def test_blank_rows_are_ignored(tmp_path: Path) -> None:
    service = _service(tmp_path)
    path = tmp_path / "blanks.xlsx"
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = _SHEET
    sheet.append(list(_HEADERS))
    sheet.append(["0786/1/25", "  ИП Ручной  ", "5001000001"])
    sheet.append([None, None, None])
    sheet.append(["", "  ", None])
    book.save(path)

    result = service.import_contract_cards(path.read_bytes(), user=_ADMIN)

    stored = service.repository.get_by_number("0786/1/25")
    assert result == {"created": 1, "skipped": 0}
    assert stored is not None
    assert stored["imported_name"] == "ИП Ручной"
    assert stored["number"] == "0786/1/25"
