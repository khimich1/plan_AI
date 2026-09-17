"""GPS-008: import filled unit prices from the GUID/price queue Excel."""

from __future__ import annotations

import sqlite3
from pathlib import Path

from openpyxl import Workbook

from core.bridge_pile_price_db import get_bridge_pile_price
from core.fbs_price_db import get_fbs_price
from core.march_price_db import get_march_price
from core.pile_price_db import GRADE_CODES as PILE_GRADE_CODES
from core.pile_price_db import get_pile_price
from core.price_import_queue import (
    PRICE_INPUT_HEADER,
    PRICE_SECTION_TITLE,
    TASKS_SHEET,
    import_queue_prices,
    parse_queue_price_rows,
)
from core.step_price_db import get_step_price

PRICE_HEADERS = [
    "Группа",
    "Наименование 1С",
    "GUID 1С",
    "Цена за шт (ввод)",
    "как у ближайшего соседа",
    "Сосед (марка)",
    "Цена 1С (Продажная)",
    "Ед.",
    "Примечание",
]

CREATE_HEADERS = [
    "Группа",
    "Марка у нас",
    "Завести в 1С как",
    "Действие",
    "GUID «у» уже есть",
    "Наименование «у»",
    "Примечание",
]


def _write_queue(path: Path, price_rows: list[list[object]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = TASKS_SHEET
    ws["A1"] = "🏭 завести в 1С"
    for col, header in enumerate(CREATE_HEADERS, start=1):
        ws.cell(row=2, column=col, value=header)
    ws.append(["Сваи", "С999.99-1", "Сваи С 999.99-1", "завести", None, None, "не импортировать"])
    gap = ws.max_row + 2
    ws.cell(row=gap, column=1, value=PRICE_SECTION_TITLE)
    hdr_row = gap + 1
    for col, header in enumerate(PRICE_HEADERS, start=1):
        ws.cell(row=hdr_row, column=col, value=header)
    for row in price_rows:
        ws.append(row)
    wb.save(path)


def _seed_pile(db_path: Path, mark: str, prices: dict[str, float]) -> None:
    from core.pile_price_db import init_pile_prices_schema

    init_pile_prices_schema(str(db_path))
    conn = sqlite3.connect(db_path)
    try:
        conn.executemany(
            "INSERT INTO pile_prices (mark, concrete_grade, price) VALUES (?, ?, ?)",
            [(mark, grade, price) for grade, price in prices.items()],
        )
        conn.commit()
    finally:
        conn.close()


def test_valid_pile_price_updates_all_existing_grades(tmp_path: Path) -> None:
    db_path = tmp_path / "pb.db"
    xlsx_path = tmp_path / "queue.xlsx"
    old_prices = {
        "B15": 100.0,
        "B20": 110.0,
        "B22_5": 120.0,
        "B25": 130.0,
        "B30_granite": 140.0,
    }
    _seed_pile(db_path, "С60.30-6", old_prices)
    _write_queue(
        xlsx_path,
        [["Сваи", "Сваи С 60.30-6", "guid-pile", 5555.5, 4000, "С60.30-10", None, "шт", ""]],
    )

    report = import_queue_prices(xlsx_path, db_path)

    assert report.skipped_empty == 0
    assert report.errors == ()
    assert len(report.written) == 1
    assert report.written[0].created is False
    assert set(report.written[0].grades) == set(PILE_GRADE_CODES)
    for grade in PILE_GRADE_CODES:
        assert get_pile_price("С60.30-6", grade, str(db_path)) == 5555.5


def test_price_le_zero_is_error_not_written(tmp_path: Path) -> None:
    db_path = tmp_path / "pb.db"
    xlsx_path = tmp_path / "queue.xlsx"
    _seed_pile(db_path, "С60.30-6", {"B25": 130.0})
    _write_queue(
        xlsx_path,
        [
            ["Сваи", "Сваи С 60.30-6", "g1", 0, None, None, None, "шт", ""],
            ["Сваи", "Сваи С 60.30-6", "g2", -10, None, None, None, "шт", ""],
        ],
    )

    report = import_queue_prices(xlsx_path, db_path)

    assert len(report.errors) == 2
    assert all("цена должна быть > 0" in err.reason for err in report.errors)
    assert report.written == ()
    assert get_pile_price("С60.30-6", "B25", str(db_path)) == 130.0


def test_unknown_mark_skipped_with_error(tmp_path: Path) -> None:
    db_path = tmp_path / "pb.db"
    xlsx_path = tmp_path / "queue.xlsx"
    _seed_pile(db_path, "С60.30-6", {"B25": 130.0})
    _write_queue(
        xlsx_path,
        [["Сваи", "непонятная номенклатура 1С", "g3", 1000, None, None, None, "шт", ""]],
    )

    report = import_queue_prices(xlsx_path, db_path)

    assert report.written == ()
    assert len(report.errors) == 1
    assert report.errors[0].reason == "неизвестная марка"
    assert get_pile_price("С60.30-6", "B25", str(db_path)) == 130.0


def test_empty_price_column_skipped_not_error(tmp_path: Path) -> None:
    db_path = tmp_path / "pb.db"
    xlsx_path = tmp_path / "queue.xlsx"
    _seed_pile(db_path, "С60.30-6", {"B25": 130.0})
    _write_queue(
        xlsx_path,
        [
            ["Сваи", "Сваи С 60.30-6", "g4", None, 9999, "С60.30-10", None, "шт", ""],
            ["Сваи", "Сваи С 50.30-4", "g5", "  ", 8888, "С50.30-10", None, "шт", ""],
        ],
    )

    report = import_queue_prices(xlsx_path, db_path)

    assert report.skipped_empty == 2
    assert report.errors == ()
    assert report.written == ()
    assert get_pile_price("С60.30-6", "B25", str(db_path)) == 130.0
    assert get_pile_price("С50.30-4", "B25", str(db_path)) is None


def test_fbs_bridge_step_march_one_row_write(tmp_path: Path) -> None:
    db_path = tmp_path / "pb.db"
    xlsx_path = tmp_path / "queue.xlsx"
    from core.bridge_pile_price_db import init_bridge_pile_prices_schema
    from core.fbs_price_db import init_fbs_prices_schema
    from core.march_price_db import init_march_prices_schema
    from core.step_price_db import init_step_prices_schema

    init_fbs_prices_schema(str(db_path))
    init_bridge_pile_prices_schema(str(db_path))
    init_march_prices_schema(str(db_path))
    init_step_prices_schema(str(db_path))
    conn = sqlite3.connect(db_path)
    try:
        conn.execute(
            "INSERT INTO fbs_prices (mark, concrete_grade, price) VALUES (?, ?, ?)",
            ("ФБС 9.3.6-Т", "B25", 1700.0),
        )
        conn.execute(
            "INSERT INTO bridge_pile_prices (mark, concrete_grade, price) "
            "VALUES (?, ?, ?)",
            ("C8-35T1", "B25", 35000.0),
        )
        conn.execute(
            "INSERT INTO march_prices (mark, concrete_grade, price) VALUES (?, ?, ?)",
            ("1ЛМ 30-11-15-4", "B25", 14000.0),
        )
        conn.execute(
            "INSERT INTO step_prices (mark, price) VALUES (?, ?)",
            ("ЛС11", 1400.0),
        )
        conn.commit()
    finally:
        conn.close()

    _write_queue(
        xlsx_path,
        [
            ["Блоки ФБС", "Блоки ФБС 9.3.6-Т", "gf", 2100.0, None, None, None, "шт", ""],
            ["Сваи мостовые", "Сваи мостовые С 8.35-Т1", "gb", 41000.0, None, None, None, "шт", ""],
            ["Марши", "Лестничные марши 1ЛМ 30-11-15-4", "gm", 15500.0, None, None, None, "шт", ""],
            ["Ступени", "Лестничные ступени ЛС11", "gs", 1600.0, None, None, None, "шт", ""],
        ],
    )

    report = import_queue_prices(xlsx_path, db_path)

    assert report.errors == ()
    assert len(report.written) == 4
    assert get_fbs_price("ФБС 9.3.6-Т", "B25", str(db_path)) == 2100.0
    assert get_bridge_pile_price("C8-35T1", "B25", str(db_path)) == 41000.0
    assert get_march_price("1ЛМ 30-11-15-4", "B25", str(db_path)) == 15500.0
    assert get_step_price("ЛС11", str(db_path)) == 1600.0


def test_new_1c_only_pile_inserts_all_known_grades(tmp_path: Path) -> None:
    db_path = tmp_path / "pb.db"
    xlsx_path = tmp_path / "queue.xlsx"
    _write_queue(
        xlsx_path,
        [["Сваи", "Сваи С 50.30-4", "g-new", 3210.0, None, None, None, "шт", ""]],
    )

    report = import_queue_prices(xlsx_path, db_path)

    assert report.errors == ()
    assert report.written[0].created is True
    assert tuple(report.written[0].grades) == PILE_GRADE_CODES
    for grade in PILE_GRADE_CODES:
        assert get_pile_price("С50.30-4", grade, str(db_path)) == 3210.0


def test_create_section_and_neighbor_hint_are_ignored(tmp_path: Path) -> None:
    db_path = tmp_path / "pb.db"
    xlsx_path = tmp_path / "queue.xlsx"
    _write_queue(
        xlsx_path,
        [["Сваи", "Сваи С 40.30-5", "g6", None, 7777.0, "С40.30-1", None, "шт", ""]],
    )
    rows = parse_queue_price_rows(xlsx_path)
    assert all(row.group != "Сваи" or "999.99" not in row.name_1c for row in rows)
    assert PRICE_INPUT_HEADER
    assert any(row.price_raw is None for row in rows)

    report = import_queue_prices(xlsx_path, db_path)
    assert report.written == ()
    assert get_pile_price("С40.30-5", "B25", str(db_path)) is None
    assert get_pile_price("С999.99-1", "B25", str(db_path)) is None


def test_pile_1c_only_a500_and_dotted_load(tmp_path: Path) -> None:
    db_path = tmp_path / "pb.db"
    xlsx_path = tmp_path / "queue.xlsx"
    _write_queue(
        xlsx_path,
        [
            ["Сваи", "Сваи С 30.15-А500", "ga", 1111.0, None, None, None, "шт", ""],
            [
                "Сваи",
                "Свая железобетонная С120.35.12 350х350, 12м",
                "gd",
                2222.0,
                None,
                None,
                None,
                "шт",
                "",
            ],
        ],
    )

    report = import_queue_prices(xlsx_path, db_path)

    assert report.errors == ()
    marks = {item.mark for item in report.written}
    assert "С30.15-А500" in marks
    assert "С120.35-12" in marks
    assert get_pile_price("С30.15-А500", "B25", str(db_path)) == 1111.0
    assert get_pile_price("С120.35-12", "B25", str(db_path)) == 2222.0


def test_unmanaged_group_with_price_is_error(tmp_path: Path) -> None:
    db_path = tmp_path / "pb.db"
    xlsx_path = tmp_path / "queue.xlsx"
    _write_queue(
        xlsx_path,
        [["Плиты", "Плиты ПБ 60-12-8п", "gp", 5000, None, None, None, "шт", ""]],
    )

    report = import_queue_prices(xlsx_path, db_path)

    assert report.written == ()
    assert len(report.errors) == 1
    assert report.errors[0].reason == "неизвестная группа"
