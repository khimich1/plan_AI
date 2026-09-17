"""GPS-015: парсер «Универсального отчёта» 1С (.xlsx)."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from core.universal_report_1c_parser import (
    UniversalReport1CFormatError,
    parse_universal_report_xlsx,
)

SAMPLE = Path("/home/roman/Загрузки/Пример Универсального отчета с УИД номенклатуры.xlsx")
GUID_RAW = "60d7ac63-68b6-11f1-9003-04d9f5387845"
GUID_BRACED = "{60d7ac63-68b6-11f1-9003-04d9f5387845}"


def _write_report(
    path: Path,
    *,
    with_filter: bool,
    header: tuple[str, ...] | None = None,
    data_rows: list[tuple] | None = None,
    merge_header: bool = True,
) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Лист_1"
    ws["A2"] = "Параметры:"
    ws["C2"] = "Тип объекта: Справочник"
    ws["C3"] = "Имя объекта: Номенклатура"
    if with_filter:
        ws["A6"] = "Отбор:"
        ws["C6"] = 'Номенклатура Равно "Сваи С 30.30-3"'
    header_row = 8
    cols = header or (
        "Номенклатура",
        None,
        None,
        "УИД",
        "Код",
        "Вес (числитель)",
        "Вес (знаменатель)",
        "Объем (числитель)",
    )
    for idx, value in enumerate(cols, start=1):
        if value is not None:
            ws.cell(header_row, idx, value)
    if merge_header:
        ws.merge_cells(start_row=header_row, start_column=1, end_row=header_row, end_column=3)
        ws.merge_cells(start_row=header_row, start_column=6, end_row=header_row, end_column=7)
    rows = data_rows if data_rows is not None else [
        ("Сваи С 30.30-3", None, None, GUID_RAW, "00-00077803", 1160, None, 0.464),
        ("Итого", None, None, None, None, 1160, None, 0.464),
    ]
    for offset, row in enumerate(rows, start=1):
        excel_row = header_row + offset
        for idx, value in enumerate(row, start=1):
            if value is not None:
                ws.cell(excel_row, idx, value)
        if merge_header:
            ws.merge_cells(
                start_row=excel_row, start_column=1, end_row=excel_row, end_column=3
            )
    wb.save(path)
    return path


def test_parses_full_report_without_filter(tmp_path: Path) -> None:
    path = _write_report(tmp_path / "full.xlsx", with_filter=False)
    result = parse_universal_report_xlsx(path)

    assert result.partial is False
    assert len(result.rows) == 1
    row = result.rows[0]
    assert row.guid == GUID_RAW
    assert row.code == "00-00077803"
    assert row.name == "Сваи С 30.30-3"
    assert row.qty is None
    assert row.weight == pytest.approx(1160)
    assert row.volume == pytest.approx(0.464)
    assert row.price is None


def test_filter_row_marks_partial(tmp_path: Path) -> None:
    path = _write_report(tmp_path / "partial.xlsx", with_filter=True)
    result = parse_universal_report_xlsx(path)
    assert result.partial is True
    assert len(result.rows) == 1


def test_skips_itogo_and_keeps_leading_zeros(tmp_path: Path) -> None:
    path = _write_report(
        tmp_path / "zeros.xlsx",
        with_filter=True,
        data_rows=[
            ("Блоки ФБС 9.3.6-Т", None, None, GUID_RAW, "00-00000007", 80, None, 0.03),
            ("Итого", None, None, None, None, 80, None, 0.03),
        ],
    )
    result = parse_universal_report_xlsx(path)
    assert [row.code for row in result.rows] == ["00-00000007"]
    assert [row.name for row in result.rows] == ["Блоки ФБС 9.3.6-Т"]


def test_guid_strips_braces_and_normalizes_case(tmp_path: Path) -> None:
    path = _write_report(
        tmp_path / "braces.xlsx",
        with_filter=False,
        data_rows=[
            ("Сваи С 40.30-6", None, None, GUID_BRACED.upper(), "00-1", 1, None, 0.1),
        ],
    )
    result = parse_universal_report_xlsx(path)
    assert result.rows[0].guid == GUID_RAW


def test_missing_header_raises_russian_error(tmp_path: Path) -> None:
    path = tmp_path / "bad.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Товар"
    ws["B1"] = "Цена"
    wb.save(path)

    with pytest.raises(UniversalReport1CFormatError, match="не универсальный отчёт 1С"):
        parse_universal_report_xlsx(path)


@pytest.mark.skipif(not SAMPLE.is_file(), reason="нет образца универсального отчёта")
def test_parses_real_sample() -> None:
    result = parse_universal_report_xlsx(SAMPLE)
    assert result.partial is True
    assert len(result.rows) == 1
    row = result.rows[0]
    assert row.name == "Лестничные площадки 2ЛП25.12-4-к"
    assert row.guid == GUID_RAW
    assert row.code == "00-00077803"
    assert row.weight == pytest.approx(1160)
    assert row.volume == pytest.approx(0.464)
    assert row.qty is None
