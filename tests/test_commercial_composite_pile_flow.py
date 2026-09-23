"""CP-601: composite piles commercial draft flow."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
import sqlite3
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.core.settings import get_settings
from app.main import create_app
from app.security.session import create_session_token
from core.composite_pile_kit import import_composite_pile_kit_from_xlsx
from core.composite_pile_price_db import import_composite_pile_prices_from_xlsx
from core.kp_db_schema import init_schema
from tests.helpers import kp_db_fixtures as fx
from tests.helpers.auth_fixtures import patch_auth_users
from tests.helpers.csrf import CsrfAwareTestClient

KIT = Path(__file__).parent / "fixtures" / "composite_pile_kit_sample.xlsx"
PRICE = Path(__file__).parent / "fixtures" / "composite_pile_price_sample.xlsx"


def _setup_env(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> Path:
    monkeypatch.setenv("APP_SECRET_KEY", "test-secret-key-for-pytest-must-be-32-chars-min")
    monkeypatch.setenv("DRAFTS_DIR", str(tmp_path / "drafts"))
    monkeypatch.setenv("OUTPUTS_DIR", str(tmp_path / "outputs"))
    (tmp_path / "drafts").mkdir(exist_ok=True)
    (tmp_path / "outputs").mkdir(exist_ok=True)

    pb_db = tmp_path / "pb.db"
    import_composite_pile_kit_from_xlsx(str(KIT), str(pb_db))
    import_composite_pile_prices_from_xlsx(str(PRICE), str(pb_db))

    import core.commercial_offer as commercial_offer
    import core.commercial_offer_xlsx as commercial_offer_xlsx
    import app.services.product_draft_config as product_draft_config

    monkeypatch.setattr(commercial_offer, "DB_PATH", str(pb_db))
    monkeypatch.setattr(commercial_offer_xlsx, "DB_PATH", str(pb_db))
    monkeypatch.setattr(product_draft_config, "DB_PATH", str(pb_db))
    get_settings.cache_clear()
    return pb_db


def _mock_manager_lookup(monkeypatch: pytest.MonkeyPatch) -> None:
    from app.repositories.manager_repository import ManagerRepository

    def fake_get_manager(self, manager_id: int):
        if manager_id == 1:
            return {
                "id": 1,
                "fio": "Tester",
                "contact_number": "+79990001122",
                "email": "tester@test.local",
            }
        return None

    monkeypatch.setattr(ManagerRepository, "get_manager", fake_get_manager)


@pytest.fixture()
def client(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> TestClient:
    _setup_env(monkeypatch, tmp_path)
    patch_auth_users(
        monkeypatch,
        [
            {
                "id": 1,
                "username": "tester",
                "role": "admin",
                "manager_id": None,
                "is_active": 1,
                "created_at": "2026-01-01 00:00:00",
            }
        ],
    )
    return CsrfAwareTestClient(create_app())


@pytest.fixture()
def auth_cookie(client: TestClient) -> dict[str, str]:
    token = create_session_token({"id": 1, "username": "tester", "role": "admin"}, ttl_seconds=300)
    client.cookies.set("app_session", token)
    return {"app_session": token}


def _draft_id(body: dict) -> str:
    return body.get("draft_id") or body["draft"]["draft_id"]


def _draft(body: dict) -> dict:
    return body.get("draft") or body


def test_create_draft_composite_piles(client: TestClient, auth_cookie: dict) -> None:
    response = client.post(
        "/api/v1/commercial/drafts",
        data={"product_type": "composite_piles", "text": "С140.30-С 5"},
        cookies=auth_cookie,
    )
    assert response.status_code == 200, response.text
    draft = _draft(response.json())
    meta = draft["metadata"]
    assert meta["product_type"] == "composite_piles"
    assert draft["wizard_state"]["current_step"] == "composite_piles"
    order = draft["order_data"]
    marks = {(r["mark"], r["qty"], r["concrete_grade"]) for r in order}
    assert ("С60.30-ВС.1", 5, "B25") in marks
    assert ("С80.30-НС.1", 5, "B25") in marks
    assert all(r["product_kind"] == "composite_pile" for r in order)
    assert "С140.30-С" not in {r["mark"] for r in order}


def test_ingest_update_and_calculate(
    client: TestClient, auth_cookie: dict, monkeypatch: pytest.MonkeyPatch
) -> None:
    _mock_manager_lookup(monkeypatch)
    created = client.post(
        "/api/v1/commercial/drafts",
        data={"product_type": "composite_piles", "text": "С140.30-С 2"},
        cookies=auth_cookie,
    )
    assert created.status_code == 200, created.text
    draft_id = _draft_id(created.json())

    updated = client.patch(
        f"/api/v1/commercial/drafts/{draft_id}/composite-piles",
        data={"mode": "replace", "text": "С140.30-С 2\nС60.30-ВС.1 3"},
        cookies=auth_cookie,
    )
    assert updated.status_code == 200, updated.text
    order = _draft(updated.json())["order_data"]
    by = {(r["mark"], r["concrete_grade"]): r["qty"] for r in order}
    assert by[("С60.30-ВС.1", "B25")] == 5
    assert by[("С80.30-НС.1", "B25")] == 2

    meta = client.patch(
        f"/api/v1/commercial/drafts/{draft_id}/meta",
        json={
            "manager_id": 1,
            "client_name": "ООО Тест",
            "conditions_mode": "standard",
        },
        cookies=auth_cookie,
    )
    assert meta.status_code == 200, meta.text

    calc = client.post(
        f"/api/v1/commercial/drafts/{draft_id}/calculate",
        cookies=auth_cookie,
    )
    assert calc.status_code == 200, calc.text
    assert _draft(calc.json())["wizard_state"]["current_step"] == "result"


def test_unparsed_blocks_calculate(
    client: TestClient, auth_cookie: dict, monkeypatch: pytest.MonkeyPatch
) -> None:
    _mock_manager_lookup(monkeypatch)
    created = client.post(
        "/api/v1/commercial/drafts",
        data={"product_type": "composite_piles", "text": "С999.30-С 1"},
        cookies=auth_cookie,
    )
    assert created.status_code == 200, created.text
    draft = _draft(created.json())
    assert draft["order_data"] == []
    assert draft["metadata"]["unparsed_lines"]
    draft_id = _draft_id(created.json())

    client.patch(
        f"/api/v1/commercial/drafts/{draft_id}/meta",
        json={"manager_id": 1, "client_name": "ООО Тест", "conditions_mode": "standard"},
        cookies=auth_cookie,
    )
    calc = client.post(f"/api/v1/commercial/drafts/{draft_id}/calculate", cookies=auth_cookie)
    assert calc.status_code in (400, 422)


def test_generate_files_sections_only(
    client: TestClient, auth_cookie: dict, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    _mock_manager_lookup(monkeypatch)
    created = client.post(
        "/api/v1/commercial/drafts",
        data={"product_type": "composite_piles", "text": "С140.30-С 5"},
        cookies=auth_cookie,
    )
    draft_id = _draft_id(created.json())
    client.patch(
        f"/api/v1/commercial/drafts/{draft_id}/meta",
        json={"manager_id": 1, "client_name": "ООО Тест", "conditions_mode": "standard"},
        cookies=auth_cookie,
    )
    assert (
        client.post(
            f"/api/v1/commercial/drafts/{draft_id}/calculate", cookies=auth_cookie
        ).status_code
        == 200
    )

    files = client.post(
        f"/api/v1/commercial/drafts/{draft_id}/generate-files",
        json={"file_types": ["pdf", "xlsx"]},
        cookies=auth_cookie,
    )
    assert files.status_code == 200, files.text
    kinds = {item["kind"] for item in files.json()["files"]}
    assert "pdf" in kinds and "xlsx" in kinds

    outputs = Path(tmp_path / "outputs")
    by_kind = {item["kind"]: item for item in files.json()["files"]}
    pdf_bytes = (outputs / by_kind["pdf"]["filename"]).read_bytes()
    assert "С140.30-С".encode("utf-8") not in pdf_bytes

    xlsx_bytes = (outputs / by_kind["xlsx"]["filename"]).read_bytes()
    wb = load_workbook(BytesIO(xlsx_bytes), read_only=True, data_only=True)
    texts: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            texts.extend(str(c) for c in row if c is not None)
    joined = "\n".join(texts)
    assert "С140.30-С" not in joined
    assert "С60.30-ВС.1" in joined
    assert "С80.30-НС.1" in joined


def test_save_to_archive_writes_kp_composite_piles(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    _setup_env(monkeypatch, tmp_path)
    plita_db = tmp_path / "plita.db"
    init_schema(str(plita_db))
    with sqlite3.connect(str(plita_db)) as conn:
        conn.execute(
            "INSERT INTO managers (id, fio, contact_number, email) "
            "VALUES (1, 'Tester', '+79990001122', 'tester@test.local')"
        )
        conn.commit()
    monkeypatch.setenv("PLITA_DB_PATH", str(plita_db))
    get_settings.cache_clear()
    _mock_manager_lookup(monkeypatch)
    patch_auth_users(
        monkeypatch,
        [
            {
                "id": 1,
                "username": "tester",
                "role": "admin",
                "manager_id": None,
                "is_active": 1,
                "created_at": "2026-01-01 00:00:00",
            }
        ],
    )

    client = CsrfAwareTestClient(create_app())
    token = create_session_token({"id": 1, "username": "tester", "role": "admin"}, ttl_seconds=300)
    client.cookies.set("app_session", token)

    create = client.post(
        "/api/v1/commercial/drafts",
        data={"product_type": "composite_piles", "text": "С140.30-С 5"},
    )
    assert create.status_code == 200, create.text
    draft_id = _draft_id(create.json())

    client.patch(
        f"/api/v1/commercial/drafts/{draft_id}/meta",
        json={
            "manager_id": 1,
            "client_name": "ООО Составные",
            "conditions_mode": "standard",
            "counterparty_id": fx.seed_test_counterparty(str(plita_db)),
        },
    )
    assert client.post(f"/api/v1/commercial/drafts/{draft_id}/calculate").status_code == 200

    save = client.post(
        f"/api/v1/commercial/drafts/{draft_id}/save",
        json={"mode": "archive", "execution_terms_input": ""},
    )
    assert save.status_code == 200, save.text
    kp_id = save.json()["saved_offer"]["kp_id"]

    with sqlite3.connect(str(plita_db)) as conn:
        meta = conn.execute(
            "SELECT product_type FROM kp_meta WHERE kp_id = ?", (kp_id,)
        ).fetchone()
        assert meta and meta[0] == "composite_piles"
        marks = {
            r[0]
            for r in conn.execute(
                "SELECT mark FROM kp_composite_piles WHERE kp_id = ?", (kp_id,)
            ).fetchall()
        }
        assert "С60.30-ВС.1" in marks
        assert "С80.30-НС.1" in marks
        assert "С140.30-С" not in marks
        assert (
            conn.execute("SELECT COUNT(*) FROM kp_piles WHERE kp_id = ?", (kp_id,)).fetchone()[0]
            == 0
        )
