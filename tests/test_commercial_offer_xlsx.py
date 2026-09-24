"""НДС в XLSX: число для нового kp_id, формула × 0,22 для старого."""

from __future__ import annotations

from decimal import Decimal
from typing import Any

from openpyxl import load_workbook

from core.commercial_offer_xlsx import calculate_total_cost, generate_commercial_offer_xlsx
from core.commercial_pricing import VAT_PER_LINE_FROM_KP_ID, vat_included_from_line_sums
from core.embed_delivery_in_unit_price import embed_delivery_in_unit_prices


def _order() -> list[dict[str, Any]]:
    return [
        {
            "name": "ПБ 60-12-8п",
            "length_m": 6.0,
            "width_m": 1.2,
            "qty": 2,
            "load_class": 800,
            "unit_price": 1000.0,
            "weight": 1000.0,
        }
    ]


def _sheet(kp_db_id: int | None, *, embed: bool = False):
    buf = generate_commercial_offer_xlsx(
        _order(),
        offer_number="N",
        offer_date="24.09.2026",
        customer_name="ООО Тест",
        discount_percent=0,
        logistics_cost=500.0,
        kp_db_id=kp_db_id,
        embed_delivery_in_unit_price=embed,
    )
    return load_workbook(buf)["КП"]


def _vat_cell(ws):
    for row in ws.iter_rows(min_row=1, max_row=80, max_col=8):
        for cell in row:
            if cell.value == "в том числе НДС (22%)":
                for col in range(1, 9):
                    value = ws.cell(row=cell.row, column=col).value
                    if value not in (None, "в том числе НДС (22%)"):
                        return value
    raise AssertionError("VAT row not found")


def test_old_kp_xlsx_keeps_vat_rate_formula() -> None:
    formula = _vat_cell(_sheet(26))
    assert isinstance(formula, str)
    assert "*0.22" in formula.replace(" ", "")


def test_new_kp_xlsx_writes_numeric_vat_matching_totals() -> None:
    order = _order()
    totals = calculate_total_cost(
        order, discount_percent=0, logistics_cost=500.0, kp_id=VAT_PER_LINE_FROM_KP_ID
    )
    value = _vat_cell(_sheet(VAT_PER_LINE_FROM_KP_ID))
    assert isinstance(value, (int, float))
    assert abs(float(value) - float(totals["vat_amount"])) < 0.001


def test_new_embed_xlsx_vat_comes_from_embedded_line_sums() -> None:
    order = _order()
    totals = calculate_total_cost(
        order, discount_percent=0, logistics_cost=500.0, kp_id=None
    )
    embedded = embed_delivery_in_unit_prices(
        qty_by_index=[2],
        product_type_by_index=["plates"],
        unit_price_by_index=[1000.0],
        discount_percent=0,
        plate_delivery_total=float(totals["plate_delivery_total"]),
        pile_delivery_total=0.0,
    )
    expected = float(
        vat_included_from_line_sums([Decimal(str(line.line_sum)) for line in embedded.lines])
    )
    value = _vat_cell(_sheet(None, embed=True))
    assert isinstance(value, (int, float))
    assert abs(float(value) - expected) < 0.001
    plain = float(_vat_cell(_sheet(None, embed=False)))
    assert abs(plain - float(totals["vat_amount"])) < 0.001
