"""Шапка PDF и XLSX: контакты менеджера и подавление заводского fallback."""

from __future__ import annotations

import subprocess
import tempfile
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from core.commercial_offer import generate_commercial_offer_pdf
from core.commercial_offer_xlsx import generate_commercial_offer_xlsx

MANAGER_NAME = "Зубов Алексей Юрьевич"
PHONE = "79621872265"
EMAIL = "zubov@gbkstart.ru"
FORMATTED_HEADER = "Тел.: +7 (962) 187-22-65, email: zubov@gbkstart.ru"
COMPANY_PHONE = "8 (4852) 595 000"
COMPANY_EMAIL = "info@zhbk.ru"
COMPANY_HEADER = f"Тел.: {COMPANY_PHONE}, Email: {COMPANY_EMAIL}"


def _order() -> list[dict[str, Any]]:
    return [
        {
            "name": "ПБ 60-12-8п",
            "length_m": 6.0,
            "width_m": 1.2,
            "qty": 1,
            "load_class": 800,
            "unit_price": 1000.0,
            "weight": 1000.0,
        }
    ]


def _pdf_text(**overrides: Any) -> str:
    buffer = generate_commercial_offer_pdf(
        _order(),
        offer_number="11",
        offer_date="21.10.2026",
        customer_name="ООО «АВТОВИД»",
        manager_name=MANAGER_NAME,
        **overrides,
    )
    with tempfile.NamedTemporaryFile(suffix=".pdf") as handle:
        handle.write(buffer.getvalue())
        handle.flush()
        result = subprocess.run(
            ["pdftotext", "-layout", handle.name, "-"],
            check=True,
            capture_output=True,
        )
    return result.stdout.decode("utf-8", errors="replace")


def _xlsx_sheet(**overrides: Any) -> Worksheet:
    buffer = generate_commercial_offer_xlsx(
        _order(),
        offer_number="11",
        offer_date="21.10.2026",
        customer_name="ООО «АВТОВИД»",
        manager_name=MANAGER_NAME,
        **overrides,
    )
    return load_workbook(buffer)["КП"]


def _xlsx_text(sheet: Worksheet) -> str:
    parts: list[str] = []
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is not None:
                parts.append(str(cell.value))
    return "\n".join(parts)


def _signature_value(sheet: Worksheet) -> str:
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value == "С уважением,":
                below = sheet.cell(row=cell.row + 1, column=cell.column).value
                return "" if below is None else str(below)
    raise AssertionError("строка «С уважением,» не найдена")


def test_pdf_header_prints_formatted_phone_and_email() -> None:
    text = _pdf_text(manager_phone=PHONE, manager_email=EMAIL)

    assert FORMATTED_HEADER in text
    assert COMPANY_PHONE not in text
    assert COMPANY_EMAIL not in text
    assert MANAGER_NAME in text
    assert text.count("+7 (962) 187-22-65") == 1


def test_xlsx_header_prints_formatted_phone_and_email() -> None:
    text = _xlsx_text(_xlsx_sheet(manager_phone=PHONE, manager_email=EMAIL))

    assert FORMATTED_HEADER in text
    assert COMPANY_PHONE not in text
    assert COMPANY_EMAIL not in text


def test_pdf_header_phone_only_omits_email_and_company() -> None:
    text = _pdf_text(manager_phone=PHONE, manager_email=None)

    assert "Тел.: +7 (962) 187-22-65" in text
    assert "email:" not in text
    assert "Email:" not in text
    assert COMPANY_PHONE not in text
    assert COMPANY_EMAIL not in text


def test_xlsx_header_email_only_omits_phone_and_company() -> None:
    text = _xlsx_text(_xlsx_sheet(manager_phone=None, manager_email=EMAIL))

    assert f"Email: {EMAIL}" in text
    assert "Тел.:" not in text
    assert COMPANY_PHONE not in text
    assert COMPANY_EMAIL not in text


def test_pdf_suppresses_company_when_both_contacts_empty() -> None:
    text = _pdf_text(
        manager_phone=None,
        manager_email=None,
        suppress_company_contacts=True,
    )

    assert COMPANY_PHONE not in text
    assert COMPANY_EMAIL not in text
    assert "Тел.:" not in text
    assert "Email:" not in text
    assert MANAGER_NAME in text


def test_xlsx_suppresses_company_when_both_contacts_empty() -> None:
    text = _xlsx_text(
        _xlsx_sheet(
            manager_phone="",
            manager_email="   ",
            suppress_company_contacts=True,
        )
    )

    assert COMPANY_PHONE not in text
    assert COMPANY_EMAIL not in text
    assert "Тел.:" not in text
    assert "Email:" not in text


def test_pdf_and_xlsx_keep_company_fallback_without_flag() -> None:
    pdf_text = _pdf_text(manager_phone=None, manager_email=None)
    xlsx_text = _xlsx_text(_xlsx_sheet(manager_phone=None, manager_email=None))

    assert COMPANY_HEADER in pdf_text
    assert COMPANY_HEADER in xlsx_text


def test_xlsx_signature_cell_is_manager_name_only() -> None:
    sheet = _xlsx_sheet(manager_phone=PHONE, manager_email=EMAIL)
    signature = _signature_value(sheet)

    assert signature == MANAGER_NAME
    assert PHONE not in signature
    assert EMAIL not in signature
    assert "+7" not in signature
