"""Сохранение спецификации, файл и ворота производства."""

from __future__ import annotations

import json
import sqlite3
from pathlib import Path

import pytest


def test_schema_specification_json_survives_repeat(tmp_path: Path) -> None:
    from core.kp_db_schema import _init_schema_impl

    db_path = str(tmp_path / "old.db")
    with sqlite3.connect(db_path) as conn:
        conn.execute(
            """
            CREATE TABLE KP_offers (
                kp_id INTEGER PRIMARY KEY AUTOINCREMENT,
                creation_date TEXT NOT NULL
            )
            """
        )

    _init_schema_impl(db_path)
    _init_schema_impl(db_path)

    with sqlite3.connect(db_path) as conn:
        columns = {row[1] for row in conn.execute("PRAGMA table_info(KP_offers)")}
    assert "specification_json" in columns


def test_specification_json_roundtrip_and_empty_reads_absent(tmp_path: Path) -> None:
    from core.kp.offers_read import get_specification_json
    from core.kp.offers_write import set_kp_specification_json
    from core.kp_db_schema import init_schema

    db_path = str(tmp_path / "plita.db")
    init_schema(db_path)
    with sqlite3.connect(db_path) as conn:
        conn.execute(
            "INSERT INTO KP_offers (creation_date, customer_name) VALUES ('29.09.2026', 'Пусто')"
        )
        conn.execute(
            "INSERT INTO KP_offers (creation_date, customer_name) VALUES ('29.09.2026', 'Есть')"
        )
        empty_id, saved_id = [
            row[0]
            for row in conn.execute("SELECT kp_id FROM KP_offers ORDER BY kp_id")
        ]

    assert get_specification_json(empty_id, db_path) is None

    document = {
        "payment": "prepay_100",
        "term": "by_date",
        "delivery": "pickup",
        "spec_date": "2026-09-29",
        "composition_hash": "abc",
    }
    assert set_kp_specification_json(saved_id, json.dumps(document, ensure_ascii=False), db_path)
    loaded = json.loads(get_specification_json(saved_id, db_path) or "")
    assert loaded == document
    assert get_specification_json(empty_id, db_path) is None


ADMIN = {"id": 1, "role": "admin"}
_CONDITION_MARK = "КЛЮЧ-УСЛОВИЙ-КП"


def _seed_plate(
    tmp_path: Path,
    *,
    status: str = "на согласовании",
    discount: float = 0,
    logistics: float = 0,
    payment_conditions: str = "",
    delivery_conditions: str = "",
    length_m: float = 0,
    width_m: float = 0,
    unit_price: float = 1000,
    qty: int = 1,
    counterparty_id: int | None = None,
) -> tuple[str, int]:
    from core.kp_db_schema import init_schema

    db_path = str(tmp_path / "plita.db")
    init_schema(db_path)
    with sqlite3.connect(db_path) as conn:
        cur = conn.execute(
            """
            INSERT INTO KP_offers (
                creation_date, customer_name, discount_percent, total_amount,
                logistics_cost, payment_conditions, delivery_conditions, counterparty_id
            ) VALUES ('29.09.2026', 'Ромашка', ?, 0, ?, ?, ?, ?)
            """,
            (discount, logistics, payment_conditions, delivery_conditions, counterparty_id),
        )
        kp_id = int(cur.lastrowid)
        conn.execute(
            """
            INSERT INTO kp_plates (
                kp_id, position_number, plate_name, qty, unit_price, length_m, width_m
            ) VALUES (?, 1, 'ПБ 60-12-8п', ?, ?, ?, ?)
            """,
            (kp_id, qty, unit_price, length_m, width_m),
        )
        conn.execute(
            "INSERT INTO kp_meta (kp_id, status, owner_user_id) VALUES (?, ?, 1)",
            (kp_id, status),
        )
    return db_path, kp_id


def _service(db_path: str, tmp_path: Path):
    from app.repositories.kp_archive_repository import KpArchiveRepository
    from app.services.archive_service import ArchiveService

    return ArchiveService(
        repository=KpArchiveRepository(db_path=db_path),
        outputs_dir=tmp_path / "out",
        export_dir=tmp_path / "exchange",
    )


def _choice(**overrides: object):
    from app.schemas.archive import SpecificationChoiceIn

    body = {
        "payment": "prepay_100",
        "term": "by_date",
        "term_date": "2026-10-01",
        "delivery": "pickup",
    }
    body.update(overrides)
    return SpecificationChoiceIn(**body)  # type: ignore[arg-type]


def _stored(db_path: str, kp_id: int) -> dict | None:
    from core.kp.offers_read import get_specification_json

    text = get_specification_json(kp_id, db_path)
    if text is None:
        return None
    return json.loads(text)


def _set_discount(db_path: str, kp_id: int, discount: float) -> None:
    with sqlite3.connect(db_path) as conn:
        conn.execute(
            "UPDATE KP_offers SET discount_percent = ? WHERE kp_id = ?",
            (discount, kp_id),
        )


def test_specification_incomplete_and_pickup_with_delivery_do_not_write(tmp_path: Path) -> None:
    from app.services.archive_service import ArchiveValidationError

    db_path, kp_id = _seed_plate(tmp_path)
    service = _service(db_path, tmp_path)

    with pytest.raises(ArchiveValidationError):
        service.save_specification(
            kp_id,
            _choice(payment="split_50_50", payment_days=3),
            user=ADMIN,
        )
    assert _stored(db_path, kp_id) is None

    delivered = tmp_path / "delivered"
    delivered.mkdir()
    db_path, kp_id = _seed_plate(
        delivered,
        logistics=2500,
        length_m=6,
        width_m=1.2,
        delivery_conditions="самовывоз",
    )
    service = _service(db_path, delivered)
    with pytest.raises(ArchiveValidationError, match="доставк"):
        service.save_specification(kp_id, _choice(), user=ADMIN)
    assert _stored(db_path, kp_id) is None


def test_specification_percent_tracks_discount_and_custom_text_goes_stale(tmp_path: Path) -> None:
    db_path, kp_id = _seed_plate(
        tmp_path,
        payment_conditions=_CONDITION_MARK,
        delivery_conditions="доставка на объект",
    )
    service = _service(db_path, tmp_path)
    service._today_override = "2026-09-18"
    saved = service.save_specification(
        kp_id,
        _choice(
            payment="split_50_50",
            payment_date="2026-09-20",
            payment_days=3,
            delivery="site",
            delivery_address="г. Кострома, ул. Лесная, 4",
        ),
        user=ADMIN,
    )

    assert "500,00" in saved.payment_paragraph
    assert _CONDITION_MARK not in saved.payment_paragraph
    assert _CONDITION_MARK not in saved.term_paragraph
    assert _CONDITION_MARK not in saved.delivery_paragraph
    assert saved.stale_custom is False
    first_hash = _stored(db_path, kp_id)["composition_hash"]
    assert _stored(db_path, kp_id)["spec_date"] == "2026-09-18"

    _set_discount(db_path, kp_id, 50)
    viewed = service.get_specification(kp_id, user=ADMIN)
    details = service.get_details(kp_id, user=ADMIN)

    assert "250,00" in viewed.payment_paragraph
    assert "500,00" not in viewed.payment_paragraph
    assert viewed.stale_custom is False
    assert details.specification_saved is True
    assert details.specification_stale_custom is False
    assert _stored(db_path, kp_id)["composition_hash"] == first_hash
    assert _CONDITION_MARK not in viewed.payment_paragraph

    service._today_override = "2026-10-02"
    custom = service.save_specification(
        kp_id,
        _choice(
            payment="custom",
            custom_text="Два транша по договорённости сторон.",
            delivery="site",
            delivery_address="г. Кострома, ул. Лесная, 4",
        ),
        user=ADMIN,
    )
    assert custom.spec_date == "2026-09-18"
    assert custom.payment_paragraph == "Два транша по договорённости сторон."
    assert custom.stale_custom is False
    fresh_hash = _stored(db_path, kp_id)["composition_hash"]
    assert fresh_hash != first_hash

    _set_discount(db_path, kp_id, 10)
    stale = service.get_specification(kp_id, user=ADMIN)
    details = service.get_details(kp_id, user=ADMIN)

    assert stale.payment_paragraph == "Два транша по договорённости сторон."
    assert stale.stale_custom is True
    assert details.specification_stale_custom is True
    assert _CONDITION_MARK not in stale.payment_paragraph
    assert _stored(db_path, kp_id)["spec_date"] == "2026-09-18"


def test_specification_put_refuses_archive_and_in_production(tmp_path: Path) -> None:
    from app.services.archive_service import ArchiveValidationError

    for status in ("в архиве", "в работе"):
        folder = tmp_path / status
        folder.mkdir()
        db_path, kp_id = _seed_plate(folder, status=status)
        service = _service(db_path, folder)
        with pytest.raises(ArchiveValidationError, match="на согласовании"):
            service.save_specification(kp_id, _choice(), user=ADMIN)
        assert _stored(db_path, kp_id) is None


def test_specification_get_suggests_prepay_and_pickup_without_copying_conditions(
    tmp_path: Path,
) -> None:
    db_path, kp_id = _seed_plate(
        tmp_path,
        payment_conditions="",
        delivery_conditions="Доставка: САМОВЫВОЗ, звонить Васе",
    )
    hinted = _service(db_path, tmp_path).get_specification(kp_id, user=ADMIN)

    assert hinted.saved is False
    assert hinted.choice is not None
    assert hinted.choice.payment == "prepay_100"
    assert hinted.choice.delivery == "pickup"
    assert "1 000,00 (одна тысяча рублей)" in hinted.payment_paragraph
    assert "Домостроителей" in hinted.delivery_paragraph
    assert "Васе" not in hinted.delivery_paragraph
    assert "Васе" not in hinted.payment_paragraph

    other = tmp_path / "custom-conditions"
    other.mkdir()
    db_path, kp_id = _seed_plate(
        other,
        payment_conditions=_CONDITION_MARK,
        delivery_conditions="доставка на объект",
    )
    plain = _service(db_path, other).get_specification(kp_id, user=ADMIN)

    assert plain.choice is not None
    assert plain.choice.payment is None
    assert plain.choice.delivery is None
    assert _CONDITION_MARK not in plain.payment_paragraph
    assert _CONDITION_MARK not in plain.delivery_paragraph
    assert plain.has_piles is False


def _active_contract(self, counterparty_id):  # noqa: ANN001
    return {
        "counterparty_id": int(counterparty_id),
        "status": "нет",
        "number": "0001/09/26",
        "contract_date": "2026-09-01",
        "legal_form": "ooo",
        "full_name": "Ромашка",
        "short_name": "Ромашка",
        "signatory_position": "Директор",
        "signatory_name": "Иванов Иван Иванович",
        "signatory_verb": "действующего",
        "authority_basis": "устав",
        "inn": "7604368921",
        "kpp": "760401001",
    }


def _forbid_invoice_guids(*_args: object, **_kwargs: object) -> None:
    raise AssertionError("guid gate")


def test_specification_download_file_uses_saved_date_and_invoice_number(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from app.services.archive_service import ArchiveValidationError
    from core.supply_contract import build_preamble

    monkeypatch.setattr(
        "app.services.archive_service.SupplyContractRepository.get_active_by_counterparty",
        _active_contract,
    )
    monkeypatch.setattr(
        "app.services.archive_service.assert_invoice_guids",
        _forbid_invoice_guids,
    )
    db_path, kp_id = _seed_plate(tmp_path, counterparty_id=5, qty=2, unit_price=1000)
    service = _service(db_path, tmp_path)
    export_dir = tmp_path / "exchange"

    with pytest.raises(ArchiveValidationError, match="Сначала сохраните"):
        service.download_specification(kp_id, user=ADMIN)
    assert _stored(db_path, kp_id) is None
    assert not export_dir.exists()

    service._today_override = "2026-09-18"
    service.save_specification(kp_id, _choice(), user=ADMIN)
    service._today_override = "2026-10-02"
    first = service.download_specification(kp_id, user=ADMIN)
    text = _workbook_text(first.content)

    assert first.filename == f"Спецификация КП {kp_id}.xlsx"
    assert "№ _____" in text
    assert "от 18.09.2026" in text
    assert "18 сентября 2026\u00a0г." in text
    assert "ЯР-15" not in text
    assert "4401082520" in text
    assert "Комбинат ЖБК" not in text
    assert "4705123456" not in text
    preamble = build_preamble(
        legal_form="ooo",
        full_name="Ромашка",
        signatory_name="Иванов Иван Иванович",
        signatory_position="Директор",
        authority_basis="устав",
    )
    assert preamble in text
    assert text.count("в лице") == preamble.count("в лице") + 1
    assert _labeled_total(first.content) == _invoice_total(db_path, kp_id)
    assert _stored(db_path, kp_id)["spec_date"] == "2026-09-18"
    assert not export_dir.exists()

    with sqlite3.connect(db_path) as conn:
        conn.execute(
            "UPDATE KP_offers SET order_number_1c = ? WHERE kp_id = ?",
            ("ЯР-15", kp_id),
        )
    second = service.download_specification(kp_id, user=ADMIN)
    numbered = _workbook_text(second.content)

    assert second.filename == "Спецификация по счету ЯР-15.xlsx"
    assert "№ ЯР-15" in numbered
    assert "от 18.09.2026" in numbered
    assert _stored(db_path, kp_id)["spec_date"] == "2026-09-18"
    assert not export_dir.exists()
    assert list(tmp_path.glob("*.xlsx")) == []


_STALE = (
    "Условия оплаты в спецификации устарели: откройте спецификацию и сохраните их снова"
)


def _mark_invoice_ready(db_path: str, kp_id: int, *, paid: bool = True, number: str | None = "ЯР-15") -> None:
    with sqlite3.connect(db_path) as conn:
        conn.execute(
            "UPDATE kp_meta SET paid_at = ? WHERE kp_id = ?",
            ("2026-09-28T10:00:00" if paid else None, kp_id),
        )
        conn.execute(
            "UPDATE KP_offers SET order_number_1c = ? WHERE kp_id = ?",
            (number, kp_id),
        )


def _status(db_path: str, kp_id: int) -> str:
    with sqlite3.connect(db_path) as conn:
        row = conn.execute("SELECT status FROM kp_meta WHERE kp_id = ?", (kp_id,)).fetchone()
    return str(row[0])


def _allow_production(db_path: str):
    from unittest.mock import MagicMock

    from core.kp.offers_write import update_kp_status

    promise = MagicMock()

    def _commit(kp_id: int, _terms: str, *, user: dict, raw: dict | None = None) -> None:
        update_kp_status(kp_id, "в работе", db_path)

    promise.commit_move_with_gate.side_effect = _commit
    return promise


def test_move_to_production_keeps_payment_and_number_and_requires_specification(
    tmp_path: Path,
) -> None:
    from app.services.archive_service import ArchiveValidationError

    db_path, kp_id = _seed_plate(tmp_path)
    service = _service(db_path, tmp_path)
    service._promise_service = _allow_production(db_path)

    with pytest.raises(ArchiveValidationError, match="Сначала отметьте оплату"):
        service.move_to_production(kp_id, "15.10.2026", user=ADMIN)

    _mark_invoice_ready(db_path, kp_id, paid=True, number=None)
    with pytest.raises(ArchiveValidationError, match="Нужен номер счёта"):
        service.move_to_production(kp_id, "15.10.2026", user=ADMIN)

    _mark_invoice_ready(db_path, kp_id)
    with pytest.raises(ArchiveValidationError, match="Сначала сохраните спецификацию"):
        service.move_to_production(kp_id, "15.10.2026", user=ADMIN)

    assert _status(db_path, kp_id) == "на согласовании"
    service._promise_service.commit_move_with_gate.assert_not_called()


def test_move_to_production_allows_percent_after_discount_and_blocks_stale_custom(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from app.services.archive_service import ArchiveValidationError

    monkeypatch.setattr(
        "app.services.archive_service.CounterpartiesService.require_active_client",
        lambda self, cid: {"id": 1, "name": "Ромашка"},
    )
    db_path, kp_id = _seed_plate(tmp_path)
    service = _service(db_path, tmp_path)
    service._promise_service = _allow_production(db_path)
    service.save_specification(
        kp_id,
        _choice(payment="split_50_50", payment_date="2026-09-20", payment_days=3),
        user=ADMIN,
    )
    _mark_invoice_ready(db_path, kp_id)
    _set_discount(db_path, kp_id, 50)

    moved = service.move_to_production(kp_id, "15.10.2026", user=ADMIN)
    assert moved.status == "в работе"

    custom_dir = tmp_path / "custom"
    custom_dir.mkdir()
    db_path, kp_id = _seed_plate(custom_dir)
    service = _service(db_path, custom_dir)
    service._promise_service = _allow_production(db_path)
    service.save_specification(
        kp_id,
        _choice(payment="custom", custom_text="Два транша по договорённости сторон."),
        user=ADMIN,
    )
    _mark_invoice_ready(db_path, kp_id)
    _set_discount(db_path, kp_id, 10)

    with pytest.raises(ArchiveValidationError, match=_STALE):
        service.move_to_production(kp_id, "15.10.2026", user=ADMIN)
    assert _status(db_path, kp_id) == "на согласовании"

    service.save_specification(
        kp_id,
        _choice(payment="custom", custom_text="Два транша по договорённости сторон."),
        user=ADMIN,
    )
    moved = service.move_to_production(kp_id, "15.10.2026", user=ADMIN)
    assert moved.status == "в работе"


def _workbook_text(data: bytes) -> str:
    import io

    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data))
    chunks: list[str] = []
    for row in workbook.active.iter_rows(values_only=True):
        for value in row:
            if value is not None:
                chunks.append(str(value))
    return "\n".join(chunks)


def _labeled_total(data: bytes) -> float:
    import io

    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data))
    sheet = workbook.active
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.strip() == "Итого:":
                for other in sheet[cell.row]:
                    if isinstance(other.value, (int, float)) and other.column > cell.column:
                        return float(other.value)
    raise AssertionError("нет итога")


def _invoice_total(db_path: str, kp_id: int) -> float:
    from app.repositories.kp_archive_repository import KpArchiveRepository
    from app.services.archive_service import ArchiveService
    from core.invoice_export import build_invoice_document, card_offer_from_kp

    raw = KpArchiveRepository(db_path=db_path).get_by_id(kp_id)
    assert raw is not None
    service = ArchiveService(
        repository=KpArchiveRepository(db_path=db_path),
        outputs_dir=Path(db_path).parent / "out-check",
    )
    offer = card_offer_from_kp(raw)
    shares = service._invoice_delivery_kopecks(raw, offer)
    document, _digest = build_invoice_document(offer, line_delivery_kopecks=shares)
    return float(document["СуммаДокумента"])
