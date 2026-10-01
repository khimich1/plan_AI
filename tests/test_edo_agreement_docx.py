"""Соглашение об ЭДО: преамбула стороны-2 и подстановка в бланк."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document

from core.edo_agreement_docx import (
    buyer_okved_line,
    render_edo_agreement_docx,
    side2_preamble,
    side2_requisite_lines,
)

_TEMPLATE = Path(__file__).resolve().parents[1] / "templates" / "edo_agreement.docx"

_PSK = {
    "number": "1028/09/26",
    "contract_date": "2026-09-22",
    "legal_form": "ooo",
    "full_name": "Промышленно-строительные конструкции",
    "short_name": "ПромСтройКонструкции",
    "signatory_position": "Генерального директора",
    "signatory_name": "Недотко Александр Александрович",
    "signatory_verb": "действующего",
    "authority_basis": "устав",
    "inn": "7814192061",
    "kpp": "784201001",
    "ogrn": "1157847095340",
    "legal_address": "191144, г. Санкт-Петербург, ул. Новгородская, д. 13, литера Л",
    "postal_address": "191167, г. Санкт-Петербург, а/я 27",
    "email": "info@gbi-psk.ru",
    "okved": "23.61",
}


def _docx_text(data: bytes) -> str:
    document = Document(BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for section in document.sections:
        parts.extend(paragraph.text for paragraph in section.header.paragraphs)
        parts.extend(paragraph.text for paragraph in section.footer.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(paragraph.text for paragraph in cell.paragraphs)
    return "\n".join(parts)


def test_ip_preamble_uses_ogrnip_not_charter() -> None:
    text = side2_preamble(
        {
            "legal_form": "ip",
            "full_name": "Иванов Иван Иванович",
            "signatory_name": "Иванов Иван Иванович",
            "signatory_verb": "действующего",
            "authority_basis": "устав",
            "ogrn": "323330000060332",
        }
    )

    assert "Индивидуальный предприниматель Иванов Иван Иванович" in text
    assert "ОГРНИП 323330000060332" in text
    assert "Сторона-2" in text
    assert "Устава" not in text
    assert "устав" not in text.lower()


def test_ooo_preamble_uses_face_and_side_two() -> None:
    text = side2_preamble(_PSK)

    assert "в лице" in text
    assert "Генерального директора" in text
    assert "Сторона-2" in text
    assert "Покупатель" not in text


def test_kfh_and_person_preambles_name_side_two() -> None:
    kfh = side2_preamble(
        {
            "legal_form": "kfh",
            "full_name": "Заря",
            "signatory_name": "Петров Пётр",
            "signatory_verb": "действующего",
            "authority_basis": "устав",
        }
    )
    person = side2_preamble(
        {
            "legal_form": "person",
            "full_name": "Смирнова Ольга",
            "signatory_name": "Смирнова Ольга",
            "signatory_verb": "действующей",
            "authority_basis": "устав",
        }
    )

    assert "Крестьянское (фермерское) хозяйство" in kfh
    assert "Сторона-2" in kfh
    assert "Покупатель" not in kfh
    assert "Смирнова Ольга" in person
    assert "действующей" in person
    assert "Сторона-2" in person


def test_empty_okved_and_edo_do_not_print_lines() -> None:
    assert buyer_okved_line(None) is None
    assert buyer_okved_line("  ") is None
    assert buyer_okved_line("23.61") == "ОКВЭД 23.61"
    lines = side2_requisite_lines({**_PSK, "okved": "", "edo_operator": None, "edo_id": ""})
    assert all("ОКВЭД" not in line for line in lines)
    assert all("Оператор" not in line for line in lines)


def test_edo_line_appears_when_either_field_is_set() -> None:
    by_operator = side2_requisite_lines({**_PSK, "edo_operator": "ООО «Такском»", "edo_id": ""})
    by_id = side2_requisite_lines({**_PSK, "edo_operator": None, "edo_id": "2BM-7814192061"})

    assert any("ООО «Такском»" in line for line in by_operator)
    assert any("2BM-7814192061" in line for line in by_id)


def test_postal_address_is_printed_only_when_it_differs() -> None:
    different = side2_requisite_lines(_PSK)
    same = side2_requisite_lines({**_PSK, "postal_address": _PSK["legal_address"]})

    assert any("191167, г. Санкт-Петербург, а/я 27" in line for line in different)
    assert all("Почтовый адрес" not in line for line in same)


def test_template_keeps_side_one_and_has_no_filled_example() -> None:
    text = _docx_text(_TEMPLATE.read_bytes())

    assert "Бисеров" not in text
    assert "2026" not in text
    assert "{{EDO_DATE_SHORT}}" in text
    assert "{{EDO_DATE_LONG}}" in text
    assert "{{SIDE2_PREAMBLE}}" in text
    assert "{{SIDE2_REQUISITES}}" in text
    assert "start@gbkstart.ru" in text
    assert "buh@gbkstart.ru" in text
    assert "pravo@gbkstart.ru" in text
    assert "Тензор" in text
    assert "2BE8508b4c4e53411e2925a005056917125" in text
    assert "Шишов Александр Васильевич" in text
    assert "Сторона-1" in text
    assert "Сторона-2" in text


def test_psk_agreement_substitutes_buyer_and_drops_the_example_name() -> None:
    blank = _docx_text(_TEMPLATE.read_bytes())
    text = _docx_text(render_edo_agreement_docx(_PSK))

    assert "7814192061" in text
    assert "info@gbi-psk.ru" in text
    assert "ОКВЭД 23.61" in text
    assert "Бисеров" not in text
    assert "start@gbkstart.ru" in text
    assert "Тензор" in text
    assert "191167, г. Санкт-Петербург, а/я 27" in text
    assert "1028/09/26" not in text
    assert "{{" not in text
    assert text.count("ОКВЭД 23.61") == blank.count("ОКВЭД 23.61") + blank.count("{{SIDE2_REQUISITES}}")
    assert "1.1. Электронный документ" in text
    assert "ПРИЛОЖЕНИЕ 1" in text
    assert "12. ПОДПИСИ И РЕКВИЗИТЫ СТОРОН" in text


def test_date_year_comes_from_contract_date() -> None:
    text = _docx_text(
        render_edo_agreement_docx({**_PSK, "contract_date": "2027-01-05", "okved": None})
    )
    blank = _docx_text(_TEMPLATE.read_bytes())

    assert "2027" in text
    assert "«05» января 2027" in text
    assert "05.01.2027" in text
    assert "2026" not in text
    assert text.count("ОКВЭД") == blank.count("ОКВЭД")


def test_ip_agreement_prints_ogrnip_in_the_file() -> None:
    text = _docx_text(
        render_edo_agreement_docx(
            {
                **_PSK,
                "legal_form": "ip",
                "full_name": "Иванов Иван Иванович",
                "signatory_name": "Иванов Иван Иванович",
                "signatory_position": None,
                "kpp": None,
                "authority_basis": "устав",
                "ogrn": "323330000060332",
                "okved": None,
            }
        )
    )

    assert "ОГРНИП 323330000060332" in text
    assert "Иванов Иван Иванович" in text
    lines = side2_requisite_lines(
        {
            **_PSK,
            "legal_form": "ip",
            "full_name": "Иванов Иван Иванович",
            "signatory_name": "Иванов Иван Иванович",
            "signatory_position": None,
            "kpp": "784201001",
            "ogrn": "323330000060332",
        }
    )
    assert all(not line.startswith("КПП") for line in lines)
