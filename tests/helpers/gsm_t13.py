"""Build a minimal Т-13 workbook in memory (no personal files in git)."""

from __future__ import annotations

import calendar
import io

from openpyxl import Workbook
from openpyxl.utils import column_index_from_string

CODE_COLUMNS = (
    "I",
    "K",
    "M",
    "N",
    "P",
    "R",
    "T",
    "V",
    "X",
    "Z",
    "AB",
    "AD",
    "AF",
    "AH",
    "AK",
)


def build_t13_xlsx(
    people: list[dict],
    *,
    year: int = 2026,
    month: int = 7,
    include_header: bool = True,
    shift_day_columns: bool = False,
) -> bytes:
    """``people`` items: full_name, personnel, codes (day→code), optional hours."""
    last_day = calendar.monthrange(year, month)[1]
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    if include_header:
        sheet["A1"] = "ТАБЕЛЬ учета рабочего времени"
    header_row = 3
    second_row = 4
    for offset, col in enumerate(CODE_COLUMNS):
        target = "J" if shift_day_columns else col
        sheet.cell(header_row, column_index_from_string(target), offset + 1)
        if offset + 16 <= 30:
            sheet.cell(second_row, column_index_from_string(target), offset + 16)
    if not shift_day_columns:
        sheet.cell(second_row, column_index_from_string("AL"), 31)

    start = 6
    for index, person in enumerate(people):
        row = start + index * 4
        sheet.cell(row, 3, person["full_name"])
        sheet.cell(row, 5, person.get("personnel", ""))
        codes: dict[int, str] = dict(person.get("codes") or {})
        default = person.get("default_code", "Я")
        hours = person.get("hours", "8")
        for day_num in range(1, last_day + 1):
            code = codes.get(day_num, default)
            _write_code(sheet, row, day_num, code)
            _write_hours(sheet, row, day_num, hours)

    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


def _write_code(sheet, person_row: int, day_num: int, code: str) -> None:
    if day_num <= 15:
        col = column_index_from_string(CODE_COLUMNS[day_num - 1])
        sheet.cell(person_row, col, code)
        return
    second = person_row + 2
    if day_num == 31:
        sheet.cell(second, column_index_from_string("AL"), code)
        return
    col = column_index_from_string(CODE_COLUMNS[day_num - 16])
    sheet.cell(second, col, code)


def _write_hours(sheet, person_row: int, day_num: int, hours: str) -> None:
    if day_num <= 15:
        col = column_index_from_string(CODE_COLUMNS[day_num - 1])
        sheet.cell(person_row + 1, col, hours)
        return
    second_hours = person_row + 3
    if day_num == 31:
        sheet.cell(second_hours, column_index_from_string("AL"), hours)
        return
    col = column_index_from_string(CODE_COLUMNS[day_num - 16])
    sheet.cell(second_hours, col, hours)
