"""Печать договора поставки из сохранённых полей (SC-005)."""

from __future__ import annotations

import sqlite3
from io import BytesIO
from pathlib import Path

from docx import Document

from app.schemas.supply_contract import SupplyContractCreate
from app.services.supply_contract_service import SupplyContractService
from core import kp_db_schema
from core.kp_db_common import _connect

ADMIN = {"id": 7, "role": "admin", "username": "admin"}


def _fresh_db(tmp_path: Path) -> str:
    db_path = str(tmp_path / "supply.db")
    kp_db_schema._schema_ready.clear()
    kp_db_schema.ensure_schema(db_path)
    return db_path


def _seed(db_path: str) -> None:
    with _connect(db_path) as conn:
        conn.row_factory = sqlite3.Row
        cur = conn.execute(
            """
            INSERT INTO counterparties (code_1c, name, name_normalized, inn, kpp)
            VALUES ('00-1', 'РОМАШКА ООО', 'ромашка ооо', '7701000001', '770101001')
            """
        )
        client_id = int(cur.lastrowid)
        conn.execute(
            """
            INSERT INTO KP_offers (
                kp_id, creation_date, customer_name, manager_name, counterparty_id
            ) VALUES (1, '2026-09-01', 'РОМАШКА ООО', 'Иван Иванов', ?)
            """,
            (client_id,),
        )
        conn.execute(
            "INSERT INTO kp_meta (kp_id, status, owner_user_id) VALUES (1, 'в архиве', 7)"
        )


def _buyer(**overrides: object) -> SupplyContractCreate:
    data: dict[str, object] = {
        "legal_form": "ooo",
        "full_name": "Ромашка",
        "short_name": "Ромашка",
        "signatory_position": "Генерального директора",
        "signatory_name": "Иванов Иван",
        "signatory_verb": "действующего",
        "authority_basis": "устав",
        "inn": "7604010010",
        "kpp": "760401001",
        "ogrn": "1027600000001",
        "legal_address": "Ярославль",
        "email": "buyer@example.com",
        "bank_name": "Промсвязьбанк",
        "account": "40702810000000000001",
        "corr_account": "30101810000000000760",
        "bik": "044525555",
        "contract_date": "2026-09-25",
    }
    data.update(overrides)
    return SupplyContractCreate.model_validate(data)


def _document_text(payload: bytes) -> str:
    document = Document(BytesIO(payload))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for section in document.sections:
        parts.extend(paragraph.text for paragraph in section.header.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def test_docx_contains_issued_number_buyer_inn_and_supplier_inn(tmp_path: Path) -> None:
    db_path = _fresh_db(tmp_path)
    _seed(db_path)
    service = SupplyContractService(db_path=db_path)
    created = service.create_for_kp(1, _buyer(), user=ADMIN)

    payload, _filename = service.document_for_kp(1, user=ADMIN)
    again, _filename = service.document_for_kp(1, user=ADMIN)
    text = _document_text(payload)

    assert created["number"] in text
    assert created["number"] in _document_text(again)
    assert "ДОГОВОР N /26" not in text
    assert "7604010010" in text
    assert "4401082520" in text
    assert "Типовая межотраслевая форма № М-2" in text
    assert "buyer@example.com" in text


def test_docx_ip_preamble_drops_hardcoded_ooo_sentence(tmp_path: Path) -> None:
    db_path = _fresh_db(tmp_path)
    _seed(db_path)
    service = SupplyContractService(db_path=db_path)
    service.create_for_kp(
        1,
        _buyer(
            legal_form="ip",
            full_name="Иванов Иван Иванович",
            short_name="Иванов И.И.",
            signatory_position=None,
            signatory_name="Иванов Иван Иванович",
            kpp=None,
            inn="760401001012",
            ogrn="304760000000012",
        ),
        user=ADMIN,
    )

    text = _document_text(service.document_for_kp(1, user=ADMIN)[0])

    assert "Индивидуальный предприниматель Иванов Иван Иванович" in text
    assert "Общество с ограниченной ответственностью ," not in text
    assert "4401082520" in text
