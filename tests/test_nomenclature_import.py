"""GPS-016: import-1c выбирает парсер .xls/.xlsx, partial не пишет disappeared."""

from __future__ import annotations

import sqlite3
from pathlib import Path

import pytest
from openpyxl import Workbook

from app.schemas.nomenclature import Import1cResponse
from app.services.nomenclature_import_service import (
    NomenclatureImportError,
    NomenclatureImportService,
)
from core.duplicate_candidates import list_candidates
from core.guid_demand import (
    FIELD_DUPLICATE,
    FIELD_GUID,
    STATE_RESOLVED,
    list_open as list_guid_demand,
    record_guid_demand,
)
from core.guid_gate import (
    HINT_CREATE,
    HINT_CREATE_U,
    MissingItem,
    OrderLine,
    REASON_MISSING,
    REASON_MISSING_U,
)
from core.nomenclature_guid import ensure_schema, get_by_mark, upsert
from core.nomenclature_sync import sync_pricelist
from core.price_queue_db import list_open as list_price_queue
from core.pricelist_1c_parser import PricelistRow

GUID_A = "11111111-1111-1111-1111-111111111111"
GUID_B = "22222222-2222-2222-2222-222222222222"
GUID_C = "33333333-3333-3333-3333-333333333333"
GUID_XLSX = "60d7ac63-68b6-11f1-9003-04d9f5387845"
MARK_12 = "С120.35-12"
MARK_13 = "С120.35-13"
NAME_12 = "Сваи С 120.35-12"

_PRICE_TABLE = {
    "pile": "pile_prices",
    "bridge_pile": "bridge_pile_prices",
    "fbs": "fbs_prices",
    "stair_flight": "march_prices",
    "stair_step": "step_prices",
}


def _seed(path: Path) -> None:
    conn = sqlite3.connect(path)
    try:
        ensure_schema(conn)
        upsert(conn, "pile", "С30.30-3", match_status="missing")
        upsert(
            conn,
            "pile",
            "С70.30-8",
            guid_1c=GUID_B,
            match_status="auto",
        )
        conn.commit()
    finally:
        conn.close()


def _xlsx(
    path: Path,
    *,
    with_filter: bool,
    name: str = "Сваи С 30.30-3",
    guid: str = GUID_XLSX,
    extra_names: list[tuple[str, str]] | None = None,
) -> Path:
    wb = Workbook()
    ws = wb.active
    ws["A2"] = "Параметры:"
    if with_filter:
        ws["A6"] = "Отбор:"
        ws["C6"] = f'Номенклатура Равно "{name}"'
    ws["A8"] = "Номенклатура"
    ws["D8"] = "УИД"
    ws["E8"] = "Код"
    ws["F8"] = "Вес (числитель)"
    ws["H8"] = "Объем (числитель)"
    ws["A9"] = name
    ws["D9"] = guid
    ws["E9"] = "00-00000001"
    ws["F9"] = 1
    ws["H9"] = 0.1
    row = 10
    for extra_name, extra_guid in extra_names or ():
        ws.cell(row, 1, extra_name)
        ws.cell(row, 4, extra_guid)
        ws.cell(row, 5, "00-00000002")
        row += 1
    ws.cell(row, 1, "Итого")
    wb.save(path)
    return path


def test_sync_partial_skips_disappeared(tmp_path: Path) -> None:
    db = tmp_path / "pb.db"
    _seed(db)
    conn = sqlite3.connect(db)
    try:
        report = sync_pricelist(
            conn,
            [
                PricelistRow(
                    guid=GUID_A,
                    name="Сваи С 30.30-3",
                    price=None,
                    unit="шт",
                    source_file="partial.xlsx",
                    row_index=9,
                )
            ],
            product_kind="pile",
            partial=True,
        )
        leftover = get_by_mark(conn, "pile", "С70.30-8")
    finally:
        conn.close()

    assert report.disappeared == ()
    assert leftover is not None
    assert leftover.guid_1c == GUID_B


def test_xlsx_partial_import_sets_mode(tmp_path: Path) -> None:
    db = tmp_path / "pb.db"
    _seed(db)
    xlsx = _xlsx(tmp_path / "report.xlsx", with_filter=True)
    service = NomenclatureImportService(db)
    result = service.import_file(xlsx.read_bytes(), "report.xlsx")

    assert isinstance(result, Import1cResponse)
    assert result.mode == "partial"
    assert result.product_kind == "pile"
    assert result.disappeared_count == 0
    assert result.new_guids_count == 1
    conn = sqlite3.connect(db)
    try:
        row = get_by_mark(conn, "pile", "С30.30-3")
        leftover = get_by_mark(conn, "pile", "С70.30-8")
    finally:
        conn.close()
    assert row is not None
    assert row.guid_1c == GUID_XLSX
    assert leftover is not None
    assert leftover.guid_1c == GUID_B


def test_xlsx_full_without_filter_can_mark_disappeared(tmp_path: Path) -> None:
    db = tmp_path / "pb.db"
    _seed(db)
    xlsx = _xlsx(tmp_path / "full.xlsx", with_filter=False)
    result = NomenclatureImportService(db).import_file(xlsx.read_bytes(), "full.xlsx")
    assert result.mode == "full"
    assert result.disappeared_count >= 1


def test_xlsx_mixed_kinds_require_product_kind(tmp_path: Path) -> None:
    db = tmp_path / "pb.db"
    _seed(db)
    xlsx = _xlsx(
        tmp_path / "mix.xlsx",
        with_filter=True,
        extra_names=[("Блоки ФБС 9.3.6-Т", GUID_A)],
    )
    try:
        NomenclatureImportService(db).import_file(xlsx.read_bytes(), "mix.xlsx")
        raised = False
    except NomenclatureImportError as exc:
        raised = True
        assert "product_kind" in str(exc)
    assert raised


def test_import_xls_still_returns_full_mode(tmp_path: Path) -> None:
    db = tmp_path / "pb.db"
    _seed(db)
    service = NomenclatureImportService(db)
    rows = [
        PricelistRow(
            guid=GUID_A,
            name="Сваи С 30.30-3",
            price=None,
            unit="шт",
            source_file="Прайс сваи.xls",
            row_index=3,
        )
    ]
    result = service._sync_and_commit(rows, "pile")
    assert result.mode == "full"
    assert result.new_guids_count == 1


def test_xlsx_import_writes_fbs_weight(tmp_path: Path) -> None:
    db = tmp_path / "pb.db"
    conn = sqlite3.connect(db)
    try:
        ensure_schema(conn)
        upsert(conn, "fbs", "ФБС 9.3.6-Т", match_status="missing")
        conn.commit()
    finally:
        conn.close()
    xlsx = _xlsx(
        tmp_path / "fbs.xlsx",
        with_filter=True,
        name="Блоки ФБС 9.3.6-Т",
        guid=GUID_XLSX,
    )
    # override weight/volume in the helper's defaults (1 / 0.1)
    from openpyxl import load_workbook

    wb = load_workbook(xlsx)
    ws = wb.active
    ws["F9"] = 350
    ws["H9"] = 0.146
    wb.save(xlsx)

    result = NomenclatureImportService(db).import_file(xlsx.read_bytes(), "fbs.xlsx")
    assert result.product_kind == "fbs"
    assert result.weights_updated == 1
    from core.product_weight_catalog import resolve_product_weight_kg

    assert resolve_product_weight_kg("ФБС 9.3.6-Т", "fbs", str(db)) == pytest.approx(350.0)


def _price_row(name: str, guid: str, *, row_index: int = 9) -> PricelistRow:
    return PricelistRow(
        guid=guid,
        name=name,
        price=None,
        unit="шт",
        source_file="Номенклатура с УИД сваи.xlsx",
        row_index=row_index,
    )


def _seed_price_mark(conn: sqlite3.Connection, kind: str, mark: str, price: float = 1500.0) -> None:
    table = _PRICE_TABLE[kind]
    conn.execute(f"CREATE TABLE IF NOT EXISTS {table} (mark TEXT NOT NULL, price REAL)")
    conn.execute(f"INSERT INTO {table} (mark, price) VALUES (?, ?)", (mark, price))


def _price_count(conn: sqlite3.Connection, kind: str) -> int:
    table = _PRICE_TABLE[kind]
    return int(conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0])


def _open_price_guids(db: Path) -> set[str]:
    conn = sqlite3.connect(db)
    try:
        return {item.guid for item in list_price_queue(conn)}
    finally:
        conn.close()


def _demand_item(
    kind: str,
    mark: str,
    *,
    reason: str = REASON_MISSING,
) -> MissingItem:
    hint = HINT_CREATE_U if reason == REASON_MISSING_U else HINT_CREATE
    return MissingItem(line=OrderLine(kind, mark), reason=reason, action_hint=hint)


def test_price_mark_receives_single_guid_without_touching_price(tmp_path: Path) -> None:
    db = tmp_path / "pb.db"
    conn = sqlite3.connect(db)
    try:
        ensure_schema(conn)
        _seed_price_mark(conn, "pile", MARK_12)
        _seed_price_mark(conn, "pile", MARK_13)
        before = _price_count(conn, "pile")
        conn.commit()
        report = sync_pricelist(
            conn,
            [_price_row(NAME_12, GUID_A)],
            product_kind="pile",
            partial=True,
        )
        stored = get_by_mark(conn, "pile", MARK_12)
        neighbor = get_by_mark(conn, "pile", MARK_13)
        after = _price_count(conn, "pile")
    finally:
        conn.close()

    assert stored is not None
    assert stored.guid_1c == GUID_A
    assert stored.guid_1c_u is None
    assert stored.match_status == "auto"
    assert neighbor is None
    assert after == before
    assert report.unmatched_1c == ()
    assert report.disappeared == ()


def test_latin_c_and_u_suffix_use_the_same_candidate_index(tmp_path: Path) -> None:
    db = tmp_path / "pb.db"
    conn = sqlite3.connect(db)
    try:
        ensure_schema(conn)
        _seed_price_mark(conn, "pile", MARK_12)
        conn.commit()
        sync_pricelist(
            conn,
            [_price_row("Сваи C 120.35-12у", GUID_A)],
            product_kind="pile",
        )
        stored = get_by_mark(conn, "pile", MARK_12)
        plain = get_by_mark(conn, "pile", MARK_13)
    finally:
        conn.close()

    assert plain is None
    assert stored is not None
    assert stored.guid_1c is None
    assert stored.guid_1c_u == GUID_A
    assert stored.match_status == "auto"


@pytest.mark.parametrize(
    ("kind", "mark"),
    [
        ("bridge_pile", "С8.35-Т1"),
        ("fbs", "ФБС 9.3.6-Т"),
        ("stair_flight", "1ЛМ27.11.14-4"),
        ("stair_step", "ЛС12"),
    ],
)
def test_single_guid_binds_mark_from_each_price_table(
    tmp_path: Path,
    kind: str,
    mark: str,
) -> None:
    db = tmp_path / "pb.db"
    conn = sqlite3.connect(db)
    try:
        ensure_schema(conn)
        _seed_price_mark(conn, kind, mark)
        before = _price_count(conn, kind)
        conn.commit()
        sync_pricelist(conn, [_price_row(mark, GUID_A)], product_kind=kind)
        stored = get_by_mark(conn, kind, mark)
        after = _price_count(conn, kind)
    finally:
        conn.close()

    assert stored is not None
    assert stored.guid_1c == GUID_A
    assert stored.match_status == "auto"
    assert after == before


def test_repeat_file_does_not_duplicate_catalog_row(tmp_path: Path) -> None:
    db = tmp_path / "pb.db"
    conn = sqlite3.connect(db)
    try:
        ensure_schema(conn)
        _seed_price_mark(conn, "pile", MARK_12)
        conn.commit()
        rows = [_price_row(NAME_12, GUID_A)]
        sync_pricelist(conn, rows, product_kind="pile")
        second = sync_pricelist(conn, rows, product_kind="pile")
        count = conn.execute(
            "SELECT COUNT(*) FROM nomenclature_guid WHERE product_kind = ? AND mark = ?",
            ("pile", MARK_12),
        ).fetchone()[0]
        stored = get_by_mark(conn, "pile", MARK_12)
    finally:
        conn.close()

    assert count == 1
    assert stored is not None
    assert stored.guid_1c == GUID_A
    assert second.new_guids == ()


def test_manual_guid_is_not_overwritten_from_price_door(tmp_path: Path) -> None:
    db = tmp_path / "pb.db"
    conn = sqlite3.connect(db)
    try:
        ensure_schema(conn)
        upsert(conn, "pile", MARK_12, guid_1c=GUID_B, match_status="manual")
        _seed_price_mark(conn, "pile", MARK_12)
        conn.commit()
        sync_pricelist(conn, [_price_row(NAME_12, GUID_A)], product_kind="pile")
        stored = get_by_mark(conn, "pile", MARK_12)
    finally:
        conn.close()

    assert stored is not None
    assert stored.guid_1c == GUID_B
    assert stored.match_status == "manual"


def test_import_does_not_open_price_queue_for_priced_mark(tmp_path: Path) -> None:
    db = tmp_path / "pb.db"
    conn = sqlite3.connect(db)
    try:
        ensure_schema(conn)
        _seed_price_mark(conn, "pile", MARK_12)
        _seed_price_mark(conn, "pile", "С140.35-12")
        _seed_price_mark(conn, "pile", "С160.35-12")
        before = _price_count(conn, "pile")
        conn.commit()
    finally:
        conn.close()

    rows = [
        _price_row(NAME_12, GUID_A, row_index=9),
        _price_row("Сваи С 140.35-12", GUID_B, row_index=10),
        _price_row("Сваи С 160.35-12", GUID_C, row_index=11),
    ]
    result = NomenclatureImportService(db)._sync_and_commit(rows, "pile", partial=True)
    conn = sqlite3.connect(db)
    try:
        after = _price_count(conn, "pile")
        marks = {
            mark: get_by_mark(conn, "pile", mark)
            for mark in (MARK_12, "С140.35-12", "С160.35-12")
        }
    finally:
        conn.close()

    assert result.unmatched_1c_count == 0
    assert result.disappeared_count == 0
    assert after == before
    assert _open_price_guids(db) == set()
    for mark, guid in ((MARK_12, GUID_A), ("С140.35-12", GUID_B), ("С160.35-12", GUID_C)):
        stored = marks[mark]
        assert stored is not None
        assert stored.guid_1c == guid
        assert stored.match_status == "auto"


def test_open_demand_mark_gets_guid_without_growing_price(tmp_path: Path) -> None:
    db = tmp_path / "pb.db"
    mark = "С99.40-8"
    conn = sqlite3.connect(db)
    try:
        ensure_schema(conn)
        conn.execute("CREATE TABLE pile_prices (mark TEXT NOT NULL, price REAL)")
        record_guid_demand(conn, 21, [_demand_item("pile", mark)])
        conn.commit()
        before = _price_count(conn, "pile")
    finally:
        conn.close()

    result = NomenclatureImportService(db)._sync_and_commit(
        [_price_row("Сваи С 99.40-8", GUID_A)],
        "pile",
        partial=True,
    )
    conn = sqlite3.connect(db)
    try:
        stored = get_by_mark(conn, "pile", mark)
        after = _price_count(conn, "pile")
        opened = list_guid_demand(conn)
        resolved = conn.execute(
            "SELECT state FROM guid_demand WHERE product_kind = ? AND mark = ? AND field = ?",
            ("pile", mark, FIELD_GUID),
        ).fetchone()
    finally:
        conn.close()

    assert result.unmatched_1c_count == 0
    assert stored is not None
    assert stored.guid_1c == GUID_A
    assert stored.match_status == "auto"
    assert after == before
    assert GUID_A not in _open_price_guids(db)
    assert opened == []
    assert resolved is not None
    assert resolved[0] == STATE_RESOLVED


def test_reinforced_demand_mark_writes_guid_1c_u(tmp_path: Path) -> None:
    db = tmp_path / "pb.db"
    conn = sqlite3.connect(db)
    try:
        ensure_schema(conn)
        record_guid_demand(
            conn,
            4,
            [_demand_item("pile", "С70.35-9у", reason=REASON_MISSING_U)],
        )
        conn.commit()
    finally:
        conn.close()

    NomenclatureImportService(db)._sync_and_commit(
        [_price_row("Сваи С 70.35-9у", GUID_A)],
        "pile",
    )
    conn = sqlite3.connect(db)
    try:
        stored = get_by_mark(conn, "pile", "С70.35-9")
        spelled = get_by_mark(conn, "pile", "С70.35-9у")
        opened = list_guid_demand(conn)
    finally:
        conn.close()

    assert spelled is None
    assert stored is not None
    assert stored.guid_1c_u == GUID_A
    assert stored.guid_1c is None
    assert opened == []


def test_draft_without_demand_does_not_create_mark(tmp_path: Path) -> None:
    db = tmp_path / "pb.db"
    conn = sqlite3.connect(db)
    try:
        ensure_schema(conn)
        conn.commit()
    finally:
        conn.close()

    result = NomenclatureImportService(db)._sync_and_commit(
        [_price_row("Сваи С 77.30-1", GUID_A)],
        "pile",
    )
    conn = sqlite3.connect(db)
    try:
        stored = get_by_mark(conn, "pile", "С77.30-1")
    finally:
        conn.close()

    assert stored is None
    assert result.unmatched_1c_count == 1
    assert GUID_A in _open_price_guids(db)


def test_two_guids_on_price_or_demand_mark_stay_ambiguous(tmp_path: Path) -> None:
    db = tmp_path / "pb.db"
    conn = sqlite3.connect(db)
    try:
        ensure_schema(conn)
        _seed_price_mark(conn, "pile", MARK_12)
        record_guid_demand(conn, 9, [_demand_item("pile", MARK_12)])
        conn.commit()
    finally:
        conn.close()

    rows = [
        _price_row(NAME_12, GUID_A, row_index=9),
        _price_row(NAME_12, GUID_B, row_index=10),
    ]
    result = NomenclatureImportService(db)._sync_and_commit(rows, "pile", partial=True)
    conn = sqlite3.connect(db)
    try:
        stored = get_by_mark(conn, "pile", MARK_12)
        candidates = list_candidates(conn, kind="pile", mark=MARK_12)
        opened = list_guid_demand(conn)
    finally:
        conn.close()

    assert stored is not None
    assert stored.guid_1c is None
    assert stored.guid_1c_u is None
    assert stored.match_status == "ambiguous"
    assert {item.guid for item in candidates} == {GUID_A, GUID_B}
    assert result.disappeared_count == 0
    assert _open_price_guids(db) == set()
    assert len(opened) == 1
    assert opened[0].mark == MARK_12
    assert opened[0].field == FIELD_DUPLICATE


def test_two_price_marks_for_one_name_do_not_receive_guid(tmp_path: Path) -> None:
    db = tmp_path / "pb.db"
    left = "С100.40-9"
    right = "С100.40-9 *"
    conn = sqlite3.connect(db)
    try:
        ensure_schema(conn)
        _seed_price_mark(conn, "pile", left)
        _seed_price_mark(conn, "pile", right)
        conn.commit()
    finally:
        conn.close()

    report_db = NomenclatureImportService(db)._sync_and_commit(
        [_price_row("Сваи С 100.40-9", GUID_A)],
        "pile",
    )
    conn = sqlite3.connect(db)
    try:
        stored_left = get_by_mark(conn, "pile", left)
        stored_right = get_by_mark(conn, "pile", right)
    finally:
        conn.close()

    assert report_db.unmatched_1c_count == 0
    assert GUID_A not in _open_price_guids(db)
    for stored in (stored_left, stored_right):
        assert stored is not None
        assert stored.guid_1c is None
        assert stored.guid_1c_u is None
        assert stored.match_status == "ambiguous"


def test_stranger_name_stays_unmatched_and_opens_price_queue(tmp_path: Path) -> None:
    db = tmp_path / "pb.db"
    conn = sqlite3.connect(db)
    try:
        ensure_schema(conn)
        upsert(conn, "pile", "С70.30-8", guid_1c=GUID_B, match_status="auto")
        conn.commit()
    finally:
        conn.close()

    xlsx = _xlsx(tmp_path / "report.xlsx", with_filter=True, name="Сваи С 999.99-9", guid=GUID_A)
    result = NomenclatureImportService(db).import_file(xlsx.read_bytes(), "report.xlsx")
    conn = sqlite3.connect(db)
    try:
        stored = get_by_mark(conn, "pile", "С999.99-9")
        leftover = get_by_mark(conn, "pile", "С70.30-8")
    finally:
        conn.close()

    assert stored is None
    assert leftover is not None
    assert leftover.guid_1c == GUID_B
    assert result.mode == "partial"
    assert result.disappeared_count == 0
    assert result.unmatched_1c_count == 1
    assert GUID_A in _open_price_guids(db)


def test_partial_price_door_does_not_mark_other_catalog_rows_disappeared(tmp_path: Path) -> None:
    db = tmp_path / "pb.db"
    conn = sqlite3.connect(db)
    try:
        ensure_schema(conn)
        upsert(conn, "pile", "С70.30-8", guid_1c=GUID_B, match_status="auto")
        _seed_price_mark(conn, "pile", MARK_12)
        conn.commit()
    finally:
        conn.close()

    xlsx = _xlsx(tmp_path / "report.xlsx", with_filter=True, name=NAME_12, guid=GUID_XLSX)
    result = NomenclatureImportService(db).import_file(xlsx.read_bytes(), "report.xlsx")
    conn = sqlite3.connect(db)
    try:
        stored = get_by_mark(conn, "pile", MARK_12)
        leftover = get_by_mark(conn, "pile", "С70.30-8")
    finally:
        conn.close()

    assert result.mode == "partial"
    assert result.disappeared_count == 0
    assert stored is not None
    assert stored.guid_1c == GUID_XLSX
    assert leftover is not None
    assert leftover.guid_1c == GUID_B
    assert GUID_XLSX not in _open_price_guids(db)
