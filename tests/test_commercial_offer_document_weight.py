"""PDF и XLSX: «общим весом» — расчётный вес строки, не затычка 0."""

from __future__ import annotations

import subprocess
import tempfile
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from core.commercial_offer import generate_commercial_offer_pdf
from core.commercial_offer_xlsx import generate_commercial_offer_xlsx
from core.kp_plate_weight import resolve_kp_line_weight_kg
from core.pile_catalog import PileCatalogEntry, upsert_pile_catalog
from core.product_weight_catalog import (
    ProductWeightRecord,
    normalize_product_mark,
    upsert_product_weights,
)

PILE_MARK = "С80.35-8у"
CATALOG_MARK = "С80.35"
CATALOG_UNIT_KG = 2500.0


def _pdf_weight(kg: float) -> str:
    """Тот же формат, что строка «общим весом» в PDF (запятая, 3 знака)."""
    return f"{kg:,.3f}".replace(",", "X").replace(".", ",").replace("X", " ")


def _xlsx_weight(kg: float) -> str:
    """Тот же формат, что строка «общим весом» в XLSX."""
    return f"{kg:,.3f}"


def _flat(text: str) -> str:
    return " ".join(text.split())


def _seed_catalog(tmp_path: Path, *, fbs_kg: float | None = None) -> str:
    from core.kp_db_schema import init_schema

    db_path = str(tmp_path / "catalog.db")
    init_schema(db_path)
    upsert_pile_catalog(
        db_path,
        [
            PileCatalogEntry(
                mark=CATALOG_MARK,
                length_m=8.0,
                section_mm=350,
                volume_m3=1.0,
                weight_kg=CATALOG_UNIT_KG,
                pcs_per_20t=7,
            )
        ],
    )
    if fbs_kg is not None:
        mark = "ФБС 9.3.6-Т"
        upsert_product_weights(
            db_path,
            [
                ProductWeightRecord(
                    mark=mark,
                    mark_norm=normalize_product_mark(mark),
                    product_type="fbs",
                    weight_kg=fbs_kg,
                    volume_m3=0.1,
                    display_name=mark,
                )
            ],
        )
    return db_path


def _pile_line(*, qty: int = 1, product_type: str = "piles", kind: str = "pile") -> dict[str, Any]:
    return {
        "mark": PILE_MARK,
        "name": PILE_MARK,
        "qty": qty,
        "concrete_grade": "B25",
        "unit_price": 1000.0,
        "product_type": product_type,
        "product_kind": kind,
    }


def _plate_line(*, qty: int = 1) -> dict[str, Any]:
    return {
        "name": "ПБ 60-12-8п",
        "length_m": 6.0,
        "width_m": 1.2,
        "qty": qty,
        "load_class": 800,
        "unit_price": 1000.0,
        "product_type": "plates",
    }


def _pdf_text(order: list[dict[str, Any]], **overrides: Any) -> str:
    buffer = generate_commercial_offer_pdf(
        order,
        offer_number="11",
        offer_date="21.10.2026",
        customer_name="ООО «АВТОВИД»",
        discount_percent=0,
        logistics_cost=0.0,
        **overrides,
    )
    with tempfile.NamedTemporaryFile(suffix=".pdf") as handle:
        handle.write(buffer.getvalue())
        handle.flush()
        result = subprocess.run(
            ["pdftotext", "-layout", handle.name, "-"],
            check=True,
            capture_output=True,
        )
    return result.stdout.decode("utf-8", errors="replace")


def _xlsx_sheet(order: list[dict[str, Any]], **overrides: Any) -> Worksheet:
    buffer = generate_commercial_offer_xlsx(
        order,
        offer_number="11",
        offer_date="21.10.2026",
        customer_name="ООО «АВТОВИД»",
        discount_percent=0,
        logistics_cost=0.0,
        **overrides,
    )
    return load_workbook(buffer)["КП"]


def _xlsx_text(sheet: Worksheet) -> str:
    parts: list[str] = []
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is not None:
                parts.append(str(cell.value))
    return "\n".join(parts)


def _assert_weight_sentence(pdf_text: str, xlsx_text: str, kg: float) -> None:
    pdf_sentence = f"общим весом {_pdf_weight(kg)} кг"
    xlsx_sentence = f"общим весом {_xlsx_weight(kg)} кг"
    assert pdf_sentence in _flat(pdf_text)
    assert xlsx_sentence in xlsx_text
    assert "общим весом 0,000" not in _flat(pdf_text)
    assert "общим весом 0.000" not in xlsx_text


def test_pile_only_pdf_and_xlsx_show_catalog_weight(tmp_path: Path) -> None:
    db_path = _seed_catalog(tmp_path)
    order = [_pile_line(qty=1)]

    pdf_text = _pdf_text(order, pile_catalog_db_path=db_path)
    xlsx_text = _xlsx_text(_xlsx_sheet(order, pile_catalog_db_path=db_path))

    _assert_weight_sentence(pdf_text, xlsx_text, CATALOG_UNIT_KG)


def test_plate_only_pdf_and_xlsx_keep_plate_formula_weight(tmp_path: Path) -> None:
    db_path = _seed_catalog(tmp_path)
    plate = _plate_line(qty=2)
    _, plate_kg = resolve_kp_line_weight_kg(plate)
    assert plate_kg > 0

    pdf_text = _pdf_text([plate], pile_catalog_db_path=db_path)
    xlsx_text = _xlsx_text(_xlsx_sheet([plate], pile_catalog_db_path=db_path))

    _assert_weight_sentence(pdf_text, xlsx_text, plate_kg)


def test_mixed_plates_and_piles_sum_formula_and_catalog_weight(tmp_path: Path) -> None:
    db_path = _seed_catalog(tmp_path)
    plate = _plate_line(qty=1)
    pile = _pile_line(qty=1)
    _, plate_kg = resolve_kp_line_weight_kg(plate)
    expected = plate_kg + CATALOG_UNIT_KG

    order = [plate, pile]
    pdf_text = _pdf_text(order, pile_catalog_db_path=db_path)
    xlsx_text = _xlsx_text(_xlsx_sheet(order, pile_catalog_db_path=db_path))

    _assert_weight_sentence(pdf_text, xlsx_text, expected)


def test_xlsx_embed_delivery_shows_same_catalog_weight(tmp_path: Path) -> None:
    db_path = _seed_catalog(tmp_path)
    order = [_pile_line(qty=1)]
    common = {
        "pile_catalog_db_path": db_path,
        "pile_logistics_cost": 1500.0,
    }

    plain = _xlsx_text(_xlsx_sheet(order, **common))
    embedded = _xlsx_text(
        _xlsx_sheet(order, embed_delivery_in_unit_price=True, **common)
    )

    sentence = f"общим весом {_xlsx_weight(CATALOG_UNIT_KG)} кг"
    assert sentence in plain
    assert sentence in embedded
    assert "общим весом 0.000" not in plain
    assert "общим весом 0.000" not in embedded


def test_fbs_pdf_and_xlsx_include_product_weight_catalog(tmp_path: Path) -> None:
    unit_kg = 350.0
    qty = 2
    db_path = _seed_catalog(tmp_path, fbs_kg=unit_kg)
    order = [
        {
            "mark": "ФБС 9.3.6-Т",
            "name": "ФБС 9.3.6-Т",
            "qty": qty,
            "concrete_grade": "B25",
            "unit_price": 100.0,
            "product_type": "fbs",
            "product_kind": "fbs",
        }
    ]

    pdf_text = _pdf_text(order, weight_catalog_db_path=db_path)
    xlsx_text = _xlsx_text(_xlsx_sheet(order, weight_catalog_db_path=db_path))

    _assert_weight_sentence(pdf_text, xlsx_text, unit_kg * qty)


def test_composite_pile_pdf_and_xlsx_stay_zero_without_weight_source(tmp_path: Path) -> None:
    """Составные сваи не имеют справочника веса в расчёте доставки."""
    db_path = _seed_catalog(tmp_path)
    order = [_pile_line(product_type="composite_piles", kind="composite_pile")]

    pdf_text = _pdf_text(order, pile_catalog_db_path=db_path)
    xlsx_text = _xlsx_text(_xlsx_sheet(order, pile_catalog_db_path=db_path))

    assert "общим весом 0,000 кг" in _flat(pdf_text)
    assert "общим весом 0.000 кг" in xlsx_text
    assert _pdf_weight(CATALOG_UNIT_KG) not in _flat(pdf_text)
    assert _xlsx_weight(CATALOG_UNIT_KG) not in xlsx_text


def test_bridge_steps_marches_and_kind_only_pile_use_catalog_weight(tmp_path: Path) -> None:
    """Мост, ЛС и ЛМ берут уже существующий справочник; составные сваи — нет."""
    from core.cargo_delivery_pricing import kp_document_total_weight_kg

    db_path = _seed_catalog(tmp_path)
    upsert_product_weights(
        db_path,
        [
            ProductWeightRecord(
                mark="ЛС-14",
                mark_norm=normalize_product_mark("ЛС-14"),
                product_type="steps",
                weight_kg=150.0,
                volume_m3=0.06,
                display_name="ЛС-14",
            ),
            ProductWeightRecord(
                mark="ЛМ28-12",
                mark_norm=normalize_product_mark("ЛМ28-12"),
                product_type="marches",
                weight_kg=400.0,
                volume_m3=0.2,
                display_name="ЛМ28-12",
            ),
        ],
    )
    kind_only = _pile_line(qty=1)
    del kind_only["product_type"]
    order = [
        _pile_line(qty=2, product_type="bridge_piles", kind="bridge_pile"),
        {
            "mark": "ЛС14-Б",
            "name": "ЛС14-Б",
            "qty": 1,
            "unit_price": 10.0,
            "product_type": "steps",
            "product_kind": "step",
        },
        {
            "mark": "ЛМ28-12",
            "name": "ЛМ28-12",
            "qty": 3,
            "unit_price": 10.0,
            "product_type": "marches",
            "product_kind": "march",
        },
        kind_only,
        _pile_line(qty=5, product_type="composite_piles", kind="composite_pile"),
        {
            "mark": "С99.99-1",
            "name": "С99.99-1",
            "qty": 4,
            "unit_price": 10.0,
            "product_type": "piles",
            "product_kind": "pile",
        },
    ]
    expected = CATALOG_UNIT_KG * 2 + 150.0 + 400.0 * 3 + CATALOG_UNIT_KG

    total = kp_document_total_weight_kg(
        order,
        pile_catalog_db_path=db_path,
        weight_catalog_db_path=db_path,
    )

    assert total == pytest.approx(expected)
