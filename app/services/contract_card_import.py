"""Разбор xlsx из трёх колонок: номер, имя, ИНН. Не лист юриста."""

from __future__ import annotations

import math
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from core.supply_contract import contract_date_from_number, normalize_contract_number

SHEET_NAME = "ИНН без сомнений"
_HEADER_NUMBER = "номер договора"
_HEADER_NAME = "имя контрагента"
_HEADER_INN = "инн"
_HEADER_SCAN_ROWS = 15


class ContractCardImportError(Exception):
    """Файл карточек не разобран. Сообщение по-русски."""


def parse_contract_card_sheet(data: bytes) -> tuple[list[dict[str, Any]], int]:
    """Карточки к вставке и сколько повторов номера внутри файла пропущено.

    Пустое имя или номер вне нормализации — ошибка на весь файл, без строк наружу.
    """
    raw_rows = _read_rows(data)
    header_at, columns = _header(raw_rows)
    ready, skipped = _collect(
        raw_rows[header_at + 1 :],
        columns,
        first_line=header_at + 2,
    )
    if not ready:
        raise ContractCardImportError("В листе нет строк договоров")
    return ready, skipped


def _read_rows(data: bytes) -> list[tuple[Any, ...]]:
    if not data:
        raise ContractCardImportError("Файл листа пуст")
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ContractCardImportError("Не удалось прочитать xlsx") from exc
    try:
        sheet = _pick_sheet(workbook)
        return [tuple(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _pick_sheet(workbook: Any):
    if SHEET_NAME in workbook.sheetnames:
        return workbook[SHEET_NAME]
    sheet = workbook.active
    if sheet is None:
        raise ContractCardImportError("В файле нет листа")
    return sheet


def _header(rows: list[tuple[Any, ...]]) -> tuple[int, dict[str, int]]:
    for index, raw in enumerate(rows[:_HEADER_SCAN_ROWS]):
        labels = [_cell_text(cell).casefold() for cell in raw]
        found = _match_columns(labels)
        if found is not None:
            return index, found
    raise ContractCardImportError(
        "Не найдены колонки: номер договора, имя контрагента, ИНН"
    )


def _match_columns(labels: list[str]) -> dict[str, int] | None:
    number = _exact(labels, _HEADER_NUMBER)
    name = _exact(labels, _HEADER_NAME)
    inn = _exact(labels, _HEADER_INN)
    if number is None or name is None or inn is None:
        return None
    return {"number": number, "name": name, "inn": inn}


def _exact(labels: list[str], expected: str) -> int | None:
    for index, label in enumerate(labels):
        if label == expected:
            return index
    return None


def _collect(
    rows: list[tuple[Any, ...]],
    columns: dict[str, int],
    *,
    first_line: int,
) -> tuple[list[dict[str, Any]], int]:
    ready: list[dict[str, Any]] = []
    seen: set[str] = set()
    skipped = 0
    for offset, raw in enumerate(rows):
        cells = _data_cells(raw, columns)
        if cells is None:
            continue
        item = _validated(*cells, line=first_line + offset)
        if item["number"] in seen:
            skipped += 1
            continue
        seen.add(str(item["number"]))
        ready.append(item)
    return ready, skipped


def _data_cells(
    raw: tuple[Any, ...],
    columns: dict[str, int],
) -> tuple[str, str, str] | None:
    number = _cell(raw, columns["number"])
    name = _cell(raw, columns["name"])
    inn = _cell(raw, columns["inn"])
    if not number and not name and not inn:
        return None
    return number, name, inn


def _validated(number: str, name: str, inn: str, line: int) -> dict[str, Any]:
    normalized = _number(number, line)
    imported_name = name.strip()
    if not imported_name:
        raise ContractCardImportError(f"Строка {line}: пустое имя контрагента")
    return {
        "imported_name": imported_name,
        "number": normalized,
        "inn": inn.strip() or None,
        "contract_date": contract_date_from_number(normalized),
    }


def _number(value: str, line: int) -> str:
    try:
        return normalize_contract_number(value)
    except ValueError as exc:
        raise ContractCardImportError(f"Строка {line}: {exc}") from exc


def _cell(raw: tuple[Any, ...], index: int) -> str:
    if index >= len(raw):
        return ""
    return _cell_text(raw[index])


def _cell_text(value: Any) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if math.isnan(value):
            return ""
        if value.is_integer():
            return str(int(value))
        return str(value).strip()
    return str(value).strip()
