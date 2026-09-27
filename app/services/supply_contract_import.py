"""Разбор листа юриста. Связка с контрагентом — отдельное действие."""

from __future__ import annotations

import re
from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from core.supply_contract import CONTRACT_STATUSES, SCAN_NOTES, contract_seq

_NOISE = frozenset(
    {
        "ооо",
        "ао",
        "зао",
        "пао",
        "оао",
        "ип",
        "кфх",
        "общество",
        "ограниченной",
        "ответственностью",
        "акционерное",
    }
)


class SupplyContractImportError(Exception):
    """Лист не разобран. Сообщение по-русски."""


def parse_lawyer_sheet(data: bytes) -> list[dict[str, Any]]:
    """Колонки листа: номер, дата, контрагент, менеджер, статус, скан."""
    if not data:
        raise SupplyContractImportError("Файл листа пуст")
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise SupplyContractImportError("Не удалось прочитать xlsx") from exc
    try:
        sheet = workbook.active
        if sheet is None:
            raise SupplyContractImportError("В файле нет листа")
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    header_at, columns = _header(rows)
    parsed: list[dict[str, Any]] = []
    seen: set[str] = set()
    for offset, raw in enumerate(rows[header_at + 1 :], start=header_at + 2):
        item = _parse_row(raw, columns, line=offset)
        if item is None:
            continue
        if item["number"] in seen:
            raise SupplyContractImportError(f"Номер {item['number']} повторяется в листе")
        seen.add(item["number"])
        parsed.append(item)
    if not parsed:
        raise SupplyContractImportError("В листе нет строк договоров")
    return parsed


def suggest_names(
    imported_name: str,
    counterparties: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """Похожие имена. Сами counterparty_id не проставляют."""
    key = _name_key(imported_name)
    if len(key) < 3:
        return []
    hits: list[dict[str, Any]] = []
    for party in counterparties:
        other = _name_key(str(party.get("name") or ""))
        if not other:
            continue
        if key in other or other in key:
            hits.append({"id": int(party["id"]), "name": str(party["name"])})
    return hits[:5]


def _header(rows: list[tuple[Any, ...]]) -> tuple[int, dict[str, int]]:
    for index, raw in enumerate(rows[:15]):
        labels = [_cell_text(cell) for cell in raw]
        found = _match_columns(labels)
        if found is not None:
            return index, found
    raise SupplyContractImportError(
        "Не найдены колонки: номер договора, дата, контрагент, менеджер, оригинал/ЭДО, наличие скана"
    )


def _match_columns(labels: list[str]) -> dict[str, int] | None:
    folded = [label.casefold().replace("ё", "е") for label in labels]
    number = _find(folded, exact=("номер договора",), contains=("номер",))
    when = _find(folded, exact=("дата", "дата договора"))
    name = _find(folded, exact=("контрагент",))
    manager = _find(folded, exact=("менеджер",))
    status = _find_status(folded)
    scan = _find(folded, exact=("наличие скана",), contains=("скан",))
    if None in (number, when, name, manager, status, scan):
        return None
    return {
        "number": int(number),
        "date": int(when),
        "name": int(name),
        "manager": int(manager),
        "status": int(status),
        "scan": int(scan),
    }


def _find(
    labels: list[str],
    *,
    exact: tuple[str, ...] = (),
    contains: tuple[str, ...] = (),
) -> int | None:
    for index, label in enumerate(labels):
        if label in exact:
            return index
    for index, label in enumerate(labels):
        if any(piece in label for piece in contains):
            return index
    return None


def _find_status(labels: list[str]) -> int | None:
    for index, label in enumerate(labels):
        if "оригинал" in label and "эдо" in label:
            return index
        if label in {"статус", "оригинал/эдо"}:
            return index
    return None


def _parse_row(
    raw: tuple[Any, ...],
    columns: dict[str, int],
    *,
    line: int,
) -> dict[str, Any] | None:
    number = _cell(raw, columns["number"])
    name = _cell(raw, columns["name"])
    if not number and not name:
        return None
    parsed_number = _number(number, line)
    parsed_date = _date(_raw(raw, columns["date"]), line)
    manager = _cell(raw, columns["manager"]) or "—"
    status = _status(_cell(raw, columns["status"]), line)
    scan = _scan(_cell(raw, columns["scan"]), line)
    return {
        "number": parsed_number,
        "contract_date": parsed_date.isoformat(),
        "imported_name": name or "—",
        "manager_name": manager,
        "status": status,
        "scan_note": scan,
    }


def _number(value: str, line: int) -> str:
    text = value.strip()
    if contract_seq(text) is None or text.count("/") != 2:
        raise SupplyContractImportError(f"Строка {line}: номер договора не распознан")
    return text


def _date(value: Any, line: int) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _cell_text(value)
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d.%m.%y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise SupplyContractImportError(f"Строка {line}: дата договора не распознана")


def _status(value: str, line: int) -> str:
    text = value.strip()
    if not text:
        return "нет"
    if text not in CONTRACT_STATUSES:
        raise SupplyContractImportError(f"Строка {line}: неизвестный статус «{text}»")
    return text


def _scan(value: str, line: int) -> str | None:
    text = value.strip()
    if not text:
        return None
    if text not in SCAN_NOTES:
        raise SupplyContractImportError(f"Строка {line}: неизвестная отметка скана «{text}»")
    return text


def _name_key(value: str) -> str:
    text = value.casefold().replace("ё", "е")
    tokens = [token for token in re.split(r"[^0-9a-zа-я]+", text) if token and token not in _NOISE]
    return " ".join(tokens)


def _cell(raw: tuple[Any, ...], index: int) -> str:
    return _cell_text(_raw(raw, index))


def _raw(raw: tuple[Any, ...], index: int) -> Any:
    if index >= len(raw):
        return None
    return raw[index]


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()
